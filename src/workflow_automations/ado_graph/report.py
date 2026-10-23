"""Output writers: export files (Parquet/CSV/Excel) and the analysis workbook."""

from __future__ import annotations

import datetime as dt
import json
import logging
import re
import time
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

from workflow_automations.ado_common.excel import excel_safe, keep_text_literal, publish_all

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink
except ImportError:
    openpyxl = None

log = logging.getLogger("ado_tool")


# Files each export format writes into the output directory.
EXPORT_FILES: dict[str, tuple[str, ...]] = {
    "parquet": ("nodes.parquet", "edges.parquet"),
    "csv": ("nodes.csv", "edges.csv"),
    "excel": ("work_items.xlsx",),
}


def known_formats(formats: list[str]) -> list[str]:
    """Return the known formats once each, in order, warning about unknown ones."""
    known: list[str] = []
    for fmt in formats:
        if fmt not in EXPORT_FILES:
            log.warning("Unknown export format '%s' – skipping.", fmt)
        elif fmt not in known:
            known.append(fmt)
    return known


def export_paths(output_dir: Path, formats: list[str]) -> list[Path]:
    """Paths an export with these formats publishes; unknown formats add none."""
    names = dict.fromkeys(name for fmt in formats for name in EXPORT_FILES.get(fmt, ()))
    return [output_dir / name for name in names]


NODE_KEY = "WI_ID"
EDGE_COLUMNS = ("Source_ID", "Target_ID", "Link_Type")

# An ISO-8601 timestamp that names its offset, as the ADO REST API returns dates.
_ZONED_TIMESTAMP = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})")


def _identity_text(value: dict) -> str | None:
    """``Display Name <unique name>`` for an ADO identity reference, else None."""
    name, unique = value.get("displayName"), value.get("uniqueName")
    if not isinstance(name, str) and not isinstance(unique, str):
        return None
    if isinstance(name, str) and isinstance(unique, str) and name != unique:
        return f"{name} <{unique}>"
    return name if isinstance(name, str) else unique


def _scalar(value: Any) -> Any:
    """Identities become text; other dicts and lists become sorted JSON."""
    if isinstance(value, dict):
        identity = _identity_text(value)
        if identity is not None:
            return identity
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, sort_keys=True, default=str)
    return value


def _is_zoned_date(value: Any) -> bool:
    if isinstance(value, dt.datetime):
        return value.tzinfo is not None
    return isinstance(value, str) and _ZONED_TIMESTAMP.fullmatch(value) is not None


def _normalize_column(column: pd.Series) -> pd.Series:
    if isinstance(column.dtype, pd.DatetimeTZDtype):
        return column.dt.tz_convert("UTC").dt.tz_localize(None)
    if not (pd.api.types.is_object_dtype(column) or pd.api.types.is_string_dtype(column)):
        return column
    present = column.dropna()
    if present.empty:
        return column
    if all(_is_zoned_date(v) for v in present):
        return pd.to_datetime(column, utc=True, format="ISO8601").dt.tz_localize(None)
    column = column.map(_scalar, na_action="ignore")
    present = column.dropna()
    if any(isinstance(v, str) for v in present) and not all(isinstance(v, str) for v in present):
        # Mixed text and numbers cannot share a Parquet column; keep them as text.
        column = column.map(str, na_action="ignore")
    return column


def _require_ids(frame: pd.DataFrame, name: str, columns: Sequence[str]) -> pd.DataFrame:
    for col in columns:
        if frame.empty:
            frame[col] = frame[col].astype("int64")
        elif not pd.api.types.is_integer_dtype(frame[col]) or frame[col].isna().any():
            raise ValueError(f"{name}: column {col} must hold integer work-item IDs")
    return frame


def prepare_export_frames(
    df_nodes: pd.DataFrame, df_edges: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Validate the export schema and normalize values every format can store.

    Raises ValueError before anything is written when a frame breaks the schema.
    """
    nodes = df_nodes.copy()
    if nodes.empty and NODE_KEY not in nodes.columns:
        nodes = pd.DataFrame(
            {NODE_KEY: pd.Series(dtype="int64"), "In_Scope": pd.Series(dtype=bool)}
        )
    if nodes.columns.duplicated().any():
        dupes = sorted(set(nodes.columns[nodes.columns.duplicated()]))
        raise ValueError(f"nodes: duplicate columns {dupes}")
    if NODE_KEY not in nodes.columns:
        raise ValueError(f"nodes: missing column {NODE_KEY}")
    nodes = _require_ids(nodes, "nodes", [NODE_KEY])
    if nodes[NODE_KEY].duplicated().any():
        dupes = sorted(nodes.loc[nodes[NODE_KEY].duplicated(), NODE_KEY].unique().tolist())
        raise ValueError(f"nodes: duplicate {NODE_KEY} values {dupes[:20]}")
    for col in nodes.columns:
        nodes[col] = _normalize_column(nodes[col])

    if tuple(df_edges.columns) != EDGE_COLUMNS:
        raise ValueError(
            f"edges: columns must be {list(EDGE_COLUMNS)}, got {list(df_edges.columns)}"
        )
    edges = _require_ids(df_edges.copy(), "edges", EDGE_COLUMNS[:2])
    edges["Link_Type"] = edges["Link_Type"].astype(str)
    return nodes, edges


# Excel's worksheet limits: rows include the heading row, text is per cell.
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
EXCEL_MAX_CELL_CHARS = 32_767


class ExcelLimitError(ValueError):
    """A sheet does not fit in an Excel worksheet; raised before anything is written."""


def excel_len(text: str) -> int:
    """Length of ``text`` as Excel counts it: UTF-16 code units, so an emoji counts 2."""
    return len(text.encode("utf-16-le")) // 2


def _long_texts(values: pd.Series, limit: int) -> pd.Series:
    """Excel length of each text value over ``limit``, indexed by row position.

    Python's length is a lower bound of Excel's and half of it an upper bound, so
    only values longer than ``limit // 2`` are measured exactly.
    """
    values = values.reset_index(drop=True)
    if isinstance(values.dtype, pd.CategoricalDtype):
        values = values.astype(object)
    if isinstance(values.dtype, pd.StringDtype):
        lengths = values.str.len()
    elif pd.api.types.is_object_dtype(values):
        # Object columns can mix text with numbers, lists or dates; only text counts.
        lengths = values.map(lambda v: len(v) if isinstance(v, str) else 0, na_action="ignore")
    else:
        return pd.Series(dtype="int64")
    suspects = values[lengths.fillna(0) > limit // 2]
    exact = suspects.map(excel_len)
    return exact[exact > limit].astype("int64")


def excel_overflow(frame: pd.DataFrame, check_text: bool = True) -> str | None:
    """Why ``frame`` (written with a heading row, no index) cannot be a worksheet, else None.

    Without ``check_text`` only rows and columns count, for callers that fit long
    text themselves (see LongText). Checking text length is one pass over the text
    and object columns, O(cells).
    """
    rows, columns = len(frame) + 1, len(frame.columns)
    if rows > EXCEL_MAX_ROWS:
        return f"{rows:,} rows with the heading exceed Excel's {EXCEL_MAX_ROWS:,}"
    if columns > EXCEL_MAX_COLUMNS:
        return f"{columns:,} columns exceed Excel's {EXCEL_MAX_COLUMNS:,}"
    if not check_text:
        return None
    for position, col in enumerate(frame.columns):
        values = frame.iloc[:, position]  # by position: step results can repeat a name
        if isinstance(values.dtype, pd.CategoricalDtype):
            values = values.cat.categories.to_series()
        long = _long_texts(values, EXCEL_MAX_CELL_CHARS)
        if not long.empty:
            return (
                f"column {col!r} holds {int(long.max()):,} characters in one cell, "
                f"over Excel's {EXCEL_MAX_CELL_CHARS:,}"
            )
    return None


# ---------------------------------------------------------------------------
# Long text: cells over the cell limit
# ---------------------------------------------------------------------------

LONG_TEXT_MODES = ("split", "trim", "keep")
LONG_TEXT_SHEET = "Long_Text"
# The smallest --max-cell-chars: a trimmed cell keeps some text beside its marker.
MIN_CELL_CHARS = 256
LONG_TEXT_COLUMNS = ("Sheet", "Cell", "Column", "Key", "Characters", "Parts")


@dataclass(frozen=True)
class LongText:
    """How a text cell longer than ``limit`` Excel characters is written.

    split: the cell keeps its start and a marker linking to its row in the
    ``Long_Text`` sheet, which holds the full text in ``Part_1``, ``Part_2``, ...
    cells of at most EXCEL_MAX_CELL_CHARS each. trim: only the start and a marker.
    keep: text is never changed, so a cell over Excel's own limit makes the sheet
    unwritable (export refuses it, analyze spills it to CSV). ``limit`` None means
    EXCEL_MAX_CELL_CHARS; keep ignores it.
    """

    mode: str = "split"
    limit: int | None = None

    @classmethod
    def from_args(cls, args: Any) -> LongText:
        return cls(getattr(args, "long_text", "split"), getattr(args, "max_cell_chars", None))

    @property
    def cell_limit(self) -> int:
        return EXCEL_MAX_CELL_CHARS if self.limit is None else self.limit

    def long_cells(self, frame: pd.DataFrame) -> int:
        """How many cells of ``frame`` this policy changes."""
        if self.mode == "keep":
            return 0
        return sum(
            len(_long_texts(frame.iloc[:, position], self.cell_limit))
            for position in range(frame.shape[1])
        )


def _cut(text: str, units: int) -> str:
    """The longest start of ``text`` within ``units`` UTF-16 units, never half an emoji."""
    if len(text) * 2 <= units:
        return text
    encoded = text.encode("utf-16-le")[: units * 2]
    if len(encoded) >= 2 and 0xD8 <= encoded[-1] <= 0xDB:  # a high surrogate ends it
        encoded = encoded[:-2]
    return encoded.decode("utf-16-le")


def _parts(text: str, size: int) -> list[str]:
    parts = []
    while text:
        part = _cut(text, size)
        parts.append(part)
        text = text[len(part) :]
    return parts


class LongTextFitter:
    """Fits a workbook's long text cells to a LongText policy, sheet by sheet.

    ``fit`` returns the sheet's frame with long cells trimmed and those cells, each
    with its Long_Text row to link to (None with trim); ``frame`` then holds the
    Long_Text rows, in fit order.
    """

    def __init__(self, policy: LongText) -> None:
        self.policy = policy
        self.rows: list[dict[str, Any]] = []
        self.links: list[tuple[str, str]] = []  # Long_Text Cell value -> source location

    def fit(
        self, sheet: str, frame: pd.DataFrame
    ) -> tuple[pd.DataFrame, list[tuple[int, int, str | None]]]:
        if self.policy.mode == "keep":
            return frame, []
        limit = self.policy.cell_limit
        links: list[tuple[int, int, str | None]] = []
        changed: dict[int, pd.Series] = {}
        trimmed = 0
        for position in range(frame.shape[1]):
            long = _long_texts(frame.iloc[:, position], limit)
            if long.empty:
                continue
            values = frame.iloc[:, position].astype(object).reset_index(drop=True)
            for row, length in long.items():
                text = values[row]
                cell = f"{get_column_letter(position + 1)}{row + 2}"
                if self.policy.mode == "split":
                    target = len(self.rows) + 2
                    marker = (
                        f" … [{length:,} characters; full text in {LONG_TEXT_SHEET} row {target}]"
                    )
                    self.rows.append(
                        {
                            "Sheet": sheet,
                            "Cell": cell,
                            "Column": str(frame.columns[position]),
                            "Key": _key(frame.iloc[row, 0]),
                            "Characters": int(length),
                            "parts": _parts(text, EXCEL_MAX_CELL_CHARS),
                        }
                    )
                    self.links.append((sheet, cell))
                    links.append((row + 2, position + 1, f"'{LONG_TEXT_SHEET}'!A{target}"))
                else:
                    marker = f" … [{length:,} characters; trimmed]"
                    links.append((row + 2, position + 1, None))
                values[row] = _cut(text, limit - excel_len(marker)) + marker
                trimmed += 1
            changed[position] = values.set_axis(frame.index)
        if not changed:
            return frame, []
        frame = frame.copy()
        for position, values in changed.items():
            frame.isetitem(position, values)
        where = f"; full text in sheet {LONG_TEXT_SHEET}" if self.policy.mode == "split" else ""
        log.warning(
            "Sheet '%s': %d cell(s) over %s characters trimmed%s.",
            sheet,
            trimmed,
            f"{limit:,}",
            where,
        )
        return frame, links

    def frame(self) -> pd.DataFrame:
        """The Long_Text sheet: one row per split cell, its text in Part_N columns."""
        width = max((len(row["parts"]) for row in self.rows), default=0)
        records = [
            {
                **{key: row[key] for key in LONG_TEXT_COLUMNS[:-1]},
                "Parts": len(row["parts"]),
                **{f"Part_{i}": part for i, part in enumerate(row["parts"], start=1)},
            }
            for row in self.rows
        ]
        columns = [*LONG_TEXT_COLUMNS, *(f"Part_{i}" for i in range(1, width + 1))]
        frame = pd.DataFrame(records, columns=columns)
        reason = excel_overflow(frame, check_text=False)
        if reason:
            raise ExcelLimitError(f"sheet {LONG_TEXT_SHEET}: {reason}")
        return frame

    def write(self, writer: Any, timings: dict[str, float] | None = None) -> None:
        """Add the Long_Text sheet to ``writer`` when any cell was split."""
        if not self.rows:
            return
        frame = excel_safe(self.frame())
        frame.to_excel(writer, sheet_name=LONG_TEXT_SHEET, index=False)
        worksheet = writer.sheets[LONG_TEXT_SHEET]
        # Parts are cut at arbitrary places, so one can start with "=" or be "#N/A".
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                _keep_literal(cell)
        start = time.perf_counter()
        for position, width in enumerate(_column_widths(LONG_TEXT_SHEET, frame), start=1):
            if width is not None:
                worksheet.column_dimensions[get_column_letter(position)].width = width
        if timings is not None:
            timings["sizing"] += time.perf_counter() - start
        _add_filter(worksheet)
        column = LONG_TEXT_COLUMNS.index("Cell") + 1
        _link_cells(
            worksheet,
            LONG_TEXT_SHEET,
            [
                (row, column, f"{_quoted(sheet)}!{cell}")
                for row, (sheet, cell) in enumerate(self.links, start=2)
            ],
        )
        log.info("  → sheet '%s' (%d long cells)", LONG_TEXT_SHEET, len(self.rows))


def _key(value: Any) -> Any:
    """The first column's value, which identifies the row (WI_ID, CR_ID, ...)."""
    if isinstance(value, str):
        return _cut(value, 255)
    if pd.api.types.is_scalar(value):
        return None if pd.isna(value) else value
    return _cut(str(value), 255)


def _quoted(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _export_dataframes(
    df_nodes: pd.DataFrame,
    df_edges: pd.DataFrame,
    output_dir: Path,
    formats: list[str],
    overwrite: bool = True,
    long_text: LongText | None = None,
) -> None:
    """Validate, normalize and publish every requested format as one staged bundle.

    ``overwrite`` defaults to True because the CLI asks for consent before the run.
    Raises ExcelLimitError before anything is written when Excel is requested and a
    sheet does not fit; Parquet and CSV have no such limit and keep every text whole.
    ``long_text`` (default split) decides how Excel holds text over its cell limit.
    """
    long_text = long_text or LongText()
    df_nodes, df_edges = prepare_export_frames(df_nodes, df_edges)
    known = known_formats(formats)
    if "excel" in known:
        for sheet, frame in (("Nodes", df_nodes), ("Edges", df_edges)):
            reason = excel_overflow(frame, check_text=long_text.mode == "keep")
            if reason:
                hint = "export with --formats parquet or csv instead"
                if long_text.mode == "keep" and "characters in one cell" in reason:
                    hint = "use --long-text split or trim, or " + hint
                raise ExcelLimitError(f"work_items.xlsx: sheet {sheet}: {reason}; {hint}")
    writers = {
        "nodes.parquet": lambda p: df_nodes.to_parquet(p, index=False),
        "edges.parquet": lambda p: df_edges.to_parquet(p, index=False),
        "nodes.csv": lambda p: df_nodes.to_csv(p, index=False),
        "edges.csv": lambda p: df_edges.to_csv(p, index=False),
        "work_items.xlsx": lambda p: _write_excel(df_nodes, df_edges, p, long_text),
    }
    paths = export_paths(output_dir, known)
    publish_all([(path, writers[path.name]) for path in paths], overwrite)
    for fmt in known:
        log.info("Exported %s.", fmt)


def _write_excel(
    df_nodes: pd.DataFrame,
    df_edges: pd.DataFrame,
    path: Path,
    long_text: LongText | None = None,
) -> None:
    fitter = LongTextFitter(long_text or LongText())
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, frame in (("Nodes", df_nodes), ("Edges", df_edges)):
            frame, links = fitter.fit(sheet, excel_safe(frame))
            frame.to_excel(writer, sheet_name=sheet, index=False)
            keep_text_literal(writer.sheets[sheet])
            _link_cells(writer.sheets[sheet], sheet, links)
        fitter.write(writer)


# Column widths come from the heading and at most this many evenly spaced rows.
SIZING_SAMPLE_ROWS = 1_000
MAX_COLUMN_WIDTH = 50


def _sample_rows(frame: pd.DataFrame) -> pd.DataFrame:
    """At most SIZING_SAMPLE_ROWS evenly spaced rows, always including the first and last."""
    limit = max(SIZING_SAMPLE_ROWS, 2)
    if len(frame) <= limit:
        return frame
    last = len(frame) - 1
    return frame.iloc[[i * last // (limit - 1) for i in range(limit)]]


def _column_widths(sheet: str, frame: pd.DataFrame) -> list[int | None]:
    """Width per column from the heading and a bounded row sample, capped at 50.

    A column whose values cannot be rendered as text keeps Excel's default width
    (None) and is reported as a warning instead of failing the report.
    """
    sample = _sample_rows(frame)
    widths: list[int | None] = []
    for position, col in enumerate(frame.columns):
        try:
            values = sample.iloc[:, position].map(str, na_action="ignore").str.len()
            longest = max(len(str(col)), int(values.max()) if values.notna().any() else 0)
        except Exception as exc:  # any __str__ can fail; report it and go on
            log.warning(
                "Sheet '%s': cannot size column %r (%s); keeping its width.", sheet, col, exc
            )
            widths.append(None)
            continue
        widths.append(min(longest + 2, MAX_COLUMN_WIDTH))
    return widths


def _spill_path(output_path: Path, sheet: str) -> Path:
    safe = re.sub(r"[^\w@.-]", "_", sheet)
    return output_path.with_name(f"{output_path.stem}.{safe}.csv")


def plan_analysis_report(
    results: dict[str, pd.DataFrame], output_path: Path, long_text: LongText | None = None
) -> dict[str, tuple[Path, str]]:
    """Sheets that do not fit in Excel, each with its spill CSV path and the reason.

    Long text makes a sheet spill only with ``long_text`` keep; otherwise its cells
    are fitted (see long_text_counts).
    """
    check_text = (long_text or LongText()).mode == "keep"
    spills: dict[str, tuple[Path, str]] = {}
    for sheet, frame in results.items():
        reason = excel_overflow(frame, check_text=check_text)
        if reason:
            spills[sheet] = (_spill_path(output_path, sheet), reason)
    return spills


def long_text_counts(
    results: dict[str, pd.DataFrame],
    spills: dict[str, tuple[Path, str]],
    long_text: LongText | None = None,
) -> dict[str, int]:
    """Result sheets whose long text cells ``long_text`` trims, with how many.

    Spilled sheets are absent: their CSV keeps every text whole.
    """
    long_text = long_text or LongText()
    counts = {
        sheet: long_text.long_cells(frame)
        for sheet, frame in results.items()
        if sheet not in spills
    }
    return {sheet: count for sheet, count in counts.items() if count}


def write_analysis_report(
    results: dict[str, pd.DataFrame],
    output_path: Path,
    overwrite: bool = True,
    timings: dict[str, float] | None = None,
    preamble: dict[str, pd.DataFrame] | None = None,
    long_text: LongText | None = None,
) -> None:
    """Publish one sheet per analysis step, with column widths sized to content.

    ``preamble`` sheets (the report's About, Config and Summary) come first and
    are never spilled; in them a ``Sheet`` column that names a sheet links to it.
    Every sheet gets a filter on its heading row, which stays frozen in view, and
    work-item URLs become hyperlinks (see _link_urls).

    A sheet that exceeds Excel's limits is written to a sibling CSV
    (``<report stem>.<sheet>.csv``) published with the workbook, and its sheet
    holds a pointer row instead. ``overwrite`` defaults to True because the CLI
    asks for consent before the run. That consent names only the report, so an
    existing spill file is replaced only along with an existing report; otherwise
    FileExistsError is raised before anything is written. ``timings``, when given,
    receives the seconds spent writing and sizing.

    Text over the cell limit is fitted by ``long_text`` (default split), whose
    Long_Text sheet comes last; with keep, such a sheet is spilled instead.
    """
    preamble = preamble or {}
    long_text = long_text or LongText()
    taken = {name[:31].lower() for name in results}
    clashes = [name for name in preamble if name[:31].lower() in taken]
    if long_text.mode == "split":
        clashes += [
            name for name in [*results, *preamble] if name[:31].lower() == LONG_TEXT_SHEET.lower()
        ]
    if clashes:
        raise ValueError(f"preamble sheets {clashes} share a name with a result sheet")
    spills = plan_analysis_report(results, output_path, long_text)
    if not output_path.exists():
        stray = [str(path) for path, _ in spills.values() if path.exists()]
        if stray:
            raise FileExistsError(
                f"{', '.join(stray)} exists and is not part of an existing {output_path.name}; "
                "remove it or choose another --output-file"
            )
    for sheet, (path, reason) in spills.items():
        log.warning("Sheet '%s': %s; writing it to %s instead.", sheet, reason, path.name)
    timings = timings if timings is not None else {}
    timings.setdefault("write", 0.0)
    timings.setdefault("sizing", 0.0)

    def spill_writer(frame: pd.DataFrame) -> Callable[[Path], None]:
        def write(path: Path) -> None:
            start = time.perf_counter()
            frame.to_csv(path, index=False)
            timings["write"] += time.perf_counter() - start

        return write

    targets: list[tuple[Path, Callable[[Path], None]]] = [
        (
            output_path,
            lambda p: _write_workbook(results, p, spills, timings, preamble, long_text),
        )
    ]
    targets += [(path, spill_writer(results[sheet])) for sheet, (path, _) in spills.items()]
    publish_all(targets, overwrite)
    log.info(
        "Analysis report written to %s (write %.2fs, sizing %.2fs)",
        output_path,
        timings["write"],
        timings["sizing"],
    )


def _write_workbook(
    results: dict[str, pd.DataFrame],
    path: Path,
    spills: dict[str, tuple[Path, str]] | None = None,
    timings: dict[str, float] | None = None,
    preamble: dict[str, pd.DataFrame] | None = None,
    long_text: LongText | None = None,
) -> None:
    spills = spills or {}
    preamble = preamble or {}
    fitter = LongTextFitter(long_text or LongText())
    timings = timings if timings is not None else {"write": 0.0, "sizing": 0.0}
    # Writing includes the save on leaving the ExcelWriter, where openpyxl serializes.
    started, sized = time.perf_counter(), timings["sizing"]
    sheet_names = [name[:31] for name in [*preamble, *results]]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in [*preamble.items(), *results.items()]:
            safe_name = sheet_name[:31]
            if sheet_name in spills and sheet_name not in preamble:
                spill, reason = spills[sheet_name]
                frame = pd.DataFrame(
                    [
                        {
                            "Spilled_To": spill.name,
                            "Rows": len(df),
                            "Columns": len(df.columns),
                            "Reason": reason,
                        }
                    ]
                )
            elif sheet_name in preamble:
                frame, links = df, []
            else:
                frame, links = fitter.fit(safe_name, df)
            frame.to_excel(writer, sheet_name=safe_name, index=False)
            start = time.perf_counter()
            worksheet = writer.sheets[safe_name]
            for position, width in enumerate(_column_widths(safe_name, frame), start=1):
                if width is not None:
                    worksheet.column_dimensions[get_column_letter(position)].width = width
            timings["sizing"] += time.perf_counter() - start
            _add_filter(worksheet)
            linked = _link_urls(worksheet, frame, safe_name)
            if sheet_name not in spills and sheet_name not in preamble:
                _link_cells(worksheet, safe_name, links, EXCEL_MAX_HYPERLINKS - linked)
            if sheet_name in preamble:
                _link_sheets(worksheet, frame, sheet_names)
                log.info("  → sheet '%s'", safe_name)
            elif sheet_name in spills:
                log.info("  → sheet '%s' (pointer to %s)", safe_name, spills[sheet_name][0].name)
            else:
                log.info("  → sheet '%s' (%d rows)", safe_name, len(df))
        fitter.write(writer, timings)
    timings["write"] += time.perf_counter() - started - (timings["sizing"] - sized)


# Excel opens at most this many hyperlinks per worksheet.
EXCEL_MAX_HYPERLINKS = 65_530
SHEET_LINK_COLUMN = "Sheet"


def _add_filter(worksheet: Any) -> None:
    """Filter buttons on the heading row, which stays in view while scrolling."""
    if worksheet.max_column and worksheet["A1"].value is not None:
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = "A2"


def _link_urls(worksheet: Any, frame: pd.DataFrame, sheet: str) -> int:
    """Make http(s) values in columns named ``*URL`` clickable; return how many.

    Excel ignores hyperlinks beyond EXCEL_MAX_HYPERLINKS per worksheet, so a sheet
    with more URL cells keeps them all as plain text and the skip is logged.
    """
    targets: list[tuple[int, int, str]] = []
    for position, col in enumerate(frame.columns, start=1):
        if not str(col).endswith("URL"):
            continue
        for row, value in enumerate(frame.iloc[:, position - 1], start=2):
            if isinstance(value, str) and value.startswith(("https://", "http://")):
                targets.append((row, position, value))
    if len(targets) > EXCEL_MAX_HYPERLINKS:
        log.warning(
            "Sheet '%s': %d work-item URLs exceed Excel's %d hyperlinks per sheet; "
            "keeping them as text.",
            sheet,
            len(targets),
            EXCEL_MAX_HYPERLINKS,
        )
        return 0
    for row, column, url in targets:
        cell = worksheet.cell(row=row, column=column)
        cell.hyperlink = url
        cell.style = "Hyperlink"
    return len(targets)


def _keep_literal(cell: Any) -> None:
    """Store text that openpyxl typed as a formula (``=...``) or error (``#N/A``) as text."""
    if cell.data_type in ("f", "e") and isinstance(cell.value, str):
        cell.data_type = "s"
        cell.quotePrefix = cell.value.startswith("=")


def _link_cells(
    worksheet: Any,
    sheet: str,
    targets: list[tuple[int, int, str | None]],
    budget: int | None = None,
) -> None:
    """Link each (row, column) cell to a location in this workbook (``'Sheet'!A1``).

    Each cell's text stays literal, never a formula, also without a location (None).
    Beyond ``budget`` (default EXCEL_MAX_HYPERLINKS) links, none is made and the
    skip is logged: the cells still name their target in text.
    """
    for row, column, _ in targets:
        _keep_literal(worksheet.cell(row=row, column=column))
    targets = [target for target in targets if target[2] is not None]
    budget = EXCEL_MAX_HYPERLINKS if budget is None else budget
    if len(targets) > budget:
        log.warning(
            "Sheet '%s': %d long-text links exceed the %d hyperlinks left; keeping them as text.",
            sheet,
            len(targets),
            budget,
        )
        return
    for row, column, location in targets:
        cell = worksheet.cell(row=row, column=column)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location)
        cell.style = "Hyperlink"


def _link_sheets(worksheet: Any, frame: pd.DataFrame, sheet_names: list[str]) -> None:
    """Link each ``Sheet`` value that names a sheet of this workbook to its first cell."""
    if SHEET_LINK_COLUMN not in frame.columns:
        return
    column = list(frame.columns).index(SHEET_LINK_COLUMN) + 1
    for row, name in enumerate(frame[SHEET_LINK_COLUMN], start=2):
        if name in sheet_names:
            cell = worksheet.cell(row=row, column=column)
            quoted = str(name).replace("'", "''")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{quoted}'!A1")
            cell.style = "Hyperlink"
