"""Unfold a Markdown outline (for example an edited xlsx-outline file) into an XLSX table.

Without --into the rows go to a new workbook. With --into, rows tagged by
`xlsx-outline --row-ids` are merged back into a copy of the original workbook.
"""

from __future__ import annotations

import argparse
import copy
import os
import re
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from workflow_automations.xlsx_outline import (
    FRONT_MATTER_KEY,
    ValidationError,
    escape,
    fill_down,
    fingerprint,
    read_table,
)

HEADING = re.compile(r"^ {0,3}(#{1,6})(?:[ \t]+(.*?))?(?:[ \t]+#+)?[ \t]*$")
BULLET = re.compile(r"^( *)(?:[-*+]|\d{1,9}[.)])(?:[ \t]+(.*))?$")
# CommonMark: a backslash before ASCII punctuation is an escape.
ESCAPED = re.compile(r"\\([!-/:-@\[-`{-~])")
ROW_TAG = re.compile(r"[ \t]*<!--[ \t]*rows?:[ \t]*([0-9][0-9 ,]*?)[ \t]*-->[ \t]*$")


@dataclass
class Item:
    text: str
    line: int
    children: list[Item] = field(default_factory=list)
    rows: list[int] = field(default_factory=list)


@dataclass
class Settings:
    levels: list[str] | None = None
    details: list[str] = field(default_factory=list)
    title: bool = False
    label_levels: bool = False
    blank_label: str = "(blank)"
    raw: bool = False
    sheet: str | None = None
    fingerprint: str | None = None
    fill_down: bool = False


def split_front_matter(lines: list[str]) -> tuple[dict, int]:
    """Return the xlsx-outline front matter settings and the first body line index."""
    if not lines or lines[0].strip() != "---":
        return {}, 0
    for end in range(1, len(lines)):
        if lines[end].strip() in {"---", "..."}:
            break
    else:
        raise ValidationError("Front matter starting on line 1 is not closed")
    try:
        data = yaml.safe_load("\n".join(lines[1:end])) or {}
    except yaml.YAMLError as exc:
        raise ValidationError(f"Front matter is not valid YAML: {exc}") from exc
    settings = data.get(FRONT_MATTER_KEY, {}) if isinstance(data, dict) else {}
    if not isinstance(settings, dict):
        raise ValidationError(f"Front matter key {FRONT_MATTER_KEY} must be a mapping")
    return settings, end + 1


def resolve_settings(front: dict, args) -> Settings:
    def names(value, key: str) -> list[str] | None:
        if value is None:
            return None
        if not isinstance(value, list) or not all(isinstance(v, str) and v for v in value):
            raise ValidationError(f"Front matter {key} must be a list of column names")
        return value

    settings = Settings(
        levels=names(front.get("levels"), "levels"),
        details=names(front.get("details"), "details") or [],
        title=bool(front.get("title", False)),
        label_levels=bool(front.get("label_levels", False)),
        blank_label=str(front.get("blank_label", "(blank)")),
        raw=bool(front.get("raw", False)),
        sheet=front.get("sheet") if isinstance(front.get("sheet"), str) else None,
        fingerprint=front.get("fingerprint") if isinstance(front.get("fingerprint"), str) else None,
        fill_down=bool(front.get("fill_down", False)),
    )
    # Command-line options override the front matter, for apps that drop it.
    if args.columns is not None:
        settings.levels = args.columns
    if args.details is not None:
        settings.details = args.details
    settings.title |= args.title
    settings.label_levels |= args.label_levels
    settings.raw |= args.raw
    settings.fill_down |= args.fill_down
    if args.blank_label is not None:
        settings.blank_label = args.blank_label
    columns = (settings.levels or []) + settings.details
    if len(set(columns)) != len(columns):
        raise ValidationError("Level and detail column names must be unique")
    return settings


def parse_outline(lines: list[str], offset: int, title: bool) -> Item:
    """Build an item tree from headings and nested bullets, in document order."""
    root = Item("", 0)
    headings: list[tuple[int, Item]] = []
    bullets: list[tuple[int, Item]] = []
    title_pending = title
    for number, raw in enumerate(lines[offset:], offset + 1):
        line = raw.expandtabs(4).rstrip()
        if not line.strip():
            continue
        rows: list[int] = []
        if tag := ROW_TAG.search(line):
            rows = [int(value) for value in re.split(r"[ ,]+", tag.group(1)) if value]
            line = line[: tag.start()]
        if match := HEADING.match(line):
            level = len(match.group(1))
            if title_pending:
                # The title heading names the document; its sections are top-level items.
                title_pending = False
                headings = [(level, root)]
                bullets = []
                continue
            while headings and headings[-1][0] >= level:
                headings.pop()
            item = Item(match.group(2) or "", number, rows=rows)
            (headings[-1][1] if headings else root).children.append(item)
            headings.append((level, item))
            bullets = []
        elif match := BULLET.match(line):
            indent = len(match.group(1))
            while bullets and bullets[-1][0] >= indent:
                bullets.pop()
            parent = bullets[-1][1] if bullets else headings[-1][1] if headings else root
            item = Item(match.group(2) or "", number, rows=rows)
            parent.children.append(item)
            bullets.append((indent, item))
        else:
            raise ValidationError(
                f"Line {number} is neither a heading nor a list item: {line.strip()[:60]}"
            )
    return root


def unescape(text: str, raw: bool) -> str:
    return text.strip() if raw else ESCAPED.sub(r"\1", text).strip()


def parse_details(text: str, settings: Settings) -> list[str | None] | None:
    """Split `Name: value; Name: value` back into detail columns, or None if it is not one."""
    if not settings.details:
        return None
    quote = (lambda value: value) if settings.raw else escape
    names = [quote(name) for name in settings.details]
    fields: list[tuple[int, str]] = []
    last = -1
    for part in split_fields(text):
        # The longest matching name wins when one name prefixes another.
        matches = [
            column
            for column, name in enumerate(names)
            if column > last and part.startswith(name + ": ")
        ]
        if matches:
            column = max(matches, key=lambda index: len(names[index]))
            fields.append((column, part[len(names[column]) + 2 :]))
            last = column
        elif fields:
            # Hand-written text may use a bare "; " inside a value.
            column, value = fields[-1]
            fields[-1] = (column, f"{value}; {part}")
        else:
            return None
    values: list[str | None] = [None] * len(names)
    for column, value in fields:
        values[column] = unescape(value, settings.raw) or None
    return values


def split_fields(text: str) -> list[str]:
    """Split on "; " unless the semicolon is backslash-escaped."""
    parts, start, index = [], 0, 0
    while index < len(text):
        if text[index] == "\\":
            index += 2
            continue
        if text.startswith("; ", index):
            parts.append(text[start:index])
            start = index = index + 2
            continue
        index += 1
    parts.append(text[start:])
    return parts


def unfold(root: Item, settings: Settings) -> tuple[list[str], list[tuple[int | None, list]]]:
    """Return column names and (source row tag or None, values) in document order."""
    rows: list[tuple[int | None, list[str | None], list[str | None]]] = []

    def value(item: Item, depth: int) -> str | None:
        text = item.text
        if settings.label_levels and settings.levels and depth < len(settings.levels):
            prefix = settings.levels[depth] if settings.raw else escape(settings.levels[depth])
            if text.startswith(prefix + ": "):
                text = text[len(prefix) + 2 :]
        # Compare before unescaping: the outline marks text equal to the blank label
        # with a leading backslash.
        text = text.strip()
        if text == settings.blank_label:
            return None
        if text == "\\" + settings.blank_label:
            return settings.blank_label
        text = unescape(text, settings.raw)
        if not text:
            raise ValidationError(f"Line {item.line} has an empty item")
        return text

    def walk(item: Item, path: list[str | None]) -> None:
        for child in item.children:
            details = None if child.children else parse_details(child.text, settings)
            if details is not None:
                rows.extend((tag, path, details) for tag in child.rows or [None])
                continue
            depth = len(path)
            if settings.levels is not None and depth >= len(settings.levels):
                raise ValidationError(
                    f"Line {child.line} is nested deeper than the "
                    f"{len(settings.levels)} level columns"
                )
            child_path = [*path, value(child, depth)]
            blank = [None] * len(settings.details)
            # A tagged parent also stands for sheet rows whose path ends at it.
            if child.rows or not child.children:
                rows.extend((tag, child_path, blank) for tag in child.rows or [None])
            walk(child, child_path)

    walk(root, [])
    if not rows:
        raise ValidationError("The outline has no items")
    depth = max(len(path) for _, path, _ in rows)
    levels = settings.levels or [f"Level {number}" for number in range(1, depth + 1)]
    table = [
        (tag, path + [None] * (len(levels) - len(path)) + details) for tag, path, details in rows
    ]
    return levels + settings.details, table


def validate_sheet_name(name: str) -> None:
    if not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name) or name[0] == "'":
        raise ValidationError("Invalid sheet name (maximum 31 characters, no \\ / * ? : [ ])")


def literal(ws, row: int, column: int, value: str | None) -> None:
    if value is None:
        ws.cell(row, column).value = None
        return
    if len(value.encode("utf-16-le")) // 2 > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
        raise ValidationError(
            "Text exceeds Excel's cell limit or contains invalid control characters"
        )
    cell = ws.cell(row, column)
    # Store literal text so values such as "=x" are never read as formulas.
    cell.value = value
    cell.data_type = "s"


def new_workbook(sheet: str, headings: list[str], rows):
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet
    for row_number, values in enumerate([headings, *(values for _, values in rows)], 1):
        for column, value in enumerate(values, 1):
            if value is not None:
                literal(ws, row_number, column, value)
    return workbook


def copy_cell(source, target) -> None:
    target._value = source._value
    target.data_type = source.data_type
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.hyperlink:
        target.hyperlink = copy.copy(source.hyperlink)
    if source.comment:
        target.comment = copy.copy(source.comment)


def merge(settings: Settings, headings: list[str], rows, into: Path, sheet: str | None):
    """Rewrite the original sheet's data rows in outline order; return workbook and counts."""
    if settings.levels is None:
        raise ValidationError("Merging needs level column names (front matter or --columns)")
    if not any(tag is not None for tag, _ in rows):
        raise ValidationError(
            "The outline has no row tags; write it with xlsx-outline --row-ids, "
            "or unfold it to a new workbook without --into"
        )
    title, original_headings, data = read_table(into, sheet or settings.sheet)
    columns = []
    for name in headings:
        if original_headings.count(name) != 1:
            raise ValidationError(f"Column {name!r} must appear exactly once in row 1 of {title}")
        columns.append(original_headings.index(name))
    selected = [[row[column] for column in columns] for row in data]
    if settings.fingerprint and settings.fingerprint != fingerprint(title, headings, selected):
        raise ValidationError(
            f"Sheet {title} changed since the outline was written; regenerate the outline"
        )
    for tag, _ in rows:
        if tag is not None and not 2 <= tag <= len(data) + 1:
            raise ValidationError(
                f"Row tag {tag} is outside the data rows 2-{len(data) + 1} of {title}"
            )

    workbook = load_workbook(into)
    ws = workbook[title]
    if any(area.max_row >= 2 for area in ws.merged_cells.ranges):
        workbook.close()
        raise ValidationError("Merged cells below the heading row are unsupported when merging")
    # Row presence comes from the stored cells, not cached values: a formula Excel never
    # calculated reads as blank in the cached view but is still data.
    present = {row for (row, _), cell in ws._cells.items() if row >= 2 and cell.value is not None}
    last = max([len(data) + 1, *present])
    selected += [[None] * len(columns)] * (last - 1 - len(selected))
    # Compare with what the outline showed, so filled-down parents stay blank in the sheet.
    shown = fill_down(selected, len(settings.levels)) if settings.fill_down else selected
    used = {tag for tag, _ in rows}
    displayed = [any(value is not None for value in row) for row in selected]
    # Rows the outline could not show keep all their data, after the outline rows.
    hidden = [
        number
        for number in range(2, last + 1)
        if number in present and number not in used and not displayed[number - 2]
    ]
    deleted = sum(
        1 for number in range(2, last + 1) if number not in used and displayed[number - 2]
    )
    # Detach the old data rows, then write them back in outline order.
    old: dict[int, dict[int, object]] = {}
    for (row, column), cell in list(ws._cells.items()):
        if 2 <= row <= last:
            old.setdefault(row, {})[column] = cell
            del ws._cells[(row, column)]
    heights = {row: ws.row_dimensions[row].height for row in range(2, last + 1)}
    updated = 0
    plan = [*rows, *((number, None) for number in hidden)]
    for target, (tag, values) in enumerate(plan, 2):
        if tag is not None:
            for column, cell in old.get(tag, {}).items():
                copy_cell(cell, ws.cell(target, column))
            ws.row_dimensions[target].height = heights.get(tag)
        if values is None:
            continue
        before = shown[tag - 2] if tag is not None else [None] * len(columns)
        changed = False
        for column, value, original in zip(columns, values, before, strict=True):
            if tag is not None and value == original:
                continue
            literal(ws, target, column + 1, value)
            changed = True
        updated += changed and tag is not None
    for row in range(len(plan) + 2, last + 1):
        ws.row_dimensions[row].height = None
    counts = {
        "rows": len(rows),
        "added": sum(1 for tag, _ in rows if tag is None),
        "updated": updated,
        "deleted": deleted,
        "kept_at_end": len(hidden),
    }
    return workbook, title, counts


def save(workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=".unfold-", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    temp = Path(name)
    try:
        workbook.save(temp)
        load_workbook(temp, read_only=True).close()
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def run(args) -> str:
    if args.output.suffix.lower() != ".xlsx":
        raise ValidationError("Output must be an .xlsx workbook")
    protected = [args.input] + ([args.into] if args.into else [])
    if any(args.output.resolve() == path.resolve() for path in protected):
        raise ValidationError("Output must be a new file, not the outline or --into workbook")
    if args.output.exists() and not args.overwrite:
        raise ValidationError(f"Output exists; pass --overwrite to replace it: {args.output}")
    if args.into is None:
        validate_sheet_name(args.sheet or "Outline")
    lines = args.input.read_text(encoding="utf-8-sig").splitlines()
    front, offset = split_front_matter(lines)
    settings = resolve_settings(front, args)
    root = parse_outline(lines, offset, settings.title)
    headings, rows = unfold(root, settings)
    if args.into is None:
        save(new_workbook(args.sheet or "Outline", headings, rows), args.output)
        return f"Unfolded {len(rows)} rows into {len(headings)} columns: {args.output}"
    workbook, title, counts = merge(settings, headings, rows, args.into, args.sheet)
    try:
        save(workbook, args.output)
    finally:
        workbook.close()
    return (
        f"Merged {counts['rows']} rows into {title}: {counts['updated']} updated, "
        f"{counts['added']} added, {counts['deleted']} deleted, "
        f"{counts['kept_at_end']} without outline values kept at the end: {args.output}"
    )


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    root.add_argument("--input", type=Path, required=True, help="Markdown outline")
    root.add_argument("--output", type=Path, required=True, help="New .xlsx workbook")
    root.add_argument(
        "--into",
        type=Path,
        help="Original workbook to merge row-tagged edits into (written to --output)",
    )
    root.add_argument(
        "--sheet",
        help="New sheet name (default: Outline); with --into, the sheet to merge into "
        "(default: the front matter sheet, else the first sheet)",
    )
    root.add_argument("--overwrite", action="store_true", help="Replace an existing --output")
    root.add_argument(
        "--columns", nargs="+", help="Level column names, outermost first (overrides front matter)"
    )
    root.add_argument(
        "--details", nargs="+", help="Detail column names in `Name: value` lines, in order"
    )
    root.add_argument("--title", action="store_true", help="The first heading is a title")
    root.add_argument(
        "--label-levels", action="store_true", help="Items carry a `Column: ` prefix to strip"
    )
    root.add_argument("--blank-label", help="Item text that stands for a blank cell")
    root.add_argument("--raw", action="store_true", help="Do not unescape Markdown characters")
    root.add_argument(
        "--fill-down",
        action="store_true",
        help="With --into: the outline was written with --fill-down",
    )
    return root


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        print(run(args))
        return 0
    except (ValidationError, UnicodeError) as exc:
        print(f"Validation error: {exc}", file=sys.stderr)
        return 2
    except (OSError, BadZipFile) as exc:
        print(f"IO error: {exc}", file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
