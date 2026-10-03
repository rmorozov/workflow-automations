"""Triage Azure DevOps work-item scope into a multi-sheet Excel decision report.

The tool has two layers. ``fetch_snapshot`` is the only code that talks to Azure
DevOps (Kerberos, on-premise). Everything else is a deterministic transformation of
a plain-data ``Snapshot`` plus an optional capacity workbook and an optional
previous report, so the whole report pipeline is testable offline.
"""

import argparse
import logging
import os
import sys
import tempfile
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger("ado_triage")

# ---------------------------------------------------------------------------
# Exit codes and errors
# ---------------------------------------------------------------------------

EXIT_OK = 0
EXIT_VALIDATION = 2
EXIT_IO = 3
EXIT_EMPTY = 4
EXIT_ADO = 5


class TriageError(ValueError):
    """Configuration or input validation failure (exit 2)."""


class AdoError(RuntimeError):
    """Azure DevOps connection, authentication or query failure (exit 5)."""


class EmptyQueryError(RuntimeError):
    """The query returned no work items; no report is written (exit 4)."""


# ---------------------------------------------------------------------------
# ADO names
# ---------------------------------------------------------------------------

REL_CHILD = "System.LinkTypes.Hierarchy-Forward"
REL_PREDECESSOR = "System.LinkTypes.Dependency-Reverse"  # this item depends on X
REL_SUCCESSOR = "System.LinkTypes.Dependency-Forward"  # X depends on this item
REL_BLOCKS_FORWARD = "Microsoft.VSTS.Common.Blocks-Forward"  # this item blocks X
REL_BLOCKS_REVERSE = "Microsoft.VSTS.Common.Blocks-Reverse"  # this item is blocked by X

RELATION_LABELS = {
    REL_CHILD: "Child",
    REL_PREDECESSOR: "Predecessor",
    REL_SUCCESSOR: "Successor",
    REL_BLOCKS_FORWARD: "Blocks",
    REL_BLOCKS_REVERSE: "Blocked By",
}
# Labels meaning "this item obstructs the destination item".
BLAST_RADIUS_LABELS = ("Successor", "Blocks")

FIELD_ID = "System.Id"
FIELD_PARENT = "System.Parent"
FIELD_WI_TYPE = "System.WorkItemType"
FIELD_AREA_PATH = "System.AreaPath"
FIELD_STATE = "System.State"
FIELD_TITLE = "System.Title"
FIELD_ASSIGNED_TO = "System.AssignedTo"
FIELD_CHANGED_DATE = "System.ChangedDate"
FIELD_ORIGINAL_ESTIMATE = "Microsoft.VSTS.Scheduling.OriginalEstimate"
FIELD_COMPLETED_WORK = "Microsoft.VSTS.Scheduling.CompletedWork"
FIELD_REMAINING_WORK = "Microsoft.VSTS.Scheduling.RemainingWork"
FIELD_EST_EXPECTED = "my.EstimationExpectedDate"
FIELD_EST_READY = "my.EstimationReadyDate"
FIELD_UAT_EXPECTED = "my.UATExpectedDate"
FIELD_UAT_READY = "my.UATReadyDate"

DATE_FIELDS = (
    FIELD_EST_EXPECTED,
    FIELD_EST_READY,
    FIELD_UAT_EXPECTED,
    FIELD_UAT_READY,
    FIELD_CHANGED_DATE,
)
HOUR_FIELDS = (FIELD_ORIGINAL_ESTIMATE, FIELD_COMPLETED_WORK, FIELD_REMAINING_WORK)

# ---------------------------------------------------------------------------
# Output column names
# ---------------------------------------------------------------------------

COL_IS_ROOT = "is_root_element"
COL_WEB_URL = "web_url"
COL_PARENT_TITLE = "parent_title"
COL_LEVEL = "Matched_Capacity_Level"
COL_CAPACITY = "Team_Capacity"
COL_MONTHLY_CAPACITY = "Monthly_Capacity"
COL_VARIANCE = "Capacity_Variance"
COL_OVER = "Is_OverCapacity"
COL_DAYS_SINCE_CHANGE = "Days_Since_Last_Change"
COL_BLOCKS_COUNT = "Blocks_Count"
COL_FLAGS = "Triage_Flags"
COL_OPP_FLAGS = "Opportunity_Flags"
COL_SCORE = "Attention_Score"
COL_ACTIONS = "Recommended_Actions"
COL_DELTA = "Delta_Status"
COL_PERSIST = "Flag_Persist_Count"
COL_LOAD_PCT = "Load_Pct"
COL_RAG = "RAG"
COL_TEAM_SCORE = "Team_Attention_Score"
UNDEFINED_CAPACITY = "UNDEFINED CAPACITY"

REL_COL_SOURCE_ID = "Source_ID"
REL_COL_SOURCE_TITLE = "Source_Title"
REL_COL_SOURCE_URL = "Source_Url"
REL_COL_SOURCE_AREA = "Source_AreaPath"
REL_COL_DEST_ID = "Destination_ID"
REL_COL_DEST_TITLE = "Destination_Title"
REL_COL_DEST_URL = "Destination_Url"
REL_COL_DEST_AREA = "Destination_AreaPath"
REL_COL_TYPE = "Relation_Type"
RELATION_COLUMNS = [
    REL_COL_SOURCE_ID,
    REL_COL_SOURCE_TITLE,
    REL_COL_SOURCE_URL,
    REL_COL_SOURCE_AREA,
    REL_COL_DEST_ID,
    REL_COL_DEST_TITLE,
    REL_COL_DEST_URL,
    REL_COL_DEST_AREA,
    REL_COL_TYPE,
]

CAPACITY_SHEET_NAME = "dev_per_area_capacity"
CAPACITY_AREA_COL = "AreaPath"
RAG_RED_PCT = 100.0
RAG_AMBER_PCT = 85.0
EXCEL_SHEET_NAME_LIMIT = 31
DATA_START_ROW = 4  # row 1 purpose, row 2 blank, row 3 header

# ---------------------------------------------------------------------------
# Plain-data snapshot (the boundary between network and pure logic)
# ---------------------------------------------------------------------------


@dataclass
class Snapshot:
    """Everything fetched from ADO, as plain data.

    ``items``: ``{"id": int, "fields": dict, "relations": [{"rel": str, "url": str}]}``.
    ``external``: fields (title/type/area) of linked items outside the query.
    ``parent_titles``: title per ``System.Parent`` id.
    ``warnings``: non-fatal degradations, surfaced on the Config sheet.
    """

    items: list[dict[str, Any]]
    external: dict[int, dict[str, Any]] = field(default_factory=dict)
    parent_titles: dict[int, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Settings:
    as_of: datetime
    base_url: str
    timebox_days: int = 60
    stale_days: int = 21
    agenda_size: int = 20
    capacity_months: int = 2
    capacity_window: tuple[str, ...] = ()
    profile: str = "full"

    @property
    def timebox_end(self) -> datetime:
        return self.as_of + timedelta(days=self.timebox_days)


def safe_int(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def relation_target_id(url: str | None) -> int | None:
    if not url:
        return None
    return safe_int(url.rstrip("/").split("/")[-1])


def work_item_url(base_url: str, wi_id: int) -> str:
    return f"{base_url.rstrip('/')}/_workitems/edit/{wi_id}"


# ---------------------------------------------------------------------------
# Network layer (Azure DevOps on-premise, Kerberos)
# ---------------------------------------------------------------------------


def chunker(seq: Sequence[int], size: int = 200) -> Iterator[list[int]]:
    for i in range(0, len(seq), size):
        yield list(seq[i : i + size])


def fetch_snapshot(base_url: str, query_id: str) -> Snapshot:
    """Execute the saved query and fetch the data the report needs.

    Missing core work items are fatal: a triage report that silently lacks part
    of the scope would look complete. Missing linked/parent titles only degrade
    labels and are recorded as warnings.
    """
    try:
        import requests
        from azure.devops.connection import Connection
        from azure.devops.exceptions import AzureDevOpsServiceError
        from msrest.authentication import Authentication
        from requests_kerberos import HTTPKerberosAuth
    except ImportError as exc:
        raise TriageError(
            "ADO access needs the 'ado-triage' and 'ado-kerberos' extras: "
            f"python -m pip install 'workflow-automations[ado-triage,ado-kerberos]' ({exc})"
        ) from exc

    class KerberosAuthentication(Authentication):
        def signed_session(self, session=None):
            session = super().signed_session(session)
            session.auth = HTTPKerberosAuth()
            return session

    service_errors = (AzureDevOpsServiceError, requests.RequestException)
    warnings: list[str] = []
    try:
        connection = Connection(base_url=base_url, creds=KerberosAuthentication())
        client = connection.clients.get_work_item_tracking_client()

        logger.info("Executing query %s", query_id)
        result = client.query_by_id(id=query_id)
        all_ids: set[int] = set()
        # Tree and one-hop queries return links; flat queries return items.
        if result.work_item_relations:
            for link in result.work_item_relations:
                for end in (link.source, link.target):
                    if end is not None:
                        all_ids.add(end.id)
        for item in result.work_items or []:
            all_ids.add(item.id)
        if not all_ids:
            raise EmptyQueryError(f"Query {query_id} returned no work items")

        items: list[dict[str, Any]] = []
        logger.info("Fetching %d work items", len(all_ids))
        for chunk in chunker(sorted(all_ids)):
            for wi in client.get_work_items(ids=chunk, expand="Relations"):
                items.append(
                    {
                        "id": wi.id,
                        "fields": dict(wi.fields or {}),
                        "relations": [{"rel": r.rel, "url": r.url} for r in wi.relations or []],
                    }
                )
        missing = all_ids - {item["id"] for item in items}
        if missing:
            raise AdoError(f"{len(missing)} queried work items could not be fetched")

        known = {item["id"] for item in items}
        external_ids = sorted(
            {
                target
                for item in items
                for rel in item["relations"]
                if rel["rel"] in RELATION_LABELS
                and (target := relation_target_id(rel["url"])) is not None
                and target not in known
            }
        )
        parent_ids = sorted(
            {
                pid
                for item in items
                if (pid := safe_int(item["fields"].get(FIELD_PARENT))) is not None
            }
        )

        def fetch_fields(ids: list[int], fields: list[str], what: str) -> dict[int, dict]:
            out: dict[int, dict] = {}
            for chunk in chunker(ids):
                try:
                    batch = client.get_work_items(ids=chunk, fields=fields, error_policy="omit")
                except service_errors as exc:
                    warnings.append(f"Could not fetch {what} for {len(chunk)} items: {exc}")
                    continue
                for wi in batch:
                    if wi is not None:
                        out[wi.id] = dict(wi.fields or {})
            unresolved = len(ids) - len(out)
            if unresolved:
                warnings.append(f"{unresolved} {what} unavailable (deleted or no permission)")
            return out

        logger.info("Fetching %d linked and %d parent items", len(external_ids), len(parent_ids))
        external = fetch_fields(
            external_ids, [FIELD_WI_TYPE, FIELD_TITLE, FIELD_AREA_PATH], "linked items"
        )
        parents = fetch_fields(parent_ids, [FIELD_TITLE], "parent titles")
    except service_errors as exc:
        raise AdoError(f"Azure DevOps request failed: {exc}") from exc

    return Snapshot(
        items=items,
        external=external,
        parent_titles={pid: f.get(FIELD_TITLE, "") for pid, f in parents.items()},
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Capacity
# ---------------------------------------------------------------------------


def _next_month(dt: datetime) -> datetime:
    return (
        dt.replace(year=dt.year + 1, month=1) if dt.month == 12 else dt.replace(month=dt.month + 1)
    )


def target_months(start_month: str | None, count: int, as_of: datetime) -> tuple[str, ...]:
    """``count`` YYYY-MM strings starting at ``start_month`` or the month after ``as_of``."""
    if count < 1:
        raise TriageError("--capacity-months must be a positive integer")
    if start_month:
        try:
            current = datetime.strptime(start_month, "%Y-%m")
        except ValueError as exc:
            raise TriageError(
                f"--capacity-start-month must be YYYY-MM, got {start_month!r}"
            ) from exc
    else:
        current = _next_month(as_of.replace(day=1, hour=0, minute=0, second=0, microsecond=0))
    months = []
    for _ in range(count):
        months.append(current.strftime("%Y-%m"))
        current = _next_month(current)
    return tuple(months)


def normalize_month_header(value: Any) -> str:
    """Excel often turns ``2026-10`` headers into datetimes; map both back to YYYY-MM."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    text = str(value).strip()
    try:
        return pd.to_datetime(text, format="%Y-%m").strftime("%Y-%m")
    except (ValueError, TypeError):
        pass
    try:
        return pd.to_datetime(text).strftime("%Y-%m")
    except (ValueError, TypeError):
        return text


def load_capacity(path: Path, months: Sequence[str]) -> dict[str, float]:
    """AreaPath -> summed capacity hours over exactly ``months``.

    Every requested month must exist: summing a partial or different window would
    silently understate or overstate capacity.
    """
    try:
        df = pd.read_excel(path, sheet_name=CAPACITY_SHEET_NAME)
    except ValueError as exc:  # missing sheet
        raise TriageError(f"{path}: {exc}") from exc
    if df.empty or len(df.columns) < 2:
        raise TriageError(f"{path}: '{CAPACITY_SHEET_NAME}' needs AreaPath plus month columns")
    df = df.rename(columns={df.columns[0]: CAPACITY_AREA_COL})
    df.columns = [CAPACITY_AREA_COL] + [normalize_month_header(c) for c in df.columns[1:]]
    if df.columns.duplicated().any():
        raise TriageError(f"{path}: duplicate month columns after normalization")

    absent = [m for m in months if m not in df.columns]
    if absent:
        available = [c for c in df.columns[1:]]
        raise TriageError(
            f"{path}: capacity months {absent} not found; available columns: {available}"
        )

    df[CAPACITY_AREA_COL] = df[CAPACITY_AREA_COL].astype(str).str.strip()
    df = df[df[CAPACITY_AREA_COL].ne("") & df[CAPACITY_AREA_COL].ne("nan")]
    duplicates = df.loc[df[CAPACITY_AREA_COL].duplicated(), CAPACITY_AREA_COL].tolist()
    if duplicates:
        raise TriageError(f"{path}: duplicate AreaPath rows: {duplicates}")

    values = df[list(months)].apply(pd.to_numeric, errors="coerce")
    if values.isna().all(axis=None):
        raise TriageError(f"{path}: no numeric capacity in months {list(months)}")
    totals = values.fillna(0).sum(axis=1)
    return dict(zip(df[CAPACITY_AREA_COL], totals.astype(float), strict=True))


def resolve_capacity(area_path: str, capacity: dict[str, float]) -> tuple[str, float | None]:
    """Closest configured ancestor (deepest first) of ``area_path``."""
    parts = str(area_path).split("\\")
    for depth in range(len(parts), 0, -1):
        candidate = "\\".join(parts[:depth])
        if candidate in capacity:
            return candidate, capacity[candidate]
    return area_path, None


# ---------------------------------------------------------------------------
# DataFrames
# ---------------------------------------------------------------------------


def _display_name(value: Any) -> Any:
    if isinstance(value, dict):
        return value.get("displayName", "")
    return value


def build_master(snapshot: Snapshot, capacity: dict[str, float], cfg: Settings) -> pd.DataFrame:
    rows = []
    for item in snapshot.items:
        row = dict(item["fields"])
        wi_id = item["id"]
        parent = safe_int(row.get(FIELD_PARENT))
        level, cap = resolve_capacity(row.get(FIELD_AREA_PATH, ""), capacity)
        row.update(
            {
                FIELD_ID: wi_id,
                # Root is a property of the item, not of the query shape.
                COL_IS_ROOT: parent is None,
                COL_WEB_URL: work_item_url(cfg.base_url, wi_id),
                COL_PARENT_TITLE: snapshot.parent_titles.get(parent, "") if parent else "",
                COL_LEVEL: level,
                COL_CAPACITY: cap,
                COL_MONTHLY_CAPACITY: cap / cfg.capacity_months if cap is not None else None,
            }
        )
        rows.append(row)

    df = pd.DataFrame(rows)
    for col in (*DATE_FIELDS, *HOUR_FIELDS, FIELD_ASSIGNED_TO, FIELD_STATE, FIELD_TITLE):
        if col not in df.columns:
            df[col] = None
    for col in DATE_FIELDS:
        # openpyxl rejects tz-aware datetimes; normalize everything to naive UTC.
        df[col] = pd.to_datetime(df[col], errors="coerce", utc=True).dt.tz_convert(None)
    for col in HOUR_FIELDS:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    for col in (COL_CAPACITY, COL_MONTHLY_CAPACITY):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df[FIELD_ASSIGNED_TO] = df[FIELD_ASSIGNED_TO].map(_display_name)
    df[COL_DAYS_SINCE_CHANGE] = (pd.Timestamp(cfg.as_of) - df[FIELD_CHANGED_DATE]).dt.days
    return df


def build_relations(snapshot: Snapshot, cfg: Settings) -> pd.DataFrame:
    own = {item["id"]: item["fields"] for item in snapshot.items}

    def info(wi_id: int) -> tuple[str, str]:
        fields = own.get(wi_id) or snapshot.external.get(wi_id)
        if fields is None:
            return "UNKNOWN", "UNKNOWN"
        return fields.get(FIELD_TITLE, ""), fields.get(FIELD_AREA_PATH, "")

    rows = []
    for item in snapshot.items:
        src_title, src_area = info(item["id"])
        for rel in item["relations"]:
            label = RELATION_LABELS.get(rel["rel"])
            dest = relation_target_id(rel["url"])
            if label is None or dest is None:
                continue
            dest_title, dest_area = info(dest)
            rows.append(
                [
                    item["id"],
                    src_title,
                    work_item_url(cfg.base_url, item["id"]),
                    src_area,
                    dest,
                    dest_title,
                    work_item_url(cfg.base_url, dest),
                    dest_area,
                    label,
                ]
            )
    return pd.DataFrame(rows, columns=RELATION_COLUMNS)


# ---------------------------------------------------------------------------
# Heuristic registry: single source of truth for every flag
# ---------------------------------------------------------------------------

Mask = Callable[[pd.DataFrame, Settings], pd.Series]


@dataclass(frozen=True)
class Heuristic:
    flag: str
    kind: str  # "risk" feeds score/persistence/resolution; "opportunity" never does
    weight: int
    action: str
    mask: Mask


def _est_debt_dates(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (
        df[FIELD_EST_READY].isna()
        & df[FIELD_EST_EXPECTED].notna()
        & (df[FIELD_EST_EXPECTED] < cfg.as_of)
    )


def _est_debt_hours(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (df[FIELD_ORIGINAL_ESTIMATE] == 0) & (df[FIELD_REMAINING_WORK] == 0)


def _over_estimate(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (df[FIELD_ORIGINAL_ESTIMATE] > 0) & (
        df[FIELD_COMPLETED_WORK] > df[FIELD_ORIGINAL_ESTIMATE]
    )


def _date_drift(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (
        df[FIELD_UAT_EXPECTED].notna()
        & (df[FIELD_UAT_EXPECTED] < cfg.timebox_end)
        & (df[FIELD_UAT_READY].isna() | (df[FIELD_UAT_READY] > df[FIELD_UAT_EXPECTED]))
    )


def _exceeds_month(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return df[COL_MONTHLY_CAPACITY].notna() & (df[FIELD_REMAINING_WORK] > df[COL_MONTHLY_CAPACITY])


def _blocks_others(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return df[COL_BLOCKS_COUNT] > 0


def _stale(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return df[COL_DAYS_SINCE_CHANGE].notna() & (df[COL_DAYS_SINCE_CHANGE] >= cfg.stale_days)


def _unassigned(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return df[FIELD_ASSIGNED_TO].isna() | (df[FIELD_ASSIGNED_TO].astype(str).str.strip() == "")


def _defer_candidate(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (
        (df[FIELD_COMPLETED_WORK] == 0)
        & (df[COL_BLOCKS_COUNT] == 0)
        & (df[FIELD_UAT_EXPECTED].isna() | (df[FIELD_UAT_EXPECTED] > cfg.timebox_end))
    )


def _quick_win(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (df[FIELD_COMPLETED_WORK] > 0) & df[FIELD_REMAINING_WORK].between(1, 20)


def _minor_progress(df: pd.DataFrame, cfg: Settings) -> pd.Series:
    return (df[FIELD_COMPLETED_WORK] > 0) & (df[FIELD_COMPLETED_WORK] < 16.0)


HEURISTICS: tuple[Heuristic, ...] = (
    Heuristic("DATE_DRIFT", "risk", 3, "Re-confirm or renegotiate UAT date", _date_drift),
    Heuristic("BLOCKS_OTHERS", "risk", 3, "Prioritize to unblock downstream items", _blocks_others),
    Heuristic("EXCEEDS_MONTH_CAP", "risk", 3, "Split CR or spread across months", _exceeds_month),
    Heuristic("EST_DEBT_DATE", "risk", 2, "Chase estimation owner; set hard date", _est_debt_dates),
    Heuristic("EST_DEBT_HOURS", "risk", 2, "Get hours estimate before committing", _est_debt_hours),
    Heuristic("STALE", "risk", 2, "Ping owner; confirm item is alive", _stale),
    Heuristic("UNASSIGNED", "risk", 2, "Assign an owner", _unassigned),
    Heuristic(
        "OVER_ESTIMATE", "risk", 1, "Re-baseline estimate; review scope creep", _over_estimate
    ),
    Heuristic("DEFER_CANDIDATE", "opportunity", 1, "Propose cut or defer", _defer_candidate),
    Heuristic("QUICK_WIN", "opportunity", 1, "Pull into sprint to finish", _quick_win),
)


def annotate(df: pd.DataFrame, rel: pd.DataFrame, cfg: Settings) -> pd.DataFrame:
    """Add blocks count, risk/opportunity flags, attention score and actions."""
    df = df.copy()
    blocking = rel[rel[REL_COL_TYPE].isin(BLAST_RADIUS_LABELS)]
    counts = blocking.groupby(REL_COL_SOURCE_ID).size()
    df[COL_BLOCKS_COUNT] = df[FIELD_ID].map(counts).fillna(0).astype(int)

    risk = pd.Series("", index=df.index, dtype=object)
    opportunity = pd.Series("", index=df.index, dtype=object)
    actions = pd.Series("", index=df.index, dtype=object)
    score = np.zeros(len(df), dtype=int)
    for h in HEURISTICS:
        hit = h.mask(df, cfg).fillna(False).to_numpy(dtype=bool)
        actions = actions + np.where(hit, h.action + "; ", "")
        if h.kind == "risk":
            risk = risk + np.where(hit, h.flag + "|", "")
            score = score + np.where(hit, h.weight, 0)
        else:
            opportunity = opportunity + np.where(hit, h.flag + "|", "")

    df[COL_FLAGS] = risk.str.rstrip("|")
    df[COL_OPP_FLAGS] = opportunity.str.rstrip("|")
    df[COL_ACTIONS] = actions.str.rstrip("; ")
    df[COL_SCORE] = score
    return df


# ---------------------------------------------------------------------------
# Deltas against a previous report
# ---------------------------------------------------------------------------

RESOLVED_COLUMNS = [FIELD_ID, FIELD_TITLE, "Previous_Flags", "Resolution"]


def read_previous_master(path: Path) -> pd.DataFrame:
    """Master Data of an earlier report (this tool's layout or the legacy layout)."""
    sheets = pd.ExcelFile(path).sheet_names
    master = next((s for s in sheets if s.endswith("Master Data")), None)
    if master is None:
        raise TriageError(f"{path}: no '* Master Data' sheet; is this a previous report?")
    for header in (2, 0):  # current layout has a purpose line above the header
        prev = pd.read_excel(path, sheet_name=master, header=header)
        if FIELD_ID in prev.columns:
            prev[FIELD_ID] = pd.to_numeric(prev[FIELD_ID], errors="coerce")
            return prev.dropna(subset=[FIELD_ID]).astype({FIELD_ID: int})
    raise TriageError(f"{path}: Master Data sheet lacks a '{FIELD_ID}' column")


def apply_deltas(
    df: pd.DataFrame, previous: pd.DataFrame | None
) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    """Add Delta_Status and Flag_Persist_Count; list risks resolved since ``previous``."""
    df = df.copy()
    notes: list[str] = []
    flagged_now = df[COL_FLAGS] != ""
    if previous is None:
        df[COL_DELTA] = "N/A"
        df[COL_PERSIST] = flagged_now.astype(int)
        return df, pd.DataFrame(columns=RESOLVED_COLUMNS), notes

    df[COL_DELTA] = np.where(df[FIELD_ID].isin(previous[FIELD_ID]), "CARRIED", "NEW")
    if COL_FLAGS not in previous.columns:
        notes.append("Previous report predates flags; persistence restarts, no resolved list")
        df[COL_PERSIST] = flagged_now.astype(int)
        return df, pd.DataFrame(columns=RESOLVED_COLUMNS), notes

    prev = previous.set_index(FIELD_ID)
    prev_flags = prev[COL_FLAGS].fillna("").astype(str)
    prev_flagged_ids = prev_flags.index[prev_flags != ""]
    if COL_PERSIST in prev.columns:
        prev_persist = pd.to_numeric(prev[COL_PERSIST], errors="coerce").fillna(0)
    else:
        prev_persist = (prev_flags != "").astype(int)
    carried_streak = df[FIELD_ID].map(prev_persist).fillna(0).to_numpy()
    was_flagged = df[FIELD_ID].isin(prev_flagged_ids).to_numpy()
    df[COL_PERSIST] = np.where(flagged_now, np.where(was_flagged, carried_streak + 1, 1), 0).astype(
        int
    )

    current = df.set_index(FIELD_ID)[COL_FLAGS]
    rows = []
    for wi_id in prev_flagged_ids:
        if wi_id not in current.index:
            resolution = "Dropped from query"
        elif current.loc[wi_id] == "":
            resolution = "Flags cleared"
        else:
            continue
        title = prev.loc[wi_id, FIELD_TITLE] if FIELD_TITLE in prev.columns else ""
        rows.append([wi_id, title, prev_flags.loc[wi_id], resolution])
    return df, pd.DataFrame(rows, columns=RESOLVED_COLUMNS), notes


# ---------------------------------------------------------------------------
# Views
# ---------------------------------------------------------------------------

View = Callable[[pd.DataFrame, pd.DataFrame, dict[str, float], Settings], pd.DataFrame]


def _select(df: pd.DataFrame, mask: pd.Series, cols: list[str]) -> pd.DataFrame:
    return df.loc[mask.fillna(False).astype(bool), cols]


def view_capacity(df, rel, capacity, cfg) -> pd.DataFrame:
    stats = (
        df.groupby(COL_LEVEL)
        .agg(Total_CRs=(FIELD_ID, "count"), Total_Remaining_Work=(FIELD_REMAINING_WORK, "sum"))
        .reset_index()
    )
    cap = stats[COL_LEVEL].map(capacity)
    stats[COL_CAPACITY] = cap.astype(object).where(cap.notna(), UNDEFINED_CAPACITY)
    variance = cap - stats["Total_Remaining_Work"]
    stats[COL_VARIANCE] = variance.astype(object).where(cap.notna(), "N/A")
    stats[COL_OVER] = (variance < 0).astype(object).where(cap.notna(), "UNKNOWN")
    return stats.sort_values("Total_Remaining_Work", ascending=False)


def view_summary(df, rel, capacity, cfg) -> pd.DataFrame:
    out = view_capacity(df, rel, capacity, cfg).set_index(COL_LEVEL)
    cap = out.index.to_series().map(capacity)
    load = (out["Total_Remaining_Work"] / cap.where(cap > 0) * 100).round(1)
    out[COL_LOAD_PCT] = load
    out[COL_RAG] = np.select(
        [load.isna(), load > RAG_RED_PCT, load >= RAG_AMBER_PCT],
        ["GREY", "RED", "AMBER"],
        default="GREEN",
    )

    def per_team(mask: pd.Series, name: str) -> pd.Series:
        return (
            df.loc[mask.fillna(False).astype(bool)]
            .groupby(COL_LEVEL)[FIELD_ID]
            .count()
            .rename(name)
        )

    rollups = [
        per_team(_est_debt_dates(df, cfg) | _est_debt_hours(df, cfg), "Estimation_Debt_Count"),
        per_team(_date_drift(df, cfg), "Date_Drift_Risk_Count"),
        per_team(_exceeds_month(df, cfg), "Exceeds_Monthly_Cap_Count"),
        per_team(_stale(df, cfg), "Stale_Items_Count"),
        per_team(_unassigned(df, cfg), "Unassigned_Count"),
        df.groupby(COL_LEVEL)[COL_SCORE].sum().rename(COL_TEAM_SCORE),
    ]
    for series in rollups:
        out = out.join(series, how="left")
    names = [s.name for s in rollups]
    out[names] = out[names].fillna(0).astype(int)

    order = {"RED": 0, "AMBER": 1, "GREY": 2, "GREEN": 3}
    out = out.reset_index().rename(columns={COL_LEVEL: "Team"})
    out["_order"] = out[COL_RAG].map(order)
    return out.sort_values(["_order", COL_TEAM_SCORE], ascending=[True, False]).drop(
        columns="_order"
    )


def view_agenda(df, rel, capacity, cfg) -> pd.DataFrame:
    flagged = df[df[COL_SCORE] > 0].sort_values(
        [COL_SCORE, COL_PERSIST, FIELD_REMAINING_WORK], ascending=False
    )
    cols = [
        FIELD_ID,
        FIELD_TITLE,
        COL_LEVEL,
        FIELD_ASSIGNED_TO,
        FIELD_REMAINING_WORK,
        COL_SCORE,
        COL_PERSIST,
        COL_DELTA,
        COL_FLAGS,
        COL_OPP_FLAGS,
        COL_ACTIONS,
        COL_WEB_URL,
    ]
    return flagged.head(cfg.agenda_size)[cols].rename(columns={COL_LEVEL: "Team"})


def view_master(df, rel, capacity, cfg) -> pd.DataFrame:
    return df


def view_est_debt_dates(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_EST_EXPECTED, COL_PERSIST, COL_WEB_URL]
    return _select(df, _est_debt_dates(df, cfg), cols).sort_values(FIELD_EST_EXPECTED)


def view_est_debt_hours(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_STATE, COL_PERSIST, COL_WEB_URL]
    return _select(df, _est_debt_hours(df, cfg), cols)


def view_over_estimate(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [
        FIELD_ID,
        FIELD_TITLE,
        FIELD_AREA_PATH,
        FIELD_ORIGINAL_ESTIMATE,
        FIELD_COMPLETED_WORK,
        COL_WEB_URL,
    ]
    return _select(df, _over_estimate(df, cfg), cols)


def view_minor_progress(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_COMPLETED_WORK, FIELD_REMAINING_WORK]
    return _select(df, _minor_progress(df, cfg), [*cols, COL_WEB_URL]).sort_values(
        FIELD_COMPLETED_WORK
    )


def view_date_drift(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [
        FIELD_ID,
        FIELD_TITLE,
        FIELD_AREA_PATH,
        FIELD_UAT_EXPECTED,
        FIELD_UAT_READY,
        COL_PERSIST,
        COL_WEB_URL,
    ]
    return _select(df, _date_drift(df, cfg), cols).sort_values(FIELD_UAT_EXPECTED)


def view_blockers(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, COL_BLOCKS_COUNT, FIELD_REMAINING_WORK]
    return _select(df, _blocks_others(df, cfg), [*cols, COL_WEB_URL]).sort_values(
        COL_BLOCKS_COUNT, ascending=False
    )


def view_cross_team(df, rel, capacity, cfg) -> pd.DataFrame:
    cross = rel[
        (rel[REL_COL_TYPE] != "Child") & (rel[REL_COL_SOURCE_AREA] != rel[REL_COL_DEST_AREA])
    ]
    return cross.sort_values([REL_COL_SOURCE_AREA, REL_COL_DEST_AREA])


def view_deferred(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_REMAINING_WORK, FIELD_UAT_EXPECTED]
    return _select(df, _defer_candidate(df, cfg), [*cols, COL_WEB_URL])


def view_quick_wins(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_COMPLETED_WORK, FIELD_REMAINING_WORK]
    return _select(df, _quick_win(df, cfg), [*cols, COL_WEB_URL]).sort_values(FIELD_REMAINING_WORK)


def view_relations(df, rel, capacity, cfg) -> pd.DataFrame:
    return rel


def view_top_20_percent(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_REMAINING_WORK, FIELD_COMPLETED_WORK]
    work = df.loc[df[FIELD_REMAINING_WORK] > 0, [*cols, COL_WEB_URL]]
    work = work.sort_values([FIELD_AREA_PATH, FIELD_REMAINING_WORK], ascending=[True, False])
    groups = work.groupby(FIELD_AREA_PATH)
    keep = (groups[FIELD_ID].transform("count") * 0.2).astype(int).clip(lower=1)
    return work[groups.cumcount() < keep]


def view_exceeds_month(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_REMAINING_WORK, COL_MONTHLY_CAPACITY]
    return _select(df, _exceeds_month(df, cfg), [*cols, COL_WEB_URL]).sort_values(
        FIELD_REMAINING_WORK, ascending=False
    )


def view_unassigned(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_STATE, FIELD_REMAINING_WORK]
    return _select(df, _unassigned(df, cfg), [*cols, COL_WEB_URL])


def view_stale(df, rel, capacity, cfg) -> pd.DataFrame:
    cols = [FIELD_ID, FIELD_TITLE, FIELD_AREA_PATH, FIELD_STATE, COL_DAYS_SINCE_CHANGE]
    return _select(df, _stale(df, cfg), [*cols, FIELD_REMAINING_WORK, COL_WEB_URL]).sort_values(
        COL_DAYS_SINCE_CHANGE, ascending=False
    )


@dataclass(frozen=True)
class SheetSpec:
    name: str
    purpose: str
    view: View | None  # None: content supplied by the pipeline (Resolved)
    always: bool = False  # written even when empty
    exec_profile: bool = False


SHEET_CONFIG = "0. Config & Legend"
SHEET_SUMMARY = "1. Executive Summary"
SHEET_AGENDA = "2. Triage Agenda"
SHEET_MASTER = "3. Master Data"
SHEET_RESOLVED = "19. Resolved Since Last Report"

SHEETS: tuple[SheetSpec, ...] = (
    SheetSpec(
        SHEET_SUMMARY,
        "One row per team: load vs capacity, RAG status and rollup counts. Start here.",
        view_summary,
        always=True,
        exec_profile=True,
    ),
    SheetSpec(
        SHEET_AGENDA,
        "Top items by Attention_Score (sum of risk-flag weights). The triage meeting agenda.",
        view_agenda,
        always=True,
        exec_profile=True,
    ),
    SheetSpec(
        SHEET_MASTER,
        "Every fetched work item with all enrichment columns; source for the other sheets.",
        view_master,
        always=True,
    ),
    SheetSpec(
        "4. Capacity Triage",
        "Remaining work vs capacity per resolved team level.",
        view_capacity,
        exec_profile=True,
    ),
    SheetSpec(
        "5. Estimation Debt (Dates)",
        "Estimation promised by a past date, no ready date. Chase the estimation owner.",
        view_est_debt_dates,
    ),
    SheetSpec(
        "6. Estimation Debt (Hours)",
        "No original estimate and no remaining work. Get hours before committing scope.",
        view_est_debt_hours,
    ),
    SheetSpec(
        "7. Over Estimation",
        "Completed work exceeds original estimate. Re-baseline; check scope creep.",
        view_over_estimate,
    ),
    SheetSpec(
        "8. Minor Progress (<16h)",
        "Token progress (<16h logged): possibly stalled or trivially started.",
        view_minor_progress,
    ),
    SheetSpec(
        "9. Date Drift Risk",
        "UAT expected inside the timebox but not ready, or ready late.",
        view_date_drift,
    ),
    SheetSpec(
        "10. Top Blockers",
        "Items others wait on (Successor/Blocks links). Prioritize to unblock.",
        view_blockers,
    ),
    SheetSpec(
        "11. Cross-Team Deps",
        "Dependency links crossing AreaPath boundaries: coordination risk.",
        view_cross_team,
    ),
    SheetSpec(
        "12. Deferred Candidates",
        "No work started, blocking nothing, UAT beyond the timebox. Safest to cut.",
        view_deferred,
    ),
    SheetSpec(
        "13. Sunk Cost Quick Wins",
        "Work invested and 1-20h left. Pull into the sprint and finish.",
        view_quick_wins,
    ),
    SheetSpec(
        "14. All Relations",
        "Raw link graph: Child, Predecessor, Successor, Blocks, Blocked By.",
        view_relations,
    ),
    SheetSpec(
        "15. Top 20% Remaining Work",
        "Per AreaPath: the top 20% of items by remaining hours (at least one).",
        view_top_20_percent,
    ),
    SheetSpec(
        "16. Exceeds Monthly Capacity",
        "Single items larger than a month of team capacity. Split or spread.",
        view_exceeds_month,
    ),
    SheetSpec(
        "17. Unassigned Items",
        "Items with no owner. Unowned scope does not move.",
        view_unassigned,
    ),
    SheetSpec(
        "18. Stale Items",
        "No field change for --stale-days days: dormant scope nobody decided to defer.",
        view_stale,
    ),
    SheetSpec(
        SHEET_RESOLVED,
        "Risk-flagged in the previous report; now clean or gone. Closure visibility.",
        None,
        exec_profile=True,
    ),
)
for _spec in SHEETS:  # Excel rejects longer names; fail at import, not at write time.
    assert len(_spec.name) <= EXCEL_SHEET_NAME_LIMIT, _spec.name


# ---------------------------------------------------------------------------
# Report assembly and Excel output
# ---------------------------------------------------------------------------


@dataclass
class Report:
    master: pd.DataFrame
    relations: pd.DataFrame
    sheets: dict[str, pd.DataFrame]
    omitted: list[str]
    notes: list[str]


def build_report(
    snapshot: Snapshot,
    capacity: dict[str, float],
    cfg: Settings,
    previous: pd.DataFrame | None = None,
) -> Report:
    if not snapshot.items:
        raise EmptyQueryError("No work items to report")
    relations = build_relations(snapshot, cfg)
    master = annotate(build_master(snapshot, capacity, cfg), relations, cfg)
    master, resolved, notes = apply_deltas(master, previous)

    sheets: dict[str, pd.DataFrame] = {}
    omitted: list[str] = []
    for spec in SHEETS:
        if cfg.profile == "exec" and not spec.exec_profile:
            continue
        frame = resolved if spec.view is None else spec.view(master, relations, capacity, cfg)
        if frame.empty and not spec.always:
            omitted.append(spec.name)
            continue
        sheets[spec.name] = frame
    return Report(master, relations, sheets, omitted, [*snapshot.warnings, *notes])


RAG_FILLS = {
    "RED": "FFC7CE",
    "AMBER": "FFEB9C",
    "GREEN": "C6EFCE",
    "GREY": "D9D9D9",
}


def _write_frame(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame, purpose: str) -> None:
    frame.to_excel(writer, sheet_name=name, index=False, startrow=DATA_START_ROW - 2)
    ws = writer.sheets[name]
    ws.cell(row=1, column=1, value=purpose).font = Font(italic=True, color="808080")
    ws.freeze_panes = f"A{DATA_START_ROW}"
    for i, col in enumerate(frame.columns, start=1):
        width = len(str(col))
        if not frame.empty:
            longest = frame[col].astype(str).str.len().max()
            width = max(width, int(longest) if pd.notna(longest) else 0)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)


def _color_scale(ws, frame: pd.DataFrame, column: str, end_color: str) -> None:
    if column not in frame.columns or frame.empty:
        return
    letter = get_column_letter(frame.columns.get_loc(column) + 1)
    last = DATA_START_ROW + len(frame) - 1
    ws.conditional_formatting.add(
        f"{letter}{DATA_START_ROW}:{letter}{last}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=end_color),
    )


def _write_config(writer: pd.ExcelWriter, report: Report, meta: dict[str, str]) -> None:
    rows = [*meta.items(), ("Empty sheets omitted", ", ".join(report.omitted) or "(none)")]
    rows += [("Warning", note) for note in report.notes]
    settings = pd.DataFrame(rows, columns=["Setting", "Value"])
    legend = pd.DataFrame(
        [(h.flag, h.kind, h.weight, h.action) for h in HEURISTICS],
        columns=["Flag", "Kind", "Weight", "Recommended Action"],
    )
    purpose = "How this report was produced: inputs, thresholds, omissions, warnings, legend."
    _write_frame(writer, SHEET_CONFIG, settings, purpose)
    legend_row = DATA_START_ROW + len(settings) + 2
    legend.to_excel(writer, sheet_name=SHEET_CONFIG, index=False, startrow=legend_row)
    ws = writer.sheets[SHEET_CONFIG]
    ws.cell(
        row=legend_row,
        column=1,
        value="Flag legend: Attention_Score and persistence use risk flags only.",
    ).font = Font(bold=True)
    for letter, width in zip("ABCD", (28, 70, 12, 45), strict=True):
        ws.column_dimensions[letter].width = width


def write_excel(path: Path, report: Report, meta: dict[str, str]) -> None:
    purposes = {spec.name: spec.purpose for spec in SHEETS}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_config(writer, report, meta)
        for name, frame in report.sheets.items():
            _write_frame(writer, name, frame, purposes[name])
        if SHEET_SUMMARY in report.sheets:
            frame = report.sheets[SHEET_SUMMARY]
            ws = writer.sheets[SHEET_SUMMARY]
            col = frame.columns.get_loc(COL_RAG) + 1
            for row, value in enumerate(frame[COL_RAG], start=DATA_START_ROW):
                color = RAG_FILLS.get(value)
                if color:
                    ws.cell(row=row, column=col).fill = PatternFill(
                        start_color=color, end_color=color, fill_type="solid"
                    )
        if SHEET_AGENDA in report.sheets:
            _color_scale(
                writer.sheets[SHEET_AGENDA], report.sheets[SHEET_AGENDA], COL_SCORE, "FF9999"
            )
        if "18. Stale Items" in report.sheets:
            _color_scale(
                writer.sheets["18. Stale Items"],
                report.sheets["18. Stale Items"],
                COL_DAYS_SINCE_CHANGE,
                "FFC000",
            )


def publish(path: Path, write: Callable[[Path], None], overwrite: bool) -> None:
    """Write to a sibling temp file, then atomically replace ``path``."""
    if path.exists() and not overwrite:
        raise TriageError(f"{path} exists; pass --overwrite to replace it")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        write(tmp)
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    commands = root.add_subparsers(dest="mode", required=True)
    ex = commands.add_parser("export", help="Query ADO and write the triage report")
    ex.add_argument("--url", required=True, help="ADO collection URL")
    ex.add_argument("--query-id", required=True, help="Saved query GUID")
    ex.add_argument(
        "--output", "--file", dest="output", required=True, type=Path, help="Report path"
    )
    ex.add_argument("--overwrite", action="store_true", help="Replace an existing report")
    ex.add_argument("--format", choices=["excel", "csv"], default="excel")
    ex.add_argument("--capacity-file", type=Path, help="XLSX with 'dev_per_area_capacity'")
    ex.add_argument("--capacity-months", type=int, default=2, help="Months to sum (default 2)")
    ex.add_argument("--capacity-start-month", help="First month, YYYY-MM (default: next month)")
    ex.add_argument("--previous-file", type=Path, help="Earlier report XLSX for deltas")
    ex.add_argument("--profile", choices=["full", "exec"], default="full")
    ex.add_argument("--agenda-size", type=int, default=20)
    ex.add_argument("--stale-days", type=int, default=21)
    ex.add_argument("--timebox-days", type=int, default=60)
    ex.add_argument("--as-of", help="Report date YYYY-MM-DD (default: today)")
    return root


def _settings(args: argparse.Namespace) -> Settings:
    for name in ("agenda_size", "stale_days", "timebox_days", "capacity_months"):
        if getattr(args, name) < 1:
            raise TriageError(f"--{name.replace('_', '-')} must be a positive integer")
    if args.as_of:
        try:
            as_of = datetime.strptime(args.as_of, "%Y-%m-%d")
        except ValueError as exc:
            raise TriageError(f"--as-of must be YYYY-MM-DD, got {args.as_of!r}") from exc
    else:
        as_of = datetime.now().replace(microsecond=0)
    window = (
        target_months(args.capacity_start_month, args.capacity_months, as_of)
        if args.capacity_file
        else ()
    )
    return Settings(
        as_of=as_of,
        base_url=args.url,
        timebox_days=args.timebox_days,
        stale_days=args.stale_days,
        agenda_size=args.agenda_size,
        capacity_months=args.capacity_months,
        capacity_window=window,
        profile=args.profile,
    )


def run_export(args: argparse.Namespace) -> Report:
    cfg = _settings(args)
    output = args.output.resolve()
    for name in ("capacity_file", "previous_file"):
        source = getattr(args, name)
        if source is not None and source.resolve() == output:
            raise TriageError(f"--output must differ from --{name.replace('_', '-')}")
    if args.format == "csv" and (args.previous_file or args.profile != "full"):
        raise TriageError("--previous-file and --profile apply to Excel output only")

    # Validate every local input before touching the network.
    capacity = load_capacity(args.capacity_file, cfg.capacity_window) if args.capacity_file else {}
    previous = read_previous_master(args.previous_file) if args.previous_file else None
    if args.output.exists() and not args.overwrite:
        raise TriageError(f"{args.output} exists; pass --overwrite to replace it")

    snapshot = fetch_snapshot(args.url, args.query_id)
    report = build_report(snapshot, capacity, cfg, previous)
    meta = {
        "Generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "As of": cfg.as_of.strftime("%Y-%m-%d"),
        "ADO URL": args.url,
        "Query ID": args.query_id,
        "Capacity file": str(args.capacity_file or "(none)"),
        "Capacity window": ", ".join(cfg.capacity_window) or "(none)",
        "Previous report": str(args.previous_file or "(none)"),
        "Profile": cfg.profile,
        "Agenda size": str(cfg.agenda_size),
        "Stale threshold (days)": str(cfg.stale_days),
        "Timebox (days)": str(cfg.timebox_days),
        "RAG (Load_Pct)": f"RED > {RAG_RED_PCT:g}, AMBER >= {RAG_AMBER_PCT:g}, else GREEN",
    }
    if args.format == "csv":
        publish(args.output, lambda p: report.master.to_csv(p, index=False), args.overwrite)
    else:
        publish(args.output, lambda p: write_excel(p, report, meta), args.overwrite)
    return report


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s", stream=sys.stderr)
    try:
        report = run_export(args)
    except TriageError as exc:
        print(f"Validation error: {exc}", file=sys.stderr)
        return EXIT_VALIDATION
    except EmptyQueryError as exc:
        print(f"Nothing written: {exc}", file=sys.stderr)
        return EXIT_EMPTY
    except AdoError as exc:
        print(f"ADO error: {exc}", file=sys.stderr)
        return EXIT_ADO
    except OSError as exc:
        print(f"IO error: {exc}", file=sys.stderr)
        return EXIT_IO
    flagged = int((report.master[COL_SCORE] > 0).sum())
    print(f"Wrote {len(report.master)} work items ({flagged} risk-flagged) to {args.output}")
    for note in report.notes:
        print(f"Warning: {note}", file=sys.stderr)
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
