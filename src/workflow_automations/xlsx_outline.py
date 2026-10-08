"""Convert a hierarchical XLSX sheet into a Markdown outline without repeated values."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

MAX_HEADING_LEVEL = 6
FRONT_MATTER_KEY = "xlsx-outline"


class ValidationError(ValueError):
    """An invalid workbook, sheet, column selection, or option combination."""


@dataclass
class Node:
    # None is a blank cell; the blank label is applied only when rendering.
    key: str | None
    children: list[Node] = field(default_factory=list)
    # Children and detail rows in the order the sheet introduced them.
    items: list[Node | Detail] = field(default_factory=list)
    index: dict[str | None, Node] = field(default_factory=dict)
    # Sheet rows whose path ends here and that carry no details.
    rows: list[int] = field(default_factory=list)


@dataclass
class Detail:
    values: list[str | None]
    row: int


def text(value) -> str | None:
    """Render a cell value as one line of text; None for blank cells."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    # Line breaks inside a cell would end a Markdown heading or list item.
    result = re.sub(r"\s*[\r\n]+\s*", " ", str(value)).strip()
    return result or None


def escape(value: str) -> str:
    """Escape characters that would otherwise turn cell text into Markdown syntax."""
    value = re.sub(r"([\\`*_\[\]<>])", r"\\\1", value)
    # Leading markers only matter when they would start a heading, list or rule.
    if re.match(r"(#{1,6}|[+-])(\s|$)|-{3,}", value):
        return "\\" + value
    return re.sub(r"^(\d{1,9})([.)])(?=\s|$)", r"\1\\\2", value)


def read_table(path: Path, sheet: str | None) -> tuple[str, list[str], list[list[str | None]]]:
    """Return the sheet title, row-1 headings and rendered data rows (from sheet row 2)."""
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("Input must be an .xlsx workbook")
    # data_only reads the values Excel cached for formulas when the file was last saved.
    workbook = load_workbook(path, data_only=True)
    try:
        if sheet is None:
            ws = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            ws = workbook[sheet]
        else:
            raise ValidationError(f"Sheet does not exist: {sheet}")
        # Visit only stored cells: iter_rows() would create every cell up to the sheet's
        # dimensions, which stray formatting can stretch far beyond the data.
        cells = {
            coordinate: cell.value
            for coordinate, cell in ws._cells.items()
            if cell.value is not None
        }
        # A merged range displays its top-left value in every covered cell.
        for area in ws.merged_cells.ranges:
            value = cells.get((area.min_row, area.min_col))
            if value is not None:
                for row in range(area.min_row, area.max_row + 1):
                    for column in range(area.min_col, area.max_col + 1):
                        cells[(row, column)] = value
        if not cells:
            raise ValidationError("Sheet is empty")
        rows = max(row for row, _ in cells)
        columns = max(column for _, column in cells)
        headings = [text(cells.get((1, column))) or "" for column in range(1, columns + 1)]
        data = [
            [text(cells.get((row, column))) for column in range(1, columns + 1)]
            for row in range(2, rows + 1)
        ]
        return ws.title, headings, data
    finally:
        workbook.close()


def select_columns(
    headings: list[str], names: list[str] | None, indices: list[int] | None
) -> list[int]:
    """Return 0-based column positions in outline order."""
    if names:
        missing = [name for name in names if name not in headings]
        if missing:
            raise ValidationError(f"Unknown columns: {', '.join(missing)}")
        if any(headings.count(name) > 1 for name in names):
            raise ValidationError("Selected column headings must be unique")
        selected = [headings.index(name) for name in names]
    elif indices:
        if any(index < 1 or index > len(headings) for index in indices):
            raise ValidationError(f"Column indices must be between 1 and {len(headings)}")
        selected = [index - 1 for index in indices]
    else:
        selected = list(range(len(headings)))
    if len(set(selected)) != len(selected):
        raise ValidationError("A column is selected more than once")
    if any(not headings[column] for column in selected):
        raise ValidationError("Every selected column must have a nonempty heading in row 1")
    return selected


def fingerprint(sheet: str, names: list[str], rows) -> str:
    """Identify the selected values of a sheet, so a merge can detect later changes."""
    payload = json.dumps([sheet, names, rows], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def fill_down(rows, levels: int) -> list[list[str | None]]:
    """Give blank hierarchy cells the value above them, within the same parent."""
    filled = []
    previous: list[str | None] = [None] * levels
    for row in rows:
        values = list(row[:levels])
        for level in range(levels):
            # Fill only within the same parent, so a new group never inherits children.
            if values[level] is None and values[:level] == previous[:level]:
                values[level] = previous[level]
        previous = values
        filled.append(values + list(row[levels:]))
    return filled


def build_tree(rows, levels: int, fill: bool, group: bool) -> Node:
    """Fold rows into a tree; equal leading values share one parent node."""
    root = Node(None)
    active: list[Node] = []
    for number, row in enumerate(fill_down(rows, levels) if fill else rows, 2):
        values = list(row[:levels])
        details = row[levels:]
        while values and values[-1] is None:
            values.pop()
        if not values and all(value is None for value in details):
            continue
        node = root
        path: list[Node] = []
        continuing = True
        for level, value in enumerate(values):
            if group:
                child = node.index.get(value)
            else:
                # Without grouping, a node is reused only while the row continues the
                # previous row's path; any difference starts a new run below it.
                child = None
                if continuing and level < len(active) and active[level].key == value:
                    child = active[level]
                continuing = child is not None
            if child is None:
                child = Node(value)
                node.children.append(child)
                node.items.append(child)
                node.index.setdefault(value, child)
            path.append(child)
            node = child
        active = path
        if any(value is not None for value in details):
            node.items.append(Detail(details, number))
        else:
            node.rows.append(number)
    return root


def render(
    root: Node,
    headings: list[str],
    detail_headings: list[str],
    *,
    heading_levels: int,
    title: str | None,
    label_levels: bool,
    indent: int,
    raw: bool,
    blank_label: str,
    row_ids: bool = False,
) -> str:
    out: list[str] = []
    quote = (lambda value: value) if raw else escape
    offset = 1 if title else 0

    def block(line: str) -> None:
        if out and out[-1] != "":
            out.append("")
        out.append(line)
        out.append("")

    def details_line(values) -> str | None:
        parts = [
            f"{quote(name)}: {quote(value)}"
            for name, value in zip(detail_headings, values, strict=True)
            if value is not None
        ]
        return "; ".join(parts) or None

    def item_label(key: str | None) -> str:
        if key is None:
            return blank_label
        label = quote(key)
        # Keep text that reads like the blank label distinct from a blank cell.
        if label == blank_label and not raw and re.match(r"[!-/:-@\[-`{-~]", label):
            label = "\\" + label
        return label

    def tag(rows: list[int]) -> str:
        # An HTML comment is invisible in rendered Markdown and names the source rows.
        return f" <!-- rows: {' '.join(map(str, rows))} -->" if row_ids and rows else ""

    def walk(node: Node, depth: int) -> None:
        bullet_depth = max(depth - heading_levels, 0)
        for child in node.items:
            if isinstance(child, Detail):
                line = details_line(child.values)
                if line:
                    out.append(" " * (indent * bullet_depth) + "- " + line + tag([child.row]))
                continue
            label = item_label(child.key)
            if label_levels:
                label = f"{quote(headings[depth])}: {label}"
            label += tag(child.rows)
            if depth < heading_levels:
                block("#" * (depth + 1 + offset) + " " + label)
            else:
                out.append(" " * (indent * (depth - heading_levels)) + "- " + label)
            walk(child, depth + 1)

    if title:
        block("# " + quote(title))
    walk(root, 0)
    while out and out[-1] == "":
        out.pop()
    return "\n".join(out) + "\n"


def front_matter(levels: list[str], details: list[str], args, sheet: str, source: str) -> str:
    """YAML front matter recording what xlsx-unfold needs to rebuild the table."""
    settings = {
        "version": 1,
        "sheet": sheet,
        "levels": levels,
        "details": details,
        "title": bool(args.title),
        "label_levels": args.label_levels,
        "blank_label": args.blank_label,
        "raw": args.raw,
        "fill_down": args.fill_down,
    }
    if args.row_ids:
        settings["fingerprint"] = source
    # JSON scalars and flow lists are valid YAML, which keeps the block dependency-free.
    lines = [f"  {key}: {json.dumps(value, ensure_ascii=False)}" for key, value in settings.items()]
    return "\n".join(["---", f"{FRONT_MATTER_KEY}:", *lines, "---", ""]) + "\n"


def write_output(output: Path, content: str, protected: Path, overwrite: bool) -> None:
    if output.resolve() == protected.resolve():
        raise ValidationError("Output must not be the input workbook")
    if output.exists() and not overwrite:
        raise ValidationError(f"Output exists; pass --overwrite to replace it: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=".outline-", suffix=".md", dir=output.parent)
    temp = Path(name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
        os.replace(temp, output)
    finally:
        temp.unlink(missing_ok=True)


def convert(args) -> tuple[str, int, int]:
    sheet, headings, data = read_table(args.input, args.sheet)
    selected = select_columns(headings, args.columns, args.column_indices)
    levels = len(selected) if args.levels is None else args.levels
    if not 1 <= levels <= len(selected):
        raise ValidationError(f"--levels must be between 1 and {len(selected)}")
    heading_levels = args.heading_levels
    if heading_levels < 0 or heading_levels > levels:
        raise ValidationError(f"--heading-levels must be between 0 and {levels}")
    if heading_levels + (1 if args.title else 0) > MAX_HEADING_LEVEL:
        raise ValidationError("Markdown supports six heading levels, including the title")
    if args.indent < 1:
        raise ValidationError("--indent must be a positive integer")
    rows = [[row[column] for column in selected] for row in data]
    root = build_tree(rows, levels, args.fill_down, args.group)
    names = [headings[column] for column in selected]
    content = render(
        root,
        names[:levels],
        names[levels:],
        heading_levels=heading_levels,
        title=args.title,
        label_levels=args.label_levels,
        indent=args.indent,
        raw=args.raw,
        blank_label=args.blank_label,
        row_ids=args.row_ids,
    )
    if args.front_matter:
        source = fingerprint(sheet, names, rows)
        content = front_matter(names[:levels], names[levels:], args, sheet, source) + content

    def count(node: Node) -> int:
        return sum(1 + count(child) for child in node.children)

    used = sum(1 for row in rows if any(value is not None for value in row))
    return content, used, count(root)


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    root.add_argument("--input", type=Path, required=True)
    root.add_argument("--sheet", help="Sheet name (default: the first sheet)")
    root.add_argument("--output", type=Path, help="Markdown file (default: standard output)")
    root.add_argument("--overwrite", action="store_true", help="Replace an existing --output")
    selection = root.add_mutually_exclusive_group()
    selection.add_argument("--columns", nargs="+", help="Headings in outline order")
    selection.add_argument("--column-indices", nargs="+", type=int, help="1-based positions")
    root.add_argument(
        "--levels",
        type=int,
        help="Selected columns that form the hierarchy (default: all); the rest are details",
    )
    root.add_argument(
        "--heading-levels",
        type=int,
        default=0,
        help="Leading hierarchy levels rendered as # headings; deeper levels are bullets",
    )
    root.add_argument("--title", help="Top-level # heading; shifts other headings down")
    root.add_argument(
        "--label-levels", action="store_true", help="Prefix hierarchy items with the heading"
    )
    root.add_argument(
        "--fill-down",
        action="store_true",
        help="Treat a blank hierarchy cell as the value above it within the same parent",
    )
    root.add_argument(
        "--group",
        action="store_true",
        help="Merge non-adjacent rows with the same parent values (first-occurrence order)",
    )
    root.add_argument("--blank-label", default="(blank)", help="Label for an inner blank level")
    root.add_argument("--indent", type=int, default=2, help="Spaces per nested bullet level")
    root.add_argument("--raw", action="store_true", help="Do not escape Markdown characters")
    root.add_argument(
        "--row-ids",
        action="store_true",
        help="Tag items with their sheet rows, so xlsx-unfold --into can merge edits back",
    )
    root.add_argument(
        "--front-matter",
        action="store_true",
        help="Start with YAML front matter that lets xlsx-unfold rebuild the table",
    )
    return root


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        content, rows, items = convert(args)
        summary = f"Outlined {rows} rows into {items} items"
        if args.output is None:
            sys.stdout.write(content)
            print(summary, file=sys.stderr)
        else:
            write_output(args.output, content, args.input, args.overwrite)
            print(f"{summary}: {args.output}")
        return 0
    except ValidationError as exc:
        print(f"Validation error: {exc}", file=sys.stderr)
        return 2
    except (OSError, BadZipFile) as exc:
        print(f"IO error: {exc}", file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
