"""Behavior contracts for ado-triage using synthetic snapshots; no ADO access."""

import json
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from workflow_automations import ado_triage as tool

AS_OF = datetime(2026, 10, 1)
BASE = "https://ado.example.com/tfs/Collection/Project"
TEAM_X = "Proj\\TeamX"
TEAM_Y = "Proj\\TeamY"


def rel(kind, target):
    return {"rel": kind, "url": f"{BASE}/_apis/wit/workItems/{target}"}


def item(wi_id, area, title, *, oe=0, cw=0, rw=0, owner="Owner", changed_days=2, **extra):
    fields = {
        tool.FIELD_TITLE: title,
        tool.FIELD_AREA_PATH: area,
        tool.FIELD_STATE: "Active",
        tool.FIELD_ORIGINAL_ESTIMATE: oe,
        tool.FIELD_COMPLETED_WORK: cw,
        tool.FIELD_REMAINING_WORK: rw,
        tool.FIELD_CHANGED_DATE: (AS_OF - timedelta(days=changed_days)).isoformat() + "Z",
    }
    if owner:
        fields[tool.FIELD_ASSIGNED_TO] = {"displayName": owner}
    fields.update(extra)
    return {"id": wi_id, "fields": fields, "relations": []}


def iso(days):
    return (AS_OF + timedelta(days=days)).isoformat() + "Z"


@pytest.fixture
def snapshot():
    items = [
        item(1, TEAM_X, "Big risky CR", oe=50, cw=10, rw=120, **{tool.FIELD_UAT_EXPECTED: iso(20)}),
        item(2, TEAM_X, "Waiting child", rw=30, **{tool.FIELD_PARENT: 1}),
        item(3, TEAM_X, "Almost done", oe=20, cw=15, rw=5),
        item(4, TEAM_Y, "Dormant and unowned", rw=10, owner=None, changed_days=40),
        item(5, TEAM_Y, "Estimate overdue", rw=8, **{tool.FIELD_EST_EXPECTED: iso(-30)}),
        item(6, TEAM_Y, "Far future", rw=15, **{tool.FIELD_UAT_EXPECTED: iso(120)}),
    ]
    # 1 blocks 2 (Successor); 1 depends on 3 (Predecessor); 1 parents 2 (Child).
    items[0]["relations"] = [rel(tool.REL_SUCCESSOR, 2), rel(tool.REL_PREDECESSOR, 3)]
    items[0]["relations"].append(rel(tool.REL_CHILD, 2))
    return tool.Snapshot(items=items, parent_titles={1: "Big risky CR"})


CAPACITY = {TEAM_X: 100.0, TEAM_Y: 500.0}


def settings(**overrides):
    values = {"as_of": AS_OF, "base_url": BASE, "capacity_months": 2}
    values.update(overrides)
    return tool.Settings(**values)


def by_id(report):
    return report.master.set_index(tool.FIELD_ID)


def test_risk_and_opportunity_flags_are_separated(snapshot):
    rows = by_id(tool.build_report(snapshot, CAPACITY, settings()))
    assert set(rows.loc[1, tool.COL_FLAGS].split("|")) == {
        "DATE_DRIFT",
        "BLOCKS_OTHERS",
        "EXCEEDS_MONTH_CAP",
    }
    assert rows.loc[1, tool.COL_SCORE] == 9
    assert set(rows.loc[4, tool.COL_FLAGS].split("|")) == {"STALE", "UNASSIGNED"}
    assert rows.loc[5, tool.COL_FLAGS] == "EST_DEBT_DATE"
    # Opportunities never contribute to attention.
    assert rows.loc[3, tool.COL_OPP_FLAGS] == "QUICK_WIN"
    assert rows.loc[6, tool.COL_OPP_FLAGS] == "DEFER_CANDIDATE"
    assert rows.loc[3, tool.COL_SCORE] == rows.loc[6, tool.COL_SCORE] == 0


def test_blast_radius_counts_only_items_this_item_obstructs(snapshot):
    rows = by_id(tool.build_report(snapshot, CAPACITY, settings()))
    # The Predecessor link (1 depends on 3) must not count as 1 blocking anything.
    assert rows.loc[1, tool.COL_BLOCKS_COUNT] == 1
    assert rows.loc[3, tool.COL_BLOCKS_COUNT] == 0


def chain_snapshot(*links):
    """Items 1-6 in one team; each link is (source, relation, target)."""
    items = {i: item(i, TEAM_X, f"Item {i}", rw=10) for i in range(1, 7)}
    for src, kind, dst in links:
        items[src]["relations"].append(rel(kind, dst))
    return tool.Snapshot(items=list(items.values()))


def test_transitive_chain_reports_length_and_path():
    # A=1 blocks B=2 blocks C=3.
    snap = chain_snapshot((1, tool.REL_BLOCKS_FORWARD, 2), (2, tool.REL_BLOCKS_FORWARD, 3))
    rows = by_id(tool.build_report(snap, {}, settings()))
    assert rows.loc[1, tool.COL_CHAIN_LENGTH] == 2
    assert rows.loc[1, tool.COL_CHAIN] == "1 > 2 > 3"
    assert rows.loc[1, tool.COL_BLOCKS_COUNT] == 1  # one-hop count is unchanged
    assert rows.loc[2, tool.COL_CHAIN_LENGTH] == 1
    assert rows.loc[3, tool.COL_CHAIN_LENGTH] == 0
    assert rows.loc[3, tool.COL_CHAIN] == ""


def test_chain_takes_longest_branch_and_reads_reverse_links():
    # 1 blocks 2 and 4; 4 -> 5 -> 6 is recorded only as reverse links on the blocked items.
    snap = chain_snapshot(
        (1, tool.REL_SUCCESSOR, 2),
        (1, tool.REL_BLOCKS_FORWARD, 4),
        (5, tool.REL_BLOCKS_REVERSE, 4),
        (6, tool.REL_PREDECESSOR, 5),
    )
    rows = by_id(tool.build_report(snap, {}, settings()))
    assert rows.loc[1, tool.COL_CHAIN] == "1 > 4 > 5 > 6"
    assert rows.loc[1, tool.COL_CHAIN_LENGTH] == 3


def test_chain_crosses_items_outside_the_query():
    # 2 blocks 99 (not queried) and 3 is blocked by 99: the chain still reaches 3.
    snap = chain_snapshot((2, tool.REL_BLOCKS_FORWARD, 99), (3, tool.REL_BLOCKS_REVERSE, 99))
    rows = by_id(tool.build_report(snap, {}, settings()))
    assert rows.loc[2, tool.COL_CHAIN] == "2 > 99 > 3"


def test_mirrored_links_are_not_counted_twice():
    snap = chain_snapshot((1, tool.REL_BLOCKS_FORWARD, 2), (2, tool.REL_BLOCKS_REVERSE, 1))
    rows = by_id(tool.build_report(snap, {}, settings()))
    assert rows.loc[1, tool.COL_CHAIN_LENGTH] == 1


def test_dependency_cycle_terminates_and_counts_members_once():
    # 1 -> 2 -> 3 -> 1 cycle, with 3 also blocking 4.
    snap = chain_snapshot(
        (1, tool.REL_BLOCKS_FORWARD, 2),
        (2, tool.REL_BLOCKS_FORWARD, 3),
        (3, tool.REL_BLOCKS_FORWARD, 1),
        (3, tool.REL_BLOCKS_FORWARD, 4),
    )
    rows = by_id(tool.build_report(snap, {}, settings()))
    assert rows.loc[1, tool.COL_CHAIN] == "1 > 2 > 3 > 4"
    assert rows.loc[2, tool.COL_CHAIN] == "2 > 1 > 3 > 4"
    assert rows.loc[4, tool.COL_CHAIN_LENGTH] == 0


def test_top_blockers_sorts_by_chain_length():
    # 3 blocks two items directly; 1 heads a longer chain with one direct link.
    snap = chain_snapshot(
        (3, tool.REL_BLOCKS_FORWARD, 5),
        (3, tool.REL_BLOCKS_FORWARD, 6),
        (1, tool.REL_BLOCKS_FORWARD, 2),
        (2, tool.REL_BLOCKS_FORWARD, 4),
        (4, tool.REL_BLOCKS_FORWARD, 6),
    )
    report = tool.build_report(snap, {}, settings())
    view = tool.view_blockers(report.master, report.relations, {}, settings())
    assert list(view[tool.FIELD_ID])[:2] == [1, 2]
    assert {tool.COL_CHAIN_LENGTH, tool.COL_CHAIN} <= set(view.columns)


def test_root_follows_parent_field_and_child_links_are_kept(snapshot):
    report = tool.build_report(snapshot, CAPACITY, settings())
    rows = by_id(report)
    assert rows.loc[1, tool.COL_IS_ROOT] and not rows.loc[2, tool.COL_IS_ROOT]
    assert rows.loc[2, tool.COL_PARENT_TITLE] == "Big risky CR"
    assert "Child" in set(report.relations[tool.REL_COL_TYPE])


def test_unknown_linked_items_are_labelled_not_dropped():
    snap = tool.Snapshot(items=[item(1, TEAM_X, "A")])
    snap.items[0]["relations"] = [
        rel(tool.REL_BLOCKS_FORWARD, 99),
        rel(tool.REL_BLOCKS_FORWARD, 98),
    ]
    snap.external = {98: {tool.FIELD_TITLE: "Outside", tool.FIELD_AREA_PATH: TEAM_Y}}
    relations = tool.build_report(snap, {}, settings()).relations.set_index(tool.REL_COL_DEST_ID)
    assert relations.loc[99, tool.REL_COL_DEST_TITLE] == "UNKNOWN"
    assert relations.loc[98, tool.REL_COL_DEST_AREA] == TEAM_Y


def test_hierarchical_capacity_load_and_rag(snapshot):
    capacity = {"Proj": 2000.0, TEAM_X: 100.0}
    report = tool.build_report(snapshot, capacity, settings())
    summary = report.sheets[tool.SHEET_SUMMARY].set_index("Team")
    assert summary.loc[TEAM_X, tool.COL_RAG] == "RED"
    assert summary.loc[TEAM_X, tool.COL_LOAD_PCT] == 155.0
    # TeamY has no row of its own and resolves to the parent level.
    assert summary.loc["Proj", tool.COL_RAG] == "GREEN"
    assert list(summary.index)[0] == TEAM_X  # RED sorts first


def test_undefined_capacity_is_grey_not_green(snapshot):
    summary = tool.build_report(snapshot, {}, settings()).sheets[tool.SHEET_SUMMARY]
    assert set(summary[tool.COL_RAG]) == {"GREY"}
    assert set(summary[tool.COL_CAPACITY]) == {tool.UNDEFINED_CAPACITY}


def test_top_20_percent_keeps_at_least_one_per_area():
    items = [item(i, TEAM_X, f"x{i}", rw=i) for i in range(1, 11)] + [item(20, TEAM_Y, "y", rw=3)]
    sheet = tool.build_report(tool.Snapshot(items=items), {}, settings()).sheets[
        "15. Top 20% Remaining Work"
    ]
    assert sorted(sheet[tool.FIELD_ID]) == [9, 10, 20]


def test_deltas_persistence_new_and_resolved(snapshot, tmp_path):
    first = tool.build_report(snapshot, CAPACITY, settings())
    path = tmp_path / "week1.xlsx"
    tool.write_excel(path, first, {"As of": "2026-10-01"})

    # Next week: item 4 gets an owner and activity; a new oversized item appears.
    snapshot.items[3] = item(4, TEAM_Y, "Dormant and unowned", rw=10)
    snapshot.items.append(item(7, TEAM_Y, "Brand new", rw=600))
    previous = tool.read_previous_master(path)
    second = tool.build_report(snapshot, CAPACITY, settings(), previous)
    rows = by_id(second)

    assert rows.loc[1, tool.COL_DELTA] == "CARRIED" and rows.loc[1, tool.COL_PERSIST] == 2
    assert rows.loc[7, tool.COL_DELTA] == "NEW" and rows.loc[7, tool.COL_PERSIST] == 1
    assert rows.loc[4, tool.COL_PERSIST] == 0
    resolved = second.sheets[tool.SHEET_RESOLVED].set_index(tool.FIELD_ID)
    assert list(resolved.index) == [4]
    assert resolved.loc[4, "Resolution"] == "Flags cleared"


def test_legacy_previous_report_without_flags_degrades(snapshot, tmp_path):
    path = tmp_path / "legacy.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({tool.FIELD_ID: [1, 2]}).to_excel(
            writer, sheet_name="1. Master Data", index=False
        )
    report = tool.build_report(snapshot, CAPACITY, settings(), tool.read_previous_master(path))
    assert by_id(report).loc[1, tool.COL_DELTA] == "CARRIED"
    assert by_id(report).loc[5, tool.COL_DELTA] == "NEW"
    assert any("predates flags" in note for note in report.notes)


def test_workbook_layout_omissions_and_profiles(snapshot, tmp_path):
    report = tool.build_report(snapshot, CAPACITY, settings())
    path = tmp_path / "full.xlsx"
    tool.write_excel(path, report, {"As of": "2026-10-01"})
    wb = load_workbook(path)
    assert wb.sheetnames[0] == tool.SHEET_CONFIG
    assert all(len(name) <= 31 for name in wb.sheetnames)
    assert tool.SHEET_RESOLVED not in wb.sheetnames  # empty without a previous report
    assert tool.SHEET_RESOLVED in report.omitted
    agenda = wb[tool.SHEET_AGENDA]
    assert agenda["A1"].value.startswith("Top items") and agenda.freeze_panes == "A4"
    first_agenda = pd.read_excel(path, sheet_name=tool.SHEET_AGENDA, header=2)
    assert first_agenda.iloc[0][tool.FIELD_ID] == 1

    exec_report = tool.build_report(snapshot, CAPACITY, settings(profile="exec"))
    assert tool.SHEET_SUMMARY in exec_report.sheets
    assert "5. Estimation Debt (Dates)" not in exec_report.sheets
    # Master Data stays so an exec report can serve as the next --previous-file.
    exec_path = tmp_path / "exec.xlsx"
    tool.write_excel(exec_path, exec_report, {"As of": "2026-10-01"})
    previous = tool.read_previous_master(exec_path)
    nxt = tool.build_report(snapshot, CAPACITY, settings(), previous)
    assert set(nxt.master[tool.COL_DELTA]) == {"CARRIED"}


def test_non_xlsx_previous_report_is_a_validation_error(tmp_path):
    for name, data in (("prev.csv", b"System.Id\n1\n"), ("prev.xlsx", b"not a zip")):
        path = tmp_path / name
        path.write_bytes(data)
        with pytest.raises(tool.TriageError, match="not a readable XLSX"):
            tool.read_previous_master(path)


def test_cross_team_lists_each_dependency_once():
    a = item(1, TEAM_X, "A")
    b = item(2, TEAM_Y, "B")
    c = item(3, TEAM_Y, "C")
    a["relations"] = [rel(tool.REL_SUCCESSOR, 2), rel(tool.REL_BLOCKS_FORWARD, 3)]
    b["relations"] = [rel(tool.REL_PREDECESSOR, 1)]
    c["relations"] = [rel(tool.REL_BLOCKS_REVERSE, 1), rel(tool.REL_PREDECESSOR, 99)]
    report = tool.build_report(tool.Snapshot(items=[a, b, c]), {}, settings())
    cross = report.sheets["11. Cross-Team Deps"]
    pairs = set(zip(cross[tool.REL_COL_SOURCE_ID], cross[tool.REL_COL_DEST_ID], strict=True))
    # Mirrored reverse links are dropped; 3 -> 99 has no forward row, so it stays.
    assert pairs == {(1, 2), (1, 3), (3, 99)}
    assert len(report.relations) == 5  # All Relations remains the raw graph


def capacity_workbook(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = tool.CAPACITY_SHEET_NAME
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_capacity_accepts_excel_date_headers_and_sums_window(tmp_path):
    path = capacity_workbook(
        tmp_path / "cap.xlsx",
        ["AreaPath", datetime(2026, 11, 1), "2026-12", "2027-01"],
        [[TEAM_X, 40, 60, 999], [" Proj ", 10, None, 0]],
    )
    months = tool.target_months(None, 2, AS_OF)
    assert months == ("2026-11", "2026-12")
    assert tool.load_capacity(path, months) == {TEAM_X: 100.0, "Proj": 10.0}


@pytest.mark.parametrize(
    ("headers", "rows", "message"),
    [
        (["AreaPath", "2026-11"], [[TEAM_X, 1]], "not found"),
        (["AreaPath", "2026-11", "2026-12"], [[TEAM_X, 1, 1], [TEAM_X, 2, 2]], "duplicate"),
    ],
)
def test_capacity_rejects_partial_window_and_duplicates(tmp_path, headers, rows, message):
    path = capacity_workbook(tmp_path / "cap.xlsx", headers, rows)
    with pytest.raises(tool.TriageError, match=message):
        tool.load_capacity(path, ("2026-11", "2026-12"))


def test_target_months_rolls_over_year_and_validates():
    assert tool.target_months("2026-12", 2, AS_OF) == ("2026-12", "2027-01")
    with pytest.raises(tool.TriageError):
        tool.target_months("12/2026", 2, AS_OF)


def cli(tmp_path, *extra):
    return [
        "export",
        "--url",
        BASE,
        "--query-id",
        "00000000-0000-0000-0000-000000000000",
        "--as-of",
        "2026-10-01",
        *extra,
    ]


def test_cli_end_to_end_with_network_layer_replaced(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    out = tmp_path / "out" / "report.xlsx"
    assert tool.main(cli(tmp_path, "--output", str(out))) == tool.EXIT_OK
    assert tool.SHEET_SUMMARY in load_workbook(out).sheetnames
    assert not [p for p in out.parent.iterdir() if p.name.startswith(".")]  # temp cleaned

    # Existing output is preserved without --overwrite.
    before = out.read_bytes()
    assert tool.main(cli(tmp_path, "--output", str(out))) == tool.EXIT_VALIDATION
    assert out.read_bytes() == before
    second = cli(tmp_path, "--output", str(out), "--overwrite", "--previous-file", str(out))
    assert tool.main(second) == tool.EXIT_VALIDATION  # output may not replace its own input

    nxt = tmp_path / "next.xlsx"
    assert tool.main(cli(tmp_path, "--output", str(nxt), "--previous-file", str(out))) == 0


def test_cli_validates_local_inputs_before_network(tmp_path, monkeypatch):
    def unreachable(url, query):
        raise AssertionError("network must not be touched")

    monkeypatch.setattr(tool, "fetch_snapshot", unreachable)
    cap = capacity_workbook(tmp_path / "cap.xlsx", ["AreaPath", "2026-11"], [[TEAM_X, 1]])
    args = cli(tmp_path, "--output", str(tmp_path / "r.xlsx"), "--capacity-file", str(cap))
    assert tool.main(args) == tool.EXIT_VALIDATION


@pytest.mark.parametrize(
    ("error", "code"),
    [(tool.AdoError("down"), tool.EXIT_ADO), (tool.EmptyQueryError("none"), tool.EXIT_EMPTY)],
)
def test_cli_exit_codes_for_ado_failures(tmp_path, monkeypatch, error, code):
    def failing(url, query):
        raise error

    monkeypatch.setattr(tool, "fetch_snapshot", failing)
    out = tmp_path / "r.xlsx"
    assert tool.main(cli(tmp_path, "--output", str(out))) == code
    assert not out.exists()


def test_csv_output_is_master_data(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    out = tmp_path / "r.csv"
    assert tool.main(cli(tmp_path, "--output", str(out), "--format", "csv")) == 0
    frame = pd.read_csv(out)
    assert {tool.COL_FLAGS, tool.COL_SCORE} <= set(frame.columns)
    assert len(frame) == len(snapshot.items)


# --- saved snapshots (--save-snapshot, --from-snapshot) ---


def offline_args(snap_path, out, *extra):
    return ["export", "--from-snapshot", str(snap_path), "--as-of", "2026-10-01",
            "--output", str(out), *extra]  # fmt: skip


def data_sheets(path):
    """Every sheet's values, except the run-specific Config sheet."""
    sheets = pd.read_excel(path, sheet_name=None)
    return {name: frame for name, frame in sheets.items() if not name.startswith("0.")}


@pytest.fixture
def no_network(monkeypatch):
    def unreachable(url, query):
        raise AssertionError("network must not be touched")

    monkeypatch.setattr(tool, "fetch_snapshot", unreachable)


def test_saved_snapshot_rebuilds_the_same_report_offline(snapshot, tmp_path, monkeypatch):
    snapshot.items[0]["relations"].append(rel(tool.REL_BLOCKS_FORWARD, 77))  # outside the query
    snapshot.external = {77: {tool.FIELD_TITLE: "Elsewhere", tool.FIELD_AREA_PATH: "Proj\\Z"}}
    snapshot.warnings = ["1 parent titles unavailable (deleted or no permission)"]
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    live, snap = tmp_path / "live.xlsx", tmp_path / "snap.json"
    assert tool.main(cli(tmp_path, "--output", str(live), "--save-snapshot", str(snap))) == 0

    def unreachable(url, query):
        raise AssertionError("network must not be touched")

    monkeypatch.setattr(tool, "fetch_snapshot", unreachable)
    replay = tmp_path / "replay.xlsx"
    assert tool.main(offline_args(snap, replay)) == tool.EXIT_OK

    expected, actual = data_sheets(live), data_sheets(replay)
    assert expected.keys() == actual.keys()
    for name, frame in expected.items():
        pd.testing.assert_frame_equal(actual[name], frame, obj=name)
    rows = load_workbook(replay)["0. Config & Legend"].iter_rows(values_only=True)
    config = {row[0]: row[1] for row in rows if row}
    assert config["ADO URL"] == BASE
    assert config["Query ID"] == "00000000-0000-0000-0000-000000000000"
    assert config["Snapshot"] == str(snap)


def test_save_snapshot_refuses_existing_file_and_input_paths(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    snap = tmp_path / "snap.json"
    snap.write_text("keep")
    out = tmp_path / "r.xlsx"
    assert tool.main(cli(tmp_path, "--output", str(out), "--save-snapshot", str(snap))) == 2
    assert snap.read_text() == "keep" and not out.exists()
    same = cli(tmp_path, "--output", str(out), "--save-snapshot", str(out), "--overwrite")
    assert tool.main(same) == tool.EXIT_VALIDATION
    replaced = cli(tmp_path, "--output", str(out), "--save-snapshot", str(snap), "--overwrite")
    assert tool.main(replaced) == tool.EXIT_OK
    assert tool.load_snapshot(snap).producer == tool.SNAPSHOT_PRODUCER


@pytest.mark.parametrize(
    ("content", "extra", "code"),
    [
        ("not json", [], tool.EXIT_VALIDATION),
        ('{"schema": "other"}', [], tool.EXIT_VALIDATION),
        (None, ["--query-id", "q"], tool.EXIT_VALIDATION),  # query comes from the snapshot
        (None, ["--save-snapshot", "again.json"], tool.EXIT_VALIDATION),  # argparse refuses
    ],
)
def test_from_snapshot_validates_before_any_write(
    snapshot, tmp_path, no_network, content, extra, code
):
    snap = tmp_path / "snap.json"
    if content is None:
        tool.save_snapshot(tool.snapshot_to_raw(snapshot, BASE, "q"), snap)
    else:
        snap.write_text(content)
    out = tmp_path / "r.xlsx"
    try:
        result = tool.main(offline_args(snap, out, *extra))
    except SystemExit as exc:  # argparse usage error
        result = exc.code
    assert result == code
    assert not out.exists()


def test_from_snapshot_output_may_not_replace_it(snapshot, tmp_path, no_network):
    snap = tmp_path / "snap.json"
    tool.save_snapshot(tool.snapshot_to_raw(snapshot, BASE, "q"), snap)
    assert tool.main(offline_args(snap, snap, "--overwrite")) == tool.EXIT_VALIDATION
    assert tool.load_snapshot(snap).roots == [1, 2, 3, 4, 5, 6]


def test_from_snapshot_empty_scope_and_missing_url(tmp_path, no_network):
    snap = tmp_path / "snap.json"
    tool.save_snapshot(tool.snapshot_to_raw(tool.Snapshot(items=[]), BASE, "q"), snap)
    assert tool.main(offline_args(snap, tmp_path / "r.xlsx")) == tool.EXIT_EMPTY
    nourl = tmp_path / "nourl.json"
    tool.save_snapshot(
        tool.snapshot_to_raw(tool.Snapshot(items=[item(1, TEAM_X, "A")]), "", "q"), nourl
    )
    assert tool.main(offline_args(nourl, tmp_path / "r.xlsx")) == tool.EXIT_VALIDATION
    with_url = offline_args(nourl, tmp_path / "r.xlsx", "--url", BASE)
    assert tool.main(with_url) == tool.EXIT_OK


def test_live_export_still_requires_url_and_query(tmp_path, no_network):
    args = ["export", "--url", BASE, "--output", str(tmp_path / "r.xlsx")]
    assert tool.main(args) == tool.EXIT_VALIDATION


# --- per-team reports (--team-dir) ---


def team_ids(team_report):
    """Work item ids per sheet, by the column that names the row's own item."""
    ids = {}
    for name, frame in team_report.sheets.items():
        for key in (tool.REL_COL_SOURCE_ID, tool.FIELD_ID):
            if key in frame.columns:
                ids[name] = set(frame[key])
                break
    return ids


def cross_pairs(report):
    cross = report.sheets["11. Cross-Team Deps"]
    return list(zip(cross[tool.REL_COL_SOURCE_ID], cross[tool.REL_COL_DEST_ID], strict=True))


def test_team_reports_partition_every_sheet_by_team(snapshot):
    snapshot.items[0]["relations"].append(rel(tool.REL_SUCCESSOR, 4))  # X item blocks a Y item
    combined = tool.build_report(snapshot, CAPACITY, settings())
    teams, notes = tool.build_team_reports(combined, CAPACITY, settings())
    assert list(teams) == [TEAM_X, TEAM_Y] and notes == []
    own = {TEAM_X: {1, 2, 3}, TEAM_Y: {4, 5, 6}}
    for team, report in teams.items():
        assert report.team == team
        sheet_ids = team_ids(report)
        assert set(report.sheets) - set(sheet_ids) == {tool.SHEET_SUMMARY, "4. Capacity Triage"}
        for name, ids in sheet_ids.items():
            assert ids <= own[team], (team, name, ids)
        assert set(report.sheets[tool.SHEET_SUMMARY]["Team"]) == {team}
        assert set(report.sheets["4. Capacity Triage"][tool.COL_LEVEL]) == {team}
    assert len(teams[TEAM_X].master) + len(teams[TEAM_Y].master) == len(combined.master)
    # The cross-team dependency is listed for the team whose item stores the link.
    assert cross_pairs(teams[TEAM_X]) == [(1, 4)]
    assert "11. Cross-Team Deps" in teams[TEAM_Y].omitted
    # Flags and scores come from the combined run; the agenda is ranked within the team.
    combined_rows = by_id(combined)
    for report in teams.values():
        for wi_id, row in report.master.set_index(tool.FIELD_ID).iterrows():
            assert row[tool.COL_SCORE] == combined_rows.loc[wi_id, tool.COL_SCORE]
    assert set(teams[TEAM_Y].sheets[tool.SHEET_AGENDA][tool.FIELD_ID]) == {4, 5}


def test_mirrored_dependency_is_listed_by_each_team():
    a = item(1, TEAM_X, "A")
    b = item(2, TEAM_Y, "B")
    a["relations"] = [rel(tool.REL_SUCCESSOR, 2)]
    b["relations"] = [rel(tool.REL_PREDECESSOR, 1)]
    combined = tool.build_report(tool.Snapshot(items=[a, b]), {}, settings())
    teams, _ = tool.build_team_reports(combined, {}, settings())
    for team, pair in ((TEAM_X, (1, 2)), (TEAM_Y, (2, 1))):
        assert cross_pairs(teams[team]) == [pair]
    combined_cross = combined.sheets["11. Cross-Team Deps"]
    assert len(combined_cross) == 1  # the combined report still lists it once


def test_resolved_items_go_to_current_or_previous_team(snapshot, tmp_path):
    first = tool.build_report(snapshot, CAPACITY, settings())
    path = tmp_path / "week1.xlsx"
    tool.write_excel(path, first, {"As of": "2026-10-01"})
    previous = tool.read_previous_master(path)
    snapshot.items[3] = item(4, TEAM_Y, "Dormant and unowned", rw=10)  # flags cleared
    del snapshot.items[4]  # item 5 (TEAM_Y, flagged) dropped from the query
    second = tool.build_report(snapshot, CAPACITY, settings(), previous)
    teams, notes = tool.build_team_reports(second, CAPACITY, settings(), previous)
    resolved = teams[TEAM_Y].sheets[tool.SHEET_RESOLVED].set_index(tool.FIELD_ID)
    assert resolved["Resolution"].to_dict() == {4: "Flags cleared", 5: "Dropped from query"}
    assert tool.SHEET_RESOLVED in teams[TEAM_X].omitted and notes == []
    assert teams[TEAM_X].notes == second.notes  # team files carry only the run's warnings

    # A dropped item whose old team has no rows now is reported, not misfiled.
    snapshot.items = [i for i in snapshot.items if i["id"] in (1, 2, 3)]
    third = tool.build_report(snapshot, CAPACITY, settings(), previous)
    teams, notes = tool.build_team_reports(third, CAPACITY, settings(), previous)
    assert list(teams) == [TEAM_X] and notes[0].endswith("[4, 5]")
    assert not any("no current team" in note for note in teams[TEAM_X].notes)
    assert set(third.sheets[tool.SHEET_RESOLVED][tool.FIELD_ID]) == {4, 5}


def test_items_without_area_path_form_their_own_team():
    items = [item(1, TEAM_X, "A"), item(2, None, "No area")]
    combined = tool.build_report(tool.Snapshot(items=items), {}, settings())
    teams, _ = tool.build_team_reports(combined, {}, settings())
    assert {t: set(r.master[tool.FIELD_ID]) for t, r in teams.items()} == {"": {2}, TEAM_X: {1}}


def test_items_without_area_path_appear_in_every_team_view(tmp_path):
    missing = item(2, None, "Missing area", rw=7)
    del missing["fields"][tool.FIELD_AREA_PATH]  # ADO omits empty fields
    items = [item(1, TEAM_X, "A", rw=5), missing, item(3, None, "Null area", rw=3)]
    report = tool.build_report(tool.Snapshot(items=items), CAPACITY, settings())

    top = report.sheets["15. Top 20% Remaining Work"]
    assert set(top[tool.FIELD_ID]) == {1, 2}  # one per team; used to raise IntCastingNaNError
    summary = report.sheets[tool.SHEET_SUMMARY].set_index("Team")
    assert summary.loc["", "Total_CRs"] == 2 and summary.loc["", "Total_Remaining_Work"] == 10
    triage = report.sheets["4. Capacity Triage"].set_index(tool.COL_LEVEL).loc[""]
    assert triage["Total_CRs"] == 2 and triage[tool.COL_CAPACITY] == tool.UNDEFINED_CAPACITY
    tool.write_excel(tmp_path / "report.xlsx", report, {})


def test_dropped_item_without_area_path_stays_with_the_blank_team(tmp_path):
    items = [item(1, None, "No area", owner=None), item(2, None, "Also no area")]
    first = tool.build_report(tool.Snapshot(items=items), {}, settings())
    path = tmp_path / "week1.xlsx"
    tool.write_excel(path, first, {})
    previous = tool.read_previous_master(path)  # blank team reads back as NaN
    second = tool.build_report(tool.Snapshot(items=items[1:]), {}, settings(), previous)
    teams, notes = tool.build_team_reports(second, {}, settings(), previous)
    assert notes == [] and list(teams[""].sheets[tool.SHEET_RESOLVED][tool.FIELD_ID]) == [1]


@pytest.mark.parametrize(
    ("team", "name"),
    [
        ("Proj\\Dept\\Team A", "Proj_Dept_Team A"),
        ('a/b:c*d?"e<f>g|h', "a_b_c_d_e_f_g_h"),
        ("", "NO_AREA"),
        (" .Team. ", "Team"),
        ("Proj\\", "Proj"),
    ],
)
def test_team_filename_is_safe_on_windows(team, name):
    assert tool.team_filename(team) == name


def test_team_paths_reject_names_that_collide():
    paths = tool.team_paths(Path("r/w1.xlsx"), Path("teams"), ["P\\A", "P\\B"])
    assert paths == {"P\\A": Path("teams/w1.P_A.xlsx"), "P\\B": Path("teams/w1.P_B.xlsx")}
    with pytest.raises(tool.TriageError, match="same file"):
        tool.team_paths(Path("w1.xlsx"), Path("t"), ["P\\A", "P/A"])
    with pytest.raises(tool.TriageError, match="same file"):
        tool.team_paths(Path("w1.xlsx"), Path("t"), ["P\\a", "P\\A"])
    # Unicode: composed and decomposed forms are one name on macOS; ß and ss are two.
    with pytest.raises(tool.TriageError, match="same file"):
        tool.team_paths(Path("w1.xlsx"), Path("t"), ["P\\Caf\u00e9", "P\\Cafe\u0301"])
    assert len(tool.team_paths(Path("w1.xlsx"), Path("t"), ["P\\Stra\u00dfe", "P\\Strasse"])) == 2
    with pytest.raises(tool.TriageError, match="over 255 bytes"):
        tool.team_paths(Path("w1.xlsx"), Path("t"), ["P\\" + "x" * 250])


def workbook_sheets(path):
    sheets = pd.read_excel(path, sheet_name=None, header=None)
    config = sheets[tool.SHEET_CONFIG]
    sheets[tool.SHEET_CONFIG] = config[config[0] != "Generated"].reset_index(drop=True)
    return sheets


def test_cli_team_dir_writes_one_workbook_per_team_and_keeps_combined(
    snapshot, tmp_path, monkeypatch
):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    plain = tmp_path / "plain" / "w1.xlsx"
    split = tmp_path / "split" / "w1.xlsx"
    teams = tmp_path / "split" / "teams"
    prev = tmp_path / "w0.xlsx"
    tool.write_excel(prev, tool.build_report(snapshot, CAPACITY, settings()), {})
    snapshot.items[3] = item(4, TEAM_Y, "Dormant and unowned", rw=10)  # resolved since w0
    since = ("--previous-file", str(prev))
    assert tool.main(cli(tmp_path, "--output", str(plain), *since)) == 0
    assert tool.main(cli(tmp_path, "--output", str(split), *since, "--team-dir", str(teams))) == 0
    assert tool.SHEET_RESOLVED in load_workbook(plain).sheetnames

    expected, actual = workbook_sheets(plain), workbook_sheets(split)
    assert list(expected) == list(actual)
    for name in expected:
        pd.testing.assert_frame_equal(expected[name], actual[name])

    files = sorted(p.name for p in teams.iterdir())
    assert files == ["w1.Proj_TeamX.xlsx", "w1.Proj_TeamY.xlsx"]
    team_x = pd.read_excel(teams / files[0], sheet_name=None, header=2)
    assert set(team_x[tool.SHEET_MASTER][tool.FIELD_ID]) == {1, 2, 3}
    config = pd.read_excel(teams / files[0], sheet_name=tool.SHEET_CONFIG, header=2)
    assert config.iloc[0, :2].tolist() == ["Team", TEAM_X]
    # A team workbook is a valid --previous-file for that team's next run.
    assert set(tool.read_previous_master(teams / files[1])[tool.FIELD_ID]) == {4, 5, 6}


def test_cli_team_dir_refuses_existing_files_before_writing_any(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    out = tmp_path / "w1.csv"
    teams = tmp_path / "teams"
    teams.mkdir()
    existing = teams / "w1.Proj_TeamY.csv"
    existing.write_text("keep")
    args = cli(tmp_path, "--output", str(out), "--format", "csv", "--team-dir", str(teams))
    assert tool.main(args) == tool.EXIT_VALIDATION
    assert not out.exists() and existing.read_text() == "keep"
    assert sorted(p.name for p in teams.iterdir()) == ["w1.Proj_TeamY.csv"]

    assert tool.main([*args, "--overwrite"]) == 0
    assert set(pd.read_csv(existing)[tool.FIELD_ID]) == {4, 5, 6}
    assert set(pd.read_csv(teams / "w1.Proj_TeamX.csv")[tool.FIELD_ID]) == {1, 2, 3}
    assert len(pd.read_csv(out)) == len(snapshot.items)


def test_cli_team_dir_guards_inputs(snapshot, tmp_path, monkeypatch):
    def unreachable(url, query):
        raise AssertionError("network must not be touched")

    monkeypatch.setattr(tool, "fetch_snapshot", unreachable)
    not_a_dir = tmp_path / "teams"
    not_a_dir.write_text("")
    for team_dir in (not_a_dir, not_a_dir / "nested"):
        args = cli(tmp_path, "--output", str(tmp_path / "r.xlsx"), "--team-dir", str(team_dir))
        assert tool.main(args) == tool.EXIT_VALIDATION

    # A team file may not replace the previous report it is computed from.
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    prev = tmp_path / "w1.Proj_TeamX.xlsx"
    tool.write_excel(prev, tool.build_report(snapshot, CAPACITY, settings()), {})
    before = prev.read_bytes()
    args = cli(
        tmp_path,
        *("--output", str(tmp_path / "out" / "w1.xlsx"), "--team-dir", str(tmp_path)),
        *("--previous-file", str(prev), "--overwrite"),
    )
    assert tool.main(args) == tool.EXIT_VALIDATION
    assert prev.read_bytes() == before and not (tmp_path / "out").exists()


def test_cli_team_write_failure_leaves_no_output(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    real_write = tool.write_excel

    def failing_for_team_y(path, report, meta):
        if report.team == TEAM_Y:
            raise OSError("disk full")
        real_write(path, report, meta)

    monkeypatch.setattr(tool, "write_excel", failing_for_team_y)
    out = tmp_path / "w1.xlsx"
    teams = tmp_path / "teams"
    args = cli(tmp_path, "--output", str(out), "--team-dir", str(teams))
    assert tool.main(args) == tool.EXIT_IO
    assert not out.exists() and list(teams.iterdir()) == []  # no outputs, no temp files
    assert [p.name for p in tmp_path.iterdir() if p.is_file()] == []

    monkeypatch.setattr(tool, "write_excel", real_write)
    assert tool.main(args) == 0  # a plain rerun succeeds


# --- fetch_snapshot against a fake ADO client (no server, no Kerberos libraries) ---

azure_connection = pytest.importorskip("azure.devops.connection")
from types import ModuleType, SimpleNamespace  # noqa: E402

from azure.devops.exceptions import (  # noqa: E402
    AzureDevOpsAuthenticationError,
    AzureDevOpsClientRequestError,
)
from msrest.exceptions import ClientRequestError  # noqa: E402


def ado_item(wi_id, fields, relations=()):
    rels = [SimpleNamespace(rel=r["rel"], url=r["url"]) for r in relations]
    return SimpleNamespace(id=wi_id, fields=fields, relations=rels)


class FakeClient:
    """Mimics WorkItemTrackingClient: get_work_items raises on unknown ids unless omitted."""

    def __init__(self, result, store, fail_on=None):
        self.result, self.store, self.fail_on = result, store, fail_on
        self.calls = []

    def query_by_id(self, id):
        if self.fail_on == "query":
            raise self.error
        return self.result

    def get_work_items(self, ids, fields=None, expand=None, error_policy=None):
        self.calls.append({"ids": ids, "fields": fields, "expand": expand})
        assert len(ids) <= 200 and not (fields and expand)  # ADO API limits
        if error_policy == "omit":
            return [self.store.get(i) for i in ids]
        return [self.store[i] for i in ids if i in self.store]


@pytest.fixture
def fake_ado(monkeypatch):
    """Install a fake Connection and a stub requests_kerberos; return a setup function."""
    kerberos = ModuleType("requests_kerberos")
    kerberos.OPTIONAL = "optional"
    kerberos.HTTPKerberosAuth = type(
        "HTTPKerberosAuth", (), {"__init__": lambda self, **kw: setattr(self, "kw", kw)}
    )
    monkeypatch.setitem(__import__("sys").modules, "requests_kerberos", kerberos)
    state = {}

    def setup(client):
        def connection(base_url, creds):
            state["base_url"], state["session"] = base_url, creds.signed_session()
            if client.fail_on == "connect":
                raise client.error
            return SimpleNamespace(
                clients=SimpleNamespace(get_work_item_tracking_client=lambda: client)
            )

        monkeypatch.setattr(azure_connection, "Connection", connection)
        return state

    setup.kerberos = kerberos
    return setup


def link(source, target, rel=None):
    end = lambda i: None if i is None else SimpleNamespace(id=i)  # noqa: E731
    return SimpleNamespace(source=end(source), target=end(target), rel=rel)


def test_fetch_tree_query_collects_items_links_parents_and_warnings(fake_ado):
    store = {
        1: ado_item(1, {tool.FIELD_TITLE: "Root"}, [rel(tool.REL_SUCCESSOR, 50)]),
        2: ado_item(2, {tool.FIELD_TITLE: "Child", tool.FIELD_PARENT: 1}),
        50: ado_item(50, {tool.FIELD_TITLE: "Outside", tool.FIELD_AREA_PATH: TEAM_Y}),
    }
    store[1].relations.append(SimpleNamespace(rel=tool.REL_BLOCKS_FORWARD, url=f"{BASE}/x/77"))
    # Tree results have a root link with no source and no flat work_items list.
    result = SimpleNamespace(work_item_relations=[link(None, 1), link(1, 2)], work_items=None)
    client = FakeClient(result, store)
    state = fake_ado(client)

    snap = tool.fetch_snapshot(BASE, "q")

    assert sorted(i["id"] for i in snap.items) == [1, 2]
    assert snap.external == {50: store[50].fields}
    assert snap.parent_titles == {1: "Root"}
    assert snap.warnings == ["1 linked items unavailable (deleted or no permission)"]
    assert client.calls[0]["expand"] == "Relations"
    assert isinstance(state["session"].auth, fake_ado.kerberos.HTTPKerberosAuth)
    assert state["session"].auth.kw == {"mutual_authentication": "optional"}
    assert state["base_url"] == BASE


def test_fetch_flat_query_batches_by_200(fake_ado):
    store = {i: ado_item(i, {tool.FIELD_TITLE: str(i)}) for i in range(1, 451)}
    result = SimpleNamespace(
        work_item_relations=None, work_items=[SimpleNamespace(id=i) for i in store]
    )
    client = FakeClient(result, store)
    fake_ado(client)
    assert len(tool.fetch_snapshot(BASE, "q").items) == 450
    assert [len(c["ids"]) for c in client.calls] == [200, 200, 50]


def test_fetch_empty_query_and_missing_core_items(fake_ado):
    fake_ado(FakeClient(SimpleNamespace(work_item_relations=[], work_items=[]), {}))
    with pytest.raises(tool.EmptyQueryError):
        tool.fetch_snapshot(BASE, "q")

    result = SimpleNamespace(work_item_relations=None, work_items=[SimpleNamespace(id=1)])
    fake_ado(FakeClient(result, {}))
    with pytest.raises(tool.AdoError, match="could not be fetched"):
        tool.fetch_snapshot(BASE, "q")


@pytest.mark.parametrize(
    ("where", "error"),
    [
        ("connect", ClientRequestError("connection refused")),
        ("query", AzureDevOpsAuthenticationError("401 Unauthorized")),
        ("query", AzureDevOpsClientRequestError("404 query not found")),
    ],
)
def test_fetch_maps_client_failures_to_ado_error(fake_ado, where, error):
    client = FakeClient(None, {}, fail_on=where)
    client.error = error
    fake_ado(client)
    with pytest.raises(tool.AdoError):
        tool.fetch_snapshot(BASE, "q")


def test_control_characters_in_fields_do_not_break_the_workbook(tmp_path):
    bad = item(1, TEAM_X, "Bad\x0btitle", **{"System.Description": "a\x01b"})
    snap = tool.Snapshot(items=[bad])
    out = tmp_path / "r.xlsx"
    tool.write_excel(out, tool.build_report(snap, {}, settings()), {"As of": "2026-10-01"})
    master = load_workbook(out)[tool.SHEET_MASTER]
    titles = [c.value for c in master[tool.DATA_START_ROW - 1]]
    assert master.cell(tool.DATA_START_ROW, titles.index(tool.FIELD_TITLE) + 1).value == "Badtitle"


# ---------------------------------------------------------------------------
# JSON configuration (--config)
# ---------------------------------------------------------------------------

CUSTOM_UAT = "Custom.UATTargetDate"


def write_config(path, config):
    path.write_text(json.dumps(config), encoding="utf-8")
    return path


def test_config_changes_weights_hour_bands_and_rag(snapshot, tmp_path):
    default = tool.build_report(snapshot, CAPACITY, settings())
    config = {
        "weights": {"DATE_DRIFT": 10, "BLOCKS_OTHERS": 0},
        "hours": {"quick_win_max": 4, "minor_progress_below": 12},
        "rag": {"red_above_pct": 200, "amber_from_pct": 150},
    }
    tuned = tool.build_report(
        snapshot, CAPACITY, settings(**tool.load_config(write_config(tmp_path / "c.json", config)))
    )
    before, after = by_id(default), by_id(tuned)

    # Item 1 drifts and blocks item 2: 3 + 3 by default, 10 + 0 when configured.
    assert {"DATE_DRIFT", "BLOCKS_OTHERS"} <= set(after.loc[1, tool.COL_FLAGS].split("|"))
    assert after.loc[1, tool.COL_SCORE] - before.loc[1, tool.COL_SCORE] == 10 - 3 - 3
    # Item 3 has 5 h left: a quick win in 1-20 h, not in 1-4 h.
    assert "QUICK_WIN" in before.loc[3, tool.COL_OPP_FLAGS]
    assert "QUICK_WIN" not in after.loc[3, tool.COL_OPP_FLAGS]
    # Item 1 logged 10 h and item 3 15 h: minor progress below 16 h, only item 1 below 12 h.
    minor = "8. Minor Progress"
    assert set(default.sheets[minor][tool.FIELD_ID]) == {1, 3}
    assert set(tuned.sheets[minor][tool.FIELD_ID]) == {1}
    # TeamX carries 155 % load: RED by default, AMBER under 150-200 bands.
    rag = {
        r.sheets[tool.SHEET_SUMMARY].set_index("Team").loc[TEAM_X, tool.COL_RAG]
        for r in (default, tuned)
    }
    assert rag == {"RED", "AMBER"}
    assert {h.flag: h.weight for h in tuned.legend}["DATE_DRIFT"] == 10


def test_config_reads_renamed_fields_and_labels_them_in_the_report(tmp_path):
    items = [item(1, TEAM_X, "Late UAT", rw=10, **{CUSTOM_UAT: iso(10)})]
    config = write_config(tmp_path / "c.json", {"fields": {"uat_expected_date": CUSTOM_UAT}})
    cfg = settings(**tool.load_config(config))
    report = tool.build_report(tool.Snapshot(items=items), CAPACITY, cfg)
    assert by_id(report).loc[1, tool.COL_FLAGS] == "DATE_DRIFT"
    drift = report.sheets["9. Date Drift Risk"]
    assert CUSTOM_UAT in drift.columns and tool.FIELD_UAT_EXPECTED not in drift.columns
    master = report.sheets[tool.SHEET_MASTER]
    assert list(master.columns).count(CUSTOM_UAT) == 1


def test_cli_config_reaches_the_workbook_and_csv(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    snapshot.items[0]["fields"][CUSTOM_UAT] = snapshot.items[0]["fields"].pop(
        tool.FIELD_UAT_EXPECTED
    )
    config = write_config(
        tmp_path / "c.json",
        {"fields": {"uat_expected_date": CUSTOM_UAT}, "weights": {"DATE_DRIFT": 7}},
    )
    out = tmp_path / "r.xlsx"
    assert tool.main(cli(tmp_path, "--config", str(config), "--output", str(out))) == 0
    ws = load_workbook(out)[tool.SHEET_CONFIG]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert ["Config file", str(config)] in [row[:2] for row in values]
    assert ["DATE_DRIFT", "risk", 7] in [row[:3] for row in values]

    csv = tmp_path / "r.csv"
    args = cli(tmp_path, "--config", str(config), "--format", "csv", "--output", str(csv))
    assert tool.main(args) == 0
    frame = pd.read_csv(csv)
    assert CUSTOM_UAT in frame.columns
    assert frame.set_index(tool.FIELD_ID).loc[1, tool.COL_SCORE] >= 7


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("{not json", "not valid JSON"),
        ("[]", "top level"),
        ('{"weight": {}}', "unknown keys"),
        ('{"weights": {"QUICK_WIN": 2}}', "unknown keys"),
        ('{"weights": {"STALE": 1.5}}', "non-negative integer"),
        ('{"weights": {"STALE": true}}', "non-negative integer"),
        ('{"hours": {"quick_win_min": 30}}', "exceeds quick_win_max"),
        ('{"hours": {"minor_progress_below": -1}}', "negative"),
        ('{"hours": {"quick_win_max": NaN}}', "finite number"),
        ('{"rag": {"amber_from_pct": 120}}', "exceeds red_above_pct"),
        ('{"fields": {"uat_ready_date": ""}}', "non-empty"),
        ('{"fields": {"uat_ready_date": "my.UATExpectedDate"}}', "several fields"),
        ('{"fields": {"remaining_work": "System.State"}}', "several fields"),
        ('{"fields": {"uat_date": "Custom.X"}}', "unknown keys"),
    ],
)
def test_invalid_config_exits_before_any_service_call(
    tmp_path, no_network, capsys, content, message
):
    config = tmp_path / "c.json"
    config.write_text(content, encoding="utf-8")
    out = tmp_path / "r.xlsx"
    assert tool.main(cli(tmp_path, "--config", str(config), "--output", str(out))) == 2
    assert message in capsys.readouterr().err
    assert not out.exists()


def test_missing_config_and_config_as_output_are_validation_errors(tmp_path, no_network):
    missing = cli(
        tmp_path, "--config", str(tmp_path / "nope.json"), "--output", str(tmp_path / "r.xlsx")
    )
    assert tool.main(missing) == tool.EXIT_VALIDATION
    config = write_config(tmp_path / "c.json", {})
    same = cli(tmp_path, "--config", str(config), "--output", str(config), "--overwrite")
    assert tool.main(same) == tool.EXIT_VALIDATION
    assert json.loads(config.read_text()) == {}


def test_swapped_field_names_are_read_from_the_original_fields(tmp_path):
    fields = {
        "original_estimate": tool.FIELD_COMPLETED_WORK,
        "completed_work": tool.FIELD_ORIGINAL_ESTIMATE,
    }
    cfg = settings(**tool.load_config(write_config(tmp_path / "c.json", {"fields": fields})))
    items = [item(1, TEAM_X, "Swapped", oe=30, cw=5, rw=1)]
    master = by_id(tool.build_report(tool.Snapshot(items=items), CAPACITY, cfg))
    assert master.loc[1, tool.FIELD_ORIGINAL_ESTIMATE] == 5
    assert master.loc[1, tool.FIELD_COMPLETED_WORK] == 30
    assert "OVER_ESTIMATE" in master.loc[1, tool.COL_FLAGS]


def test_options_file_sets_thresholds_and_command_line_wins(snapshot, tmp_path, monkeypatch):
    monkeypatch.setattr(tool, "fetch_snapshot", lambda url, query: snapshot)
    config = write_config(tmp_path / "c.json", {"weights": {"DATE_DRIFT": 7}})
    options = tmp_path / "triage.ini"
    options.write_text(
        f"stale-days = 5\nagenda-size = 3\nconfig = {config}\nprofile = exec\n",
        encoding="utf-8",
    )
    out = tmp_path / "r.xlsx"
    args = cli(tmp_path, "--options", str(options), "--agenda-size", "4", "--output", str(out))
    assert tool.main(args) == tool.EXIT_OK
    rows = load_workbook(out)[tool.SHEET_CONFIG].iter_rows(values_only=True)
    meta = {row[0]: row[1] for row in rows if row}
    assert meta["Stale threshold (days)"] == "5"
    assert meta["Agenda size"] == "4"
    assert meta["Profile"] == "exec"
    assert meta["Options file"] == str(options)
    assert meta["Config file"] == str(config)


def test_written_options_and_config_reproduce_the_settings(tmp_path, no_network, capsys):
    options, config = tmp_path / "triage.ini", tmp_path / "triage.json"
    argv = ["export", "--stale-days", "30", "--profile", "exec", "--write-options", str(options)]
    assert tool.main(argv) == tool.EXIT_OK
    text = options.read_text(encoding="utf-8")
    assert "stale-days = 30\n" in text and "# output =\n" in text and "overwrite =" not in text
    assert "write-config" not in text
    parsed, _ = tool.options.parse(tool.parser(), ["export", "--options", str(options)])
    assert (parsed.stale_days, parsed.profile, parsed.timebox_days) == (30, "exec", 60)

    custom = write_config(tmp_path / "c.json", {"rag": {"amber_from_pct": 80}})
    argv = ["export", "--config", str(custom), "--write-config", str(config)]
    assert tool.main(argv) == tool.EXIT_OK
    template = json.loads(config.read_text(encoding="utf-8"))
    assert template["rag"] == {"red_above_pct": 100.0, "amber_from_pct": 80.0}
    assert template["weights"]["DATE_DRIFT"] == 3
    assert template["fields"]["remaining_work"] == tool.FIELD_REMAINING_WORK
    effective = tool.config_template(tool.load_config(custom))
    assert tool.config_template(tool.load_config(config)) == effective == template

    assert tool.main(argv) == tool.EXIT_VALIDATION  # existing file without --overwrite
    assert "exists; pass --overwrite" in capsys.readouterr().err
    assert tool.main([*argv, "--overwrite"]) == tool.EXIT_OK


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("stale-day = 5\n", "unknown option stale-day"),
        ("overwrite = true\n", "overwrite must be given on the command line"),
        ("agenda-size = many\n", "invalid int value"),
    ],
)
def test_options_file_rejects_bad_keys(tmp_path, no_network, capsys, content, message):
    options = tmp_path / "bad.ini"
    options.write_text(content, encoding="utf-8")
    out = tmp_path / "r.xlsx"
    with pytest.raises(SystemExit) as exc:
        tool.main(cli(tmp_path, "--options", str(options), "--output", str(out)))
    assert exc.value.code == tool.EXIT_VALIDATION
    assert message in capsys.readouterr().err
    assert not out.exists()


def test_output_is_required_and_may_not_be_the_options_file(tmp_path, no_network, capsys):
    with pytest.raises(SystemExit) as exc:
        tool.main(cli(tmp_path))
    assert exc.value.code == tool.EXIT_VALIDATION
    assert "required: --output" in capsys.readouterr().err
    options = tmp_path / "triage.ini"
    options.write_text("stale-days = 5\n", encoding="utf-8")
    same = cli(tmp_path, "--options", str(options), "--output", str(options), "--overwrite")
    assert tool.main(same) == tool.EXIT_VALIDATION
    snap = [*cli(tmp_path, "--output", str(tmp_path / "r.xlsx")), "--save-snapshot", str(options)]
    assert tool.main([*snap, "--options", str(options), "--overwrite"]) == tool.EXIT_VALIDATION
    assert options.read_text(encoding="utf-8") == "stale-days = 5\n"
