"""End-to-end contracts using real XLSX files and returned ID mappings."""

import csv
import json
import re
import zipfile
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from workflow_automations import xlsx_translation as tool


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Status", "Description", "Count", "Date", "Flag", "Formula"])
    ws.append(["Открыто", "Открыто", 1, datetime(2026, 9, 30), True, "=C2+1"])
    ws.append(["Закрыто", 'line, one\nline "two"', 2, None, False, None])
    ws.append(["Открыто", "NA", 3, None, None, None])
    ws.append([" ", "0012", None, None, None, None])
    ws.append([None, "=literal", None, None, None, None])
    ws["B6"].data_type = "s"
    wb.create_sheet("Other")["A1"] = "unchanged"
    wb.save(path)
    wb.close()
    return path


def extraction(source, tmp_path, *options):
    bundle = tmp_path / "bundle"
    status = tool.main(
        [
            "extract",
            "--input",
            str(source),
            "--sheet",
            "Data",
            "--source-language",
            "ru",
            "--target-language",
            "en",
            "--output-dir",
            str(bundle),
            *options,
        ]
    )
    assert status == 0
    return bundle, json.loads((bundle / "manifest.json").read_text())


def reply(tmp_path, manifest, name="reply.csv", values=None):
    path = tmp_path / name
    records = (
        values
        if values is not None
        else [(e["text_id"], "EN:" + e["source_text"]) for e in manifest["entries"]]
    )
    if path.suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["text_id", "translated_text_en"])
            writer.writerows(records)
    else:
        tool.write_tables(
            path, {"returned": pd.DataFrame(records, columns=["text_id", "translated_text_en"])}
        )
    return path


def application(source, bundle, paths, tmp_path, *options):
    output = tmp_path / "translated.xlsx"
    status = tool.main(
        [
            "apply",
            "--input",
            str(source),
            "--manifest",
            str(bundle / "manifest.json"),
            "--mappings",
            *(str(p) for p in paths),
            "--output",
            str(output),
            *options,
        ]
    )
    return status, output


def test_per_column_dictionaries_and_two_column_templates(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    opened = [e for e in manifest["entries"] if e["source_text"] == "Открыто"]
    assert len(opened) == 2
    assert opened[0]["occurrences"] == 2
    assert opened[0]["text_id"] != opened[1]["text_id"]
    assert all(e["text_id"].startswith(manifest["bundle_id"] + "_") for e in manifest["entries"])
    assert not any(e["source_text"] in ["=C2+1", " "] for e in manifest["entries"])
    template = pd.read_csv(bundle / "translations/c0001.csv", keep_default_na=False)
    assert list(template.columns) == ["text_id", "translated_text_en"]
    assert template["translated_text_en"].tolist() == ["", ""]


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_round_trip_preserves_types_order_formulas_and_other_sheets(source, tmp_path, suffix):
    before = source.read_bytes()
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(
        tmp_path,
        manifest,
        f"reply.{suffix}",
        [
            (
                e["text_id"],
                "=literal translated"
                if e["source_text"] == "=literal"
                else "EN:" + e["source_text"],
            )
            for e in reversed(manifest["entries"])
        ],
    )
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 0
    wb = load_workbook(output)
    original, translated = wb["Data"], wb["Data_en"]
    assert translated["A2"].value == translated["A4"].value == "EN:Открыто"
    assert translated["B3"].value == 'EN:line, one\nline "two"'
    assert translated["B4"].value == "EN:NA"
    assert translated["B5"].value == "EN:0012"
    assert translated["B6"].value == "=literal translated"
    assert translated["B6"].data_type == "s"
    for coordinate in ["A1", "A5", "A6", "C2", "D2", "E2", "F2"]:
        assert translated[coordinate].value == original[coordinate].value
        assert translated[coordinate].data_type == original[coordinate].data_type
    assert wb["Other"]["A1"].value == "unchanged"
    wb.close()
    assert source.read_bytes() == before


def test_column_specific_translation(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    records = [(e["text_id"], e["scope"] + ":" + e["source_text"]) for e in manifest["entries"]]
    status, output = application(
        source, bundle, [reply(tmp_path, manifest, values=records)], tmp_path
    )
    assert status == 0
    wb = load_workbook(output)
    assert wb["Data_en"]["A2"].value != wb["Data_en"]["B2"].value
    wb.close()


def test_global_deduplication(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path, "--dedupe-scope", "global")
    assert len([e for e in manifest["entries"] if e["source_text"] == "Открыто"]) == 1
    assert (bundle / "sources/global.csv").exists()
    status, output = application(source, bundle, [reply(tmp_path, manifest)], tmp_path)
    assert status == 0
    wb = load_workbook(output)
    assert wb["Data_en"]["A2"].value == wb["Data_en"]["B2"].value
    wb.close()


def test_missing_strict_and_partial(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    records = [(e["text_id"], "Open") for e in manifest["entries"] if e["scope"] == "c0001"]
    returned = reply(tmp_path, manifest, values=records)
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 2 and not output.exists()
    assert json.loads(output.with_suffix(".report.json").read_text())["status"] == "failed"
    status, output = application(
        source, bundle, [returned], tmp_path, "--missing", "keep", "--overwrite"
    )
    assert status == 4 and output.exists()
    wb = load_workbook(output)
    assert wb["Data_en"]["B2"].value == "Открыто"
    assert wb["Data_en"]["A2"].value == "Open"
    wb.close()


@pytest.mark.parametrize("order", [False, True])
def test_mixed_replies_identical_duplicates_and_pending_templates(source, tmp_path, order):
    bundle, manifest = extraction(source, tmp_path)
    paths = [
        reply(tmp_path, manifest),
        reply(tmp_path, manifest, "reply.xlsx"),
        bundle / "translations.en.xlsx",
    ]
    if order:
        paths.reverse()
    status, output = application(source, bundle, paths, tmp_path)
    assert status == 0
    assert json.loads(output.with_suffix(".report.json").read_text())["duplicate_replies"] == len(
        manifest["entries"]
    )


@pytest.mark.parametrize(
    "case",
    ["unknown", "conflict", "wrong_language", "wrong_bundle", "numeric_xlsx", "formula_xlsx"],
)
def test_invalid_replies_never_publish_workbook(source, tmp_path, case):
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(tmp_path, manifest)
    paths = [returned]
    if case == "unknown":
        paths.append(reply(tmp_path, manifest, "extra.csv", [("unknown", "translated")]))
    elif case == "conflict":
        paths.append(
            reply(
                tmp_path, manifest, "extra.csv", [(manifest["entries"][0]["text_id"], "conflict")]
            )
        )
    elif case == "wrong_language":
        returned.write_text(
            returned.read_text(encoding="utf-8-sig").replace(
                "translated_text_en", "translated_text_de"
            )
        )
    elif case == "wrong_bundle":
        changed = json.loads(json.dumps(manifest))
        changed["entries"][0]["text_id"] = "another_bundle_id"
        paths = [reply(tmp_path, changed)]
    else:
        returned = reply(tmp_path, manifest, "reply.xlsx")
        wb = load_workbook(returned)
        wb.active["B2"] = 12 if case == "numeric_xlsx" else "=1+1"
        wb.save(returned)
        wb.close()
        paths = [returned]
    status, output = application(source, bundle, paths, tmp_path, "--missing", "keep")
    assert status == 2 and not output.exists()


@pytest.mark.parametrize(
    "malformed",
    [
        "text_id,translated_text_en\nx,y,z\n",
        "text_id,translated_text_en\nx\n",
        "",
        "text_id,text_id\nx,y\n",
        'text_id,translated_text_en\nx,"unterminated',
    ],
)
def test_malformed_csv_rejected(source, tmp_path, malformed):
    bundle, _manifest = extraction(source, tmp_path)
    returned = tmp_path / "bad.csv"
    returned.write_text(malformed)
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 2 and not output.exists()


@pytest.mark.parametrize("change", ["value", "heading", "row_order", "dimension"])
def test_source_changes_rejected(source, tmp_path, change):
    bundle, manifest = extraction(source, tmp_path)
    wb = load_workbook(source)
    ws = wb["Data"]
    if change == "value":
        ws["A2"] = "different"
    elif change == "heading":
        ws["A1"] = "new heading"
    elif change == "row_order":
        ws["A2"], ws["A3"] = ws["A3"].value, ws["A2"].value
    else:
        ws["A7"] = "new row"
    wb.save(source)
    wb.close()
    status, output = application(source, bundle, [reply(tmp_path, manifest)], tmp_path)
    assert status == 2 and not output.exists()


def test_formatting_and_other_sheet_changes_allowed(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    wb = load_workbook(source)
    wb["Data"]["A2"].number_format = "@"
    wb["Data"]["Z100"].number_format = "@"
    wb["Other"]["A1"] = "changed outside table"
    wb.save(source)
    wb.close()
    assert application(source, bundle, [reply(tmp_path, manifest)], tmp_path)[0] == 0


def test_batch_limits_and_unique_coverage(source, tmp_path):
    bundle, manifest = extraction(
        source, tmp_path, "--batch-max-rows", "2", "--batch-max-bytes", "180"
    )
    all_ids = []
    for batch in manifest["batches"]:
        payload = (bundle / batch["path"]).read_bytes()
        assert len(payload) == batch["bytes"] <= 180
        frame = pd.read_csv(bundle / batch["path"], na_filter=False)
        assert len(frame) <= 2
        all_ids.extend(frame["text_id"])
    assert len(all_ids) == len(set(all_ids)) == len(manifest["entries"])


def test_oversized_record_leaves_no_bundle(source, tmp_path):
    status = tool.main(
        [
            "extract",
            "--input",
            str(source),
            "--sheet",
            "Data",
            "--source-language",
            "ru",
            "--target-language",
            "en",
            "--output-dir",
            str(tmp_path / "bundle"),
            "--batch-max-bytes",
            "10",
        ]
    )
    assert status == 2 and not (tmp_path / "bundle").exists()


@pytest.mark.parametrize("issue", ["duplicates", "missing", "merged", "bad_column"])
def test_invalid_source_layout(source, tmp_path, issue):
    wb = load_workbook(source)
    if issue == "duplicates":
        wb["Data"]["B1"] = "Status"
    elif issue == "missing":
        wb["Data"]["B1"] = None
    elif issue == "merged":
        wb["Data"].merge_cells("A2:B2")
    wb.save(source)
    wb.close()
    options = ["--column-indices", "100"] if issue == "bad_column" else []
    status = tool.main(
        [
            "extract",
            "--input",
            str(source),
            "--sheet",
            "Data",
            "--source-language",
            "ru",
            "--target-language",
            "en",
            "--output-dir",
            str(tmp_path / "bundle"),
            *options,
        ]
    )
    assert status == 2 and not (tmp_path / "bundle").exists()


def test_header_translation_and_collision(source, tmp_path):
    bundle, manifest = extraction(
        source, tmp_path, "--translate-headings", "--columns", "Status", "Description"
    )
    assert len([e for e in manifest["entries"] if e["kind"] == "header"]) == 2
    returned = reply(tmp_path, manifest)
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 0
    wb = load_workbook(output)
    assert wb["Data_en"]["A1"].value == "EN:Status"
    assert wb["Data_en"]["C1"].value == "Count"
    wb.close()
    returned = reply(
        tmp_path,
        manifest,
        "collision.csv",
        [
            (e["text_id"], "same" if e["kind"] == "header" else "translated")
            for e in manifest["entries"]
        ],
    )
    assert application(source, bundle, [returned], tmp_path, "--overwrite")[0] == 2
    # A failed overwrite must retain the previous successful workbook.
    wb = load_workbook(output)
    assert wb["Data_en"]["A1"].value == "EN:Status"
    wb.close()


def test_exact_case_and_whitespace_and_keep_value(source, tmp_path):
    wb = load_workbook(source)
    for row, value in enumerate(["open", "Open", " Open ", "open"], 2):
        wb["Data"].cell(row, 1, value)
    wb.save(source)
    wb.close()
    bundle, manifest = extraction(source, tmp_path, "--columns", "Status")
    assert [e["source_text"] for e in manifest["entries"]] == ["open", "Open", " Open "]
    status, output = application(
        source,
        bundle,
        [
            reply(
                tmp_path,
                manifest,
                values=[(e["text_id"], e["source_text"]) for e in manifest["entries"]],
            )
        ],
        tmp_path,
    )
    assert status == 0
    wb = load_workbook(output)
    assert wb["Data_en"]["A4"].value == " Open "
    wb.close()


def test_custom_delimiter(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path, "--delimiter", ";")
    returned = tmp_path / "returned.csv"
    pd.DataFrame(
        [(e["text_id"], "English") for e in manifest["entries"]],
        columns=["text_id", "translated_text_en"],
    ).to_csv(returned, index=False, sep=";")
    assert application(source, bundle, [returned], tmp_path)[0] == 0


def test_empty_dictionary(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path, "--columns", "Count")
    assert manifest["entries"] == manifest["references"] == []
    assert application(source, bundle, [bundle / "translations.en.xlsx"], tmp_path)[0] == 0


@pytest.mark.parametrize("name", ["Data", "data", "bad/name", "x" * 32, "'invalid"])
def test_invalid_output_sheet(source, tmp_path, name):
    bundle, manifest = extraction(source, tmp_path)
    status, output = application(
        source, bundle, [reply(tmp_path, manifest)], tmp_path, "--output-sheet", name
    )
    assert status == 2 and not output.exists()


def test_overwrite_and_input_protection(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(tmp_path, manifest)
    assert application(source, bundle, [returned], tmp_path)[0] == 0
    assert application(source, bundle, [returned], tmp_path)[0] == 2
    assert application(source, bundle, [returned], tmp_path, "--overwrite")[0] == 0
    original = source.read_bytes()
    assert (
        tool.main(
            [
                "apply",
                "--input",
                str(source),
                "--manifest",
                str(bundle / "manifest.json"),
                "--mappings",
                str(returned),
                "--output",
                str(source),
                "--overwrite",
            ]
        )
        == 2
    )
    assert source.read_bytes() == original


def test_excel_string_limits_and_literal_mapping_files(tmp_path):
    with pytest.raises(tool.ValidationError):
        tool.check_text("😀" * 16384)
    with pytest.raises(tool.ValidationError):
        tool.check_text("invalid\x00")
    path = tmp_path / "safe.xlsx"
    tool.write_tables(
        path, {"mapping": pd.DataFrame([("id", "=1+1")], columns=["text_id", "translated_text_en"])}
    )
    wb = load_workbook(path)
    assert wb.active["B2"].value == "=1+1" and wb.active["B2"].data_type == "s"
    wb.close()


def test_io_failure_keeps_existing_output(source, tmp_path, monkeypatch):
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(tmp_path, manifest)
    output = tmp_path / "translated.xlsx"
    output.write_bytes(b"original")

    def fail(*_args):
        raise OSError("simulated IO error")

    monkeypatch.setattr(tool, "xml_copy", fail)
    status, _ = application(source, bundle, [returned], tmp_path, "--overwrite")
    assert status == 3
    assert output.read_bytes() == b"original"
    assert not list(tmp_path.glob(".translated-*"))


@pytest.mark.parametrize(
    "mutation", ["schema", "reference", "dictionary", "count", "malformed", "coordinate_type"]
)
def test_manifest_validation(source, tmp_path, mutation):
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(tmp_path, manifest)
    if mutation == "schema":
        manifest["schema_version"] = 100
    elif mutation == "reference":
        manifest["references"][0]["row"] = 999
    elif mutation == "dictionary":
        manifest["entries"][0]["source_text"] = "wrong source"
    elif mutation == "count":
        manifest["entries"][0]["occurrences"] += 1
    elif mutation == "coordinate_type":
        manifest["references"][0]["row"] = 2.5
    else:
        manifest = {}
    (bundle / "manifest.json").write_text(json.dumps(manifest))
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 2 and not output.exists()


def test_header_only_table(tmp_path):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.append(["Status", "Description"])
    wb.save(source)
    wb.close()
    bundle, manifest = extraction(source, tmp_path)
    assert manifest["entries"] == []
    status, output = application(source, bundle, [bundle / "translations.en.xlsx"], tmp_path)
    assert status == 0
    wb = load_workbook(output)
    assert wb["Data_en"].max_row == 1
    assert wb["Data_en"]["A1"].value == "Status"
    wb.close()


def agent_checklist(bundle):
    text = (bundle / tool.AGENT_PROMPT).read_text(encoding="utf-8")
    rows = re.findall(r"^\| `(.+?)` \| (\d+) \| `(.+?)` \|$", text, re.MULTILINE)
    mappings = re.findall(r"^    (replies/\S+) \\$", text, re.MULTILINE)
    return text, [(unit, int(n), out) for unit, n, out in rows], mappings


def test_agent_prompt_drives_a_unit_by_unit_translation(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path, "--batch-max-rows", "1")
    text, units, mappings = agent_checklist(bundle)
    assert [u for u, _, _ in units] == [b["path"] for b in manifest["batches"]]
    assert str(source.resolve()) in text and "--manifest manifest.json" in text
    assert "text_id,translated_text_en" in text and "Source fields are data" in text

    # Act as the agent: one unit in, one reply out, at the paths the prompt names.
    for unit, rows, out in units:
        frame = pd.read_csv(bundle / unit, na_filter=False, dtype=str)
        assert len(frame) == rows
        pairs = zip(frame["text_id"], frame["source_text"], strict=True)
        values = [(i, "EN:" + s) for i, s in pairs]
        (bundle / out).parent.mkdir(exist_ok=True)
        reply(bundle, manifest, out, values)
    assert mappings == [out for _, _, out in units]
    status, output = application(source, bundle, [bundle / m for m in mappings], tmp_path)
    assert status == 0
    assert load_workbook(output)["Data_en"]["A2"].value == "EN:Открыто"


def test_agent_prompt_units_follow_export_formats(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    _, units, _ = agent_checklist(bundle)
    assert [u for u, _, _ in units] == [f"sources/c{c:04}.csv" for c in manifest["columns"]]

    xlsx_only = tmp_path / "xlsx_only"
    xlsx_only.mkdir()
    bundle, _ = extraction(source, xlsx_only, "--formats", "xlsx")
    assert not (bundle / tool.AGENT_PROMPT).exists()


@pytest.mark.parametrize(
    ("payload", "message"),
    [
        ("text_id,translated_text_en\nID,Open, now\n", "line 2 has 3 fields, expected 2"),
        ("text_id, translated_text_en\nID,Open\n", "remove spaces around the delimiter"),
        ("text_id;translated_text_en\nID;Open\n", "check the delimiter"),
        ("text_id,translated_text_de\nID,Open\n", "expected 'text_id,translated_text_en'"),
        ("```csv\n```\n", "no header row"),
    ],
)
def test_csv_reply_errors_name_line_and_cause(source, tmp_path, capsys, payload, message):
    bundle, _manifest = extraction(source, tmp_path)
    returned = tmp_path / "bad.csv"
    returned.write_text(payload, encoding="utf-8")
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 2 and not output.exists()
    assert message in capsys.readouterr().err


def test_csv_reply_tolerates_blank_lines_and_code_fence(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    rows = [
        f'{e["text_id"]},"EN:{e["source_text"].replace(chr(34), chr(34) * 2)}"'
        for e in manifest["entries"]
    ]
    body = "\n\n".join(rows)
    returned = tmp_path / "fenced.csv"
    returned.write_text(f"```csv\ntext_id,translated_text_en\n{body}\n   \n```\n\n", "utf-8")
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 0
    assert load_workbook(output)["Data_en"]["A2"].value == "EN:Открыто"


def sheet_cells(path, name):
    ws = load_workbook(path)[name]
    return [[(c.value, c.data_type, c.number_format) for c in row] for row in ws.iter_rows()]


def test_xml_writer_matches_openpyxl_writer(source, tmp_path, monkeypatch):
    bundle, manifest = extraction(source, tmp_path, "--translate-headings")
    returned = reply(tmp_path, manifest)
    status, fast = application(source, bundle, [returned], tmp_path)
    assert status == 0
    assert json.loads(fast.with_suffix(".report.json").read_text())["writer"] == "xml"
    fast = fast.rename(tmp_path / "fast.xlsx")

    monkeypatch.setattr(tool, "xml_copy", lambda *args: False)
    status, slow = application(source, bundle, [returned], tmp_path, "--overwrite")
    assert status == 0
    assert json.loads(slow.with_suffix(".report.json").read_text())["writer"] == "openpyxl"
    assert load_workbook(fast).sheetnames == load_workbook(slow).sheetnames
    for name in load_workbook(fast).sheetnames:
        assert sheet_cells(fast, name) == sheet_cells(slow, name), name
    assert sheet_cells(fast, "Data") == sheet_cells(source, "Data")


def test_sheet_parts_a_copy_cannot_share_are_dropped_from_it(tmp_path):
    from openpyxl.worksheet.table import Table

    path = tmp_path / "table.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Status"])
    ws.append(["Открыто"])
    ws.add_table(Table(displayName="Statuses", ref="A1:A2"))
    wb.save(path)
    bundle, manifest = extraction(path, tmp_path)
    status, output = application(path, bundle, [reply(tmp_path, manifest)], tmp_path)
    assert status == 0
    result = load_workbook(output)
    assert result["Data_en"]["A2"].value == "EN:Открыто"
    assert list(result["Data"].tables) == ["Statuses"] and not result["Data_en"].tables


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml"


def excel_like_workbook(path, prefix="", cell_attrs='r="{ref}" s="{s}" t="s"'):
    """A minimal workbook written the way Excel and the OpenXML SDK write them: shared
    strings, a bold style, an external hyperlink, optionally prefixed elements."""
    p = f"{prefix}:" if prefix else ""
    ns = f'xmlns{":" + prefix if prefix else ""}="{MAIN_NS}" xmlns:r="{DOC_NS}"'
    strings = ["Status", "Открыто", "Link", "a &amp; &lt;b&gt;"]

    def cell(ref, index, style=0):
        attrs = cell_attrs.format(ref=ref, s=style)
        return f"<{p}c {attrs}><{p}v>{index}</{p}v></{p}c>"

    def rel(rid, kind, target, extra=""):
        return f'<Relationship Id="{rid}" Type="{DOC_NS}/{kind}" Target="{target}"{extra}/>'

    def override(part, kind):
        return f'<Override PartName="/xl/{part}" ContentType="{TYPE}.{kind}+xml"/>'

    files = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            + override("workbook.xml", "sheet.main")
            + override("worksheets/sheet1.xml", "worksheet")
            + override("sharedStrings.xml", "sharedStrings")
            + override("styles.xml", "styles")
            + "</Types>"
        ),
        "_rels/.rels": (
            f'<Relationships xmlns="{PKG_NS}">'
            + rel("rId1", "officeDocument", "xl/workbook.xml")
            + "</Relationships>"
        ),
        "xl/workbook.xml": (
            f"<{p}workbook {ns}><{p}sheets>"
            f'<{p}sheet name="Data" sheetId="1" r:id="rId1"/></{p}sheets></{p}workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{PKG_NS}">'
            + rel("rId1", "worksheet", "worksheets/sheet1.xml")
            + rel("rId2", "sharedStrings", "sharedStrings.xml")
            + rel("rId3", "styles", "styles.xml")
            + "</Relationships>"
        ),
        "xl/sharedStrings.xml": (
            f'<{p}sst {ns} count="4" uniqueCount="4">'
            + "".join(f"<{p}si><{p}t>{text}</{p}t></{p}si>" for text in strings)
            + f"</{p}sst>"
        ),
        "xl/styles.xml": (
            f'<{p}styleSheet {ns}><{p}fonts count="2"><{p}font/><{p}font><{p}b/></{p}font>'
            f'</{p}fonts><{p}fills count="1"><{p}fill><{p}patternFill patternType="none"/>'
            f'</{p}fill></{p}fills><{p}borders count="1"><{p}border/></{p}borders>'
            f'<{p}cellStyleXfs count="1"><{p}xf/></{p}cellStyleXfs><{p}cellXfs count="2">'
            f'<{p}xf fontId="0"/><{p}xf fontId="1" applyFont="1"/></{p}cellXfs>'
            f'<{p}cellStyles count="1"><{p}cellStyle name="Normal" xfId="0" builtinId="0"/>'
            f"</{p}cellStyles></{p}styleSheet>"
        ),
        "xl/worksheets/sheet1.xml": (
            f"<{p}worksheet {ns}><{p}sheetViews>"
            f'<{p}sheetView tabSelected="1" workbookViewId="0"/></{p}sheetViews><{p}sheetData>'
            f'<{p}row r="1">{cell("A1", 0, 1)}{cell("B1", 2)}</{p}row>'
            f'<{p}row r="2">{cell("A2", 1, 1)}{cell("B2", 3)}</{p}row></{p}sheetData>'
            f'<{p}hyperlinks><{p}hyperlink ref="B2" r:id="rId1"/></{p}hyperlinks>'
            f"</{p}worksheet>"
        ),
        "xl/worksheets/_rels/sheet1.xml.rels": (
            f'<Relationships xmlns="{PKG_NS}">'
            + rel("rId1", "hyperlink", "https://example.com/", ' TargetMode="External"')
            + "</Relationships>"
        ),
    }
    with zipfile.ZipFile(path, "w") as archive:
        for name, text in files.items():
            archive.writestr(name, text)
    return path


@pytest.mark.parametrize(
    ("prefix", "cell_attrs", "writer"),
    [
        ("", 'r="{ref}" s="{s}" t="s"', "xml"),
        ("x", 'r="{ref}" s="{s}" t="s"', "xml"),
        ("", 's="{s}" t="s" r="{ref}"', "openpyxl"),  # r not first: rewrite cannot find it
    ],
)
def test_excel_style_workbook_translation(tmp_path, prefix, cell_attrs, writer):
    path = excel_like_workbook(tmp_path / "excel.xlsx", prefix, cell_attrs)
    before = path.read_bytes()
    bundle, manifest = extraction(path, tmp_path, "--translate-headings")
    status, output = application(path, bundle, [reply(tmp_path, manifest)], tmp_path)
    assert status == 0 and path.read_bytes() == before
    assert json.loads(output.with_suffix(".report.json").read_text())["writer"] == writer
    wb = load_workbook(output)
    assert wb.sheetnames == ["Data", "Data_en"]
    copy = wb["Data_en"]
    assert [c.value for c in copy[1]] == ["EN:Status", "EN:Link"]
    assert [c.value for c in copy[2]] == ["EN:Открыто", "EN:a & <b>"]
    assert copy["A2"].font.b and copy["B2"].hyperlink.target == "https://example.com/"
    assert [c.value for c in wb["Data"][2]] == ["Открыто", "a & <b>"]


def test_untouched_source_skips_cell_verification(source, tmp_path):
    bundle, manifest = extraction(source, tmp_path)
    returned = reply(tmp_path, manifest)
    status, output = application(source, bundle, [returned], tmp_path)
    assert status == 0
    assert json.loads(output.with_suffix(".report.json").read_text())["source_check"] == "unchanged"

    # Any byte change, even formatting only, takes the full semantic check.
    wb = load_workbook(source)
    wb["Data"]["A2"].number_format = "@"
    wb.save(source)
    status, output = application(source, bundle, [returned], tmp_path, "--overwrite")
    assert status == 0
    assert json.loads(output.with_suffix(".report.json").read_text())["source_check"] == "cells"
