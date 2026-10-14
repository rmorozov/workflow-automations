"""Outline contracts using real XLSX files with synthetic hierarchical data."""

from datetime import datetime

import pytest
from openpyxl import Workbook

from workflow_automations import xlsx_outline as tool


def workbook(path, rows, sheet="Data", merges=()):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    for area in merges:
        ws.merge_cells(area)
    wb.create_sheet("Other")["A1"] = "ignored"
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def plan(tmp_path):
    return workbook(
        tmp_path / "plan.xlsx",
        [
            ["Area", "Team", "Item", "Owner"],
            ["Platform", "Core", "Scheduler", "Ann"],
            ["Platform", "Core", "Memory", "Bob"],
            ["Platform", "Net", "Routing", None],
            ["Apps", "Mobile", "Login", "Eve"],
            ["Platform", "Core", "IPC", "Ann"],
        ],
    )


def run(capsys, *argv):
    status = tool.main([str(arg) for arg in argv])
    captured = capsys.readouterr()
    return status, captured.out, captured.err


def test_repeated_leading_values_become_parents(plan, capsys):
    status, out, err = run(capsys, "--input", plan, "--levels", "3")
    assert status == 0
    # A non-adjacent repeat keeps row order and starts a new run.
    assert out == (
        "- Platform\n"
        "  - Core\n"
        "    - Scheduler\n"
        "      - Owner: Ann\n"
        "    - Memory\n"
        "      - Owner: Bob\n"
        "  - Net\n"
        "    - Routing\n"
        "- Apps\n"
        "  - Mobile\n"
        "    - Login\n"
        "      - Owner: Eve\n"
        "- Platform\n"
        "  - Core\n"
        "    - IPC\n"
        "      - Owner: Ann\n"
    )
    assert err == "Outlined 5 rows into 12 items\n"


def test_group_merges_nonadjacent_parents(plan, capsys):
    _, out, _ = run(capsys, "--input", plan, "--columns", "Area", "Team", "Item", "--group")
    assert out == (
        "- Platform\n"
        "  - Core\n"
        "    - Scheduler\n"
        "    - Memory\n"
        "    - IPC\n"
        "  - Net\n"
        "    - Routing\n"
        "- Apps\n"
        "  - Mobile\n"
        "    - Login\n"
    )


def test_headings_title_and_labels(plan, capsys):
    _, out, _ = run(
        capsys,
        "--input",
        plan,
        "--column-indices",
        "1",
        "2",
        "3",
        "--group",
        "--heading-levels",
        "2",
        "--title",
        "Plan",
        "--label-levels",
    )
    assert out.startswith(
        "# Plan\n\n## Area: Platform\n\n### Team: Core\n\n"
        "- Item: Scheduler\n- Item: Memory\n- Item: IPC\n\n### Team: Net\n\n- Item: Routing\n\n"
    )
    assert out.endswith("## Area: Apps\n\n### Team: Mobile\n\n- Item: Login\n")


def test_column_order_reorders_hierarchy(plan, capsys):
    _, out, _ = run(capsys, "--input", plan, "--columns", "Owner", "Item", "--group")
    assert out.splitlines()[:4] == ["- Ann", "  - Scheduler", "  - IPC", "- Bob"]
    # A trailing blank shortens the path; an inner blank keeps a placeholder.
    assert "- (blank)\n  - Routing\n" in out


def test_fill_down_stays_within_parent(tmp_path, capsys):
    path = workbook(
        tmp_path / "sparse.xlsx",
        [
            ["Team", "Project", "Task"],
            ["Alpha", "X", "Design"],
            [None, None, "Build"],
            [None, "Y", "Test"],
            ["Beta", None, None],
        ],
    )
    _, out, _ = run(capsys, "--input", path, "--fill-down")
    assert out == "- Alpha\n  - X\n    - Design\n    - Build\n  - Y\n    - Test\n- Beta\n"
    _, out, _ = run(capsys, "--input", path)
    assert "- (blank)\n  - (blank)\n    - Build\n" in out


def test_merged_cells_repeat_top_left_value(tmp_path, capsys):
    path = workbook(
        tmp_path / "merged.xlsx",
        [["Team", "Task"], ["Alpha", "One"], [None, "Two"], ["Beta", "Three"]],
        merges=["A2:A3"],
    )
    _, out, _ = run(capsys, "--input", path)
    assert out == "- Alpha\n  - One\n  - Two\n- Beta\n  - Three\n"


def test_values_are_formatted_and_escaped(tmp_path, capsys):
    path = workbook(
        tmp_path / "values.xlsx",
        [
            ["Kind", "Value"],
            ["Typed", 3.0],
            ["Typed", 2.5],
            ["Typed", True],
            ["Typed", datetime(2026, 10, 8)],
            ["Markup", "# not a heading"],
            ["Markup", "1. not a list"],
            ["Markup", "*bold* [link]"],
            ["Markup", "line one\nline two"],
            ["Markup", "- dash"],
            ["Markup", "---"],
            ["Markup", "-5 #tag"],
        ],
    )
    _, out, _ = run(capsys, "--input", path, "--group")
    assert out == (
        "- Typed\n  - 3\n  - 2.5\n  - TRUE\n  - 2026-10-08\n"
        "- Markup\n  - \\# not a heading\n  - 1\\. not a list\n"
        "  - \\*bold\\* \\[link\\]\n  - line one line two\n"
        "  - \\- dash\n  - \\---\n  - -5 #tag\n"
    )
    _, out, _ = run(capsys, "--input", path, "--group", "--raw", "--indent", "4")
    assert "    - *bold* [link]\n" in out


def test_formulas_use_cached_values(tmp_path, capsys):
    path = workbook(tmp_path / "formula.xlsx", [["Team", "Total"], ["Alpha", "=1+1"]])
    _, out, _ = run(capsys, "--input", path)
    # openpyxl does not calculate formulas, so an uncached result renders as blank.
    assert out == "- Alpha\n"


def test_output_file_requires_overwrite(plan, tmp_path, capsys):
    output = tmp_path / "out" / "plan.md"
    status, out, _ = run(capsys, "--input", plan, "--output", output)
    assert status == 0
    assert out == f"Outlined 5 rows into 16 items: {output}\n"
    assert output.read_text(encoding="utf-8").startswith("- Platform\n  - Core\n")
    status, _, err = run(capsys, "--input", plan, "--output", output)
    assert status == 2
    assert "--overwrite" in err
    assert run(capsys, "--input", plan, "--output", output, "--overwrite")[0] == 0
    before = plan.read_bytes()
    status, _, err = run(capsys, "--input", plan, "--output", plan, "--overwrite")
    assert status == 2
    assert plan.read_bytes() == before


@pytest.mark.parametrize(
    ("options", "message"),
    [
        (["--sheet", "Missing"], "Sheet does not exist"),
        (["--columns", "Nope"], "Unknown columns"),
        (["--column-indices", "9"], "between 1 and 4"),
        (["--column-indices", "1", "1"], "more than once"),
        (["--levels", "5"], "--levels"),
        (["--heading-levels", "5"], "--heading-levels"),
        (["--heading-levels", "4", "--title", "T", "--levels", "4"], None),
        (["--indent", "0"], "--indent"),
    ],
)
def test_validation_errors(plan, capsys, options, message):
    status, _, err = run(capsys, "--input", plan, *options)
    if message is None:
        assert status == 0
    else:
        assert status == 2
        assert message in err


def test_heading_limit_counts_title(tmp_path, capsys):
    path = workbook(tmp_path / "deep.xlsx", [list("ABCDEF"), list("abcdef")])
    status, _, err = run(capsys, "--input", path, "--heading-levels", "6", "--title", "T")
    assert status == 2
    assert "six heading levels" in err
    assert run(capsys, "--input", path, "--heading-levels", "6")[0] == 0


def test_io_errors(tmp_path, capsys):
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a zip")
    assert run(capsys, "--input", bad)[0] == 3
    assert run(capsys, "--input", tmp_path / "absent.xlsx")[0] == 3
    assert run(capsys, "--input", tmp_path / "data.csv")[0] == 2
    empty = workbook(tmp_path / "empty.xlsx", [])
    status, _, err = run(capsys, "--input", empty)
    assert status == 2
    assert "empty" in err


def test_shortened_row_ends_the_child_run(tmp_path, capsys):
    path = workbook(
        tmp_path / "short.xlsx", [["Area", "Item"], ["A", "X"], ["A", None], ["A", "X"]]
    )
    _, out, _ = run(capsys, "--input", path)
    assert out == "- A\n  - X\n  - X\n"
    _, out, _ = run(capsys, "--input", path, "--group")
    assert out == "- A\n  - X\n"


def test_blank_cell_differs_from_blank_label_text(tmp_path, capsys):
    path = workbook(
        tmp_path / "blank.xlsx", [["Area", "Item"], [None, "X"], ["(blank)", "Y"], [None, "Z"]]
    )
    _, out, _ = run(capsys, "--input", path, "--group")
    assert out == "- (blank)\n  - X\n  - Z\n- \\(blank)\n  - Y\n"
    _, out, _ = run(capsys, "--input", path)
    assert out == "- (blank)\n  - X\n- \\(blank)\n  - Y\n- (blank)\n  - Z\n"


def test_details_keep_row_order_around_children(tmp_path, capsys):
    path = workbook(
        tmp_path / "order.xlsx",
        [
            ["Area", "Item", "Owner"],
            ["A", None, "before"],
            ["A", "X", "first"],
            ["A", None, "between"],
            ["A", "Y", "second"],
            ["A", None, "after"],
        ],
    )
    _, out, _ = run(capsys, "--input", path, "--levels", "2")
    assert out == (
        "- A\n"
        "  - Owner: before\n"
        "  - X\n"
        "    - Owner: first\n"
        "  - Owner: between\n"
        "  - Y\n"
        "    - Owner: second\n"
        "  - Owner: after\n"
    )


def test_formatting_far_from_data_is_not_materialized(tmp_path, capsys, monkeypatch):
    path = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Team", "Task"])
    ws.append(["Alpha", "One"])
    ws.cell(5000, 5000).number_format = "0.00"
    wb.save(path)
    wb.close()
    opened = []
    original = tool.load_workbook

    def spy(*args, **kwargs):
        opened.append(original(*args, **kwargs))
        return opened[-1]

    monkeypatch.setattr(tool, "load_workbook", spy)
    _, out, _ = run(capsys, "--input", path)
    assert out == "- Alpha\n  - One\n"
    assert len(opened[0].active._cells) < 10


@pytest.fixture
def two_sheets(tmp_path):
    path = tmp_path / "teams.xlsx"
    wb = Workbook()
    wb.active.title = "Plan"
    for row in [["Area", "Item", "Owner"], ["Core", "Scheduler", "Ann"], ["Core", "IPC", None]]:
        wb.active.append(row)
    wb.create_sheet("Empty")
    risks = wb.create_sheet("Risks #")
    for row in [["Risk", "Mitigation"], ["Outage", "Failover"], ["Churn", None]]:
        risks.append(row)
    wb.save(path)
    wb.close()
    return path


def test_all_sheets_outline_one_heading_per_sheet(two_sheets, capsys):
    status, out, err = run(capsys, "--input", two_sheets, "--all-sheets", "--levels", "2")
    assert status == 0
    # The empty sheet is skipped; a trailing "#" is escaped as for any heading.
    assert out == (
        "# Plan\n"
        "\n"
        "- Core\n"
        "  - Scheduler\n"
        "    - Owner: Ann\n"
        "  - IPC\n"
        "\n"
        "# Risks \\#\n"
        "\n"
        "- Outage\n"
        "  - Failover\n"
        "- Churn\n"
    )
    assert err == "Outlined 4 rows into 6 items\n"


def test_all_sheets_sections_unfold_back_separately(two_sheets, tmp_path, capsys):
    from openpyxl import load_workbook

    from workflow_automations import xlsx_unfold

    md = tmp_path / "teams.md"
    argv = ["--input", two_sheets, "--all-sheets", "--heading-levels", "1", "--output", md]
    assert run(capsys, *argv)[0] == 0
    text = md.read_text()
    starts = [i for i, line in enumerate(text.splitlines()) if line.startswith("# ")]
    assert len(starts) == 2
    lines = text.splitlines()
    sections = [lines[starts[0] : starts[1]], lines[starts[1] :]]
    # Each section is that sheet's own outline titled with its name.
    single = run(
        capsys, "--input", two_sheets, "--sheet", "Plan", "--title", "Plan", "--heading-levels", "1"
    )[1]
    assert "\n".join(sections[0]).rstrip("\n") + "\n" == single
    expected = {
        "Plan": [("Area", "Item", "Owner"), ("Core", "Scheduler", "Ann"), ("Core", "IPC", None)],
        "Risks #": [("Risk", "Mitigation"), ("Outage", "Failover"), ("Churn", None)],
    }
    for index, (sheet, rows) in enumerate(expected.items()):
        part = tmp_path / f"part{index}.md"
        part.write_text("\n".join(sections[index]) + "\n")
        result = tmp_path / f"part{index}.xlsx"
        status = xlsx_unfold.main(
            ["--input", str(part), "--output", str(result), "--title", "--columns", *rows[0]]
        )
        capsys.readouterr()
        assert status == 0
        wb = load_workbook(result)
        assert list(wb["Outline"].iter_rows(values_only=True)) == rows, sheet
        wb.close()


@pytest.mark.parametrize(
    ("options", "message"),
    [
        (["--title", "T"], "cannot be combined with --title"),
        (["--front-matter"], "omit it with --all-sheets"),
        (["--columns", "Area", "Item"], "Sheet Risks #: Unknown columns: Area, Item"),
        (["--levels", "3"], "Sheet Risks #: --levels must be between 1 and 2"),
        (["--heading-levels", "6"], "Sheet Plan: --heading-levels must be between 0 and 3"),
    ],
)
def test_all_sheets_validation_errors(two_sheets, capsys, options, message):
    status, out, err = run(capsys, "--input", two_sheets, "--all-sheets", *options)
    assert status == 2
    assert out == ""
    assert message in err


def test_all_sheets_rejects_sheet_and_empty_workbooks(two_sheets, tmp_path, capsys):
    with pytest.raises(SystemExit):
        tool.main(["--input", str(two_sheets), "--all-sheets", "--sheet", "Plan"])
    assert "not allowed with argument" in capsys.readouterr().err
    empty = tmp_path / "empty.xlsx"
    Workbook().save(empty)
    status, _, err = run(capsys, "--input", empty, "--all-sheets")
    assert status == 2
    assert "Every sheet is empty" in err


def test_options_are_checked_before_the_workbook_is_read(tmp_path, capsys):
    status, _, err = run(capsys, "--input", tmp_path / "missing.xlsx", "--indent", "0")
    assert status == 2
    assert "--indent must be a positive integer" in err


def test_all_sheets_heading_limit_counts_the_sheet_heading(tmp_path, capsys):
    path = tmp_path / "deep.xlsx"
    wb = Workbook()
    wb.active.append([f"L{level}" for level in range(6)])
    wb.active.append([f"v{level}" for level in range(6)])
    wb.save(path)
    status, _, err = run(capsys, "--input", path, "--all-sheets", "--heading-levels", "6")
    assert status == 2
    assert "including the sheet heading" in err
