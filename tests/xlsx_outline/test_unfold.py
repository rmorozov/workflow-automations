"""Reverse contracts: outlines (generated or hand-edited) unfold into XLSX tables."""

from datetime import date, datetime, time

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.datetime import MAC_EPOCH

from workflow_automations import xlsx_outline, xlsx_unfold

ROWS = [
    ["Area", "Team", "Item", "Owner", "Due"],
    ["Platform", "Core", "Scheduler", "Ann", "Q1"],
    ["Platform", "Core", "Memory", "Bob", None],
    ["Platform", "Net", "Routing", None, None],
    ["Platform", None, None, "Lead; Owner: Cy", None],
    ["Apps", "Mobile", "*Login* #1", "Eve", "Q2"],
    ["Apps", "(blank)", "Logout", None, None],
    ["Apps", None, "Theme", None, "Q3"],
]


def source(tmp_path, rows=ROWS):
    path = tmp_path / "plan.xlsx"
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    return path


def table(path, sheet="Outline"):
    wb = load_workbook(path)
    rows = [list(row) for row in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return rows


def outline(tmp_path, capsys, *options, rows=ROWS):
    output = tmp_path / "plan.md"
    argv = ["--input", str(source(tmp_path, rows)), "--output", str(output), "--front-matter"]
    assert xlsx_outline.main([*argv, *options]) == 0
    capsys.readouterr()
    return output


def unfold(capsys, *argv):
    status = xlsx_unfold.main([str(arg) for arg in argv])
    captured = capsys.readouterr()
    return status, captured.out, captured.err


def test_round_trip_restores_rows(tmp_path, capsys):
    md = outline(tmp_path, capsys, "--levels", "3")
    assert md.read_text().startswith('---\nxlsx-outline:\n  version: 1\n  sheet: "Sheet"\n')
    result = tmp_path / "back.xlsx"
    status, out, _ = unfold(capsys, "--input", md, "--output", result)
    assert status == 0
    assert out == f"Unfolded 7 rows into 5 columns: {result}\n"
    assert table(result) == ROWS


def test_round_trip_with_headings_title_and_labels(tmp_path, capsys):
    # A detail line after a child section would join that section under headings.
    rows = [row for row in ROWS if row[1] is not None or row[2] is not None]
    md = outline(
        tmp_path,
        capsys,
        "--levels",
        "3",
        "--heading-levels",
        "2",
        "--title",
        "Plan",
        "--label-levels",
        "--indent",
        "4",
        rows=rows,
    )
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result)[0] == 0
    assert table(result) == rows


def test_edits_flow_back_as_row_changes(tmp_path, capsys):
    md = outline(tmp_path, capsys, "--levels", "3")
    edited = (
        md.read_text()
        # Renaming a parent renames it for every row beneath it.
        .replace("  - Core\n", "  - Kernel\n")
        # Moving a subtree changes its parent columns.
        .replace("  - Net\n    - Routing\n", "")
        .replace("- Apps\n", "- Apps\n  - Net\n    - Routing\n      - Owner: Dee\n")
        # Deleting an item deletes its row; adding one adds a row.
        .replace("    - Memory\n      - Owner: Bob\n", "    - Timers\n")
    )
    md.write_text(edited)
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result)[0] == 0
    assert table(result) == [
        ROWS[0],
        ["Platform", "Kernel", "Scheduler", "Ann", "Q1"],
        ["Platform", "Kernel", "Timers", None, None],
        ["Platform", None, None, "Lead; Owner: Cy", None],
        ["Apps", "Net", "Routing", "Dee", None],
        *ROWS[5:],
    ]


def test_mind_map_formatting_is_accepted(tmp_path, capsys):
    md = tmp_path / "map.md"
    md.write_text(
        "# Plan\n\n## Platform\n\n"
        "* Core\n\t* Scheduler\n\t\t+ Owner: Ann\n\n"
        "1. Net\n    1. Routing\n"
        "## Apps\n#### Mobile ##\n- Login\n"
    )
    result = tmp_path / "back.xlsx"
    status, _, _ = unfold(
        capsys,
        "--input",
        md,
        "--output",
        result,
        "--title",
        "--columns",
        "Area",
        "Team",
        "Item",
        "--details",
        "Owner",
    )
    assert status == 0
    assert table(result) == [
        ["Area", "Team", "Item", "Owner"],
        ["Platform", "Core", "Scheduler", "Ann"],
        ["Platform", "Net", "Routing", None],
        ["Apps", "Mobile", "Login", None],
    ]


def test_without_front_matter_levels_get_generic_names(tmp_path, capsys):
    md = tmp_path / "plain.md"
    md.write_text("- A\n  - X\n  - Y\n- B\n")
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result, "--sheet", "Edited")[0] == 0
    assert table(result, "Edited") == [
        ["Level 1", "Level 2"],
        ["A", "X"],
        ["A", "Y"],
        ["B", None],
    ]


def test_formula_like_text_stays_text(tmp_path, capsys):
    md = tmp_path / "formula.md"
    md.write_text("- =1+1\n")
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result)[0] == 0
    cell = load_workbook(result).active["A2"]
    assert (cell.value, cell.data_type) == ("=1+1", "s")


def test_validation_and_output_protection(tmp_path, capsys):
    md = tmp_path / "plan.md"
    result = tmp_path / "back.xlsx"
    md.write_text("- A\n  - X\n")

    def error(*options):
        status, _, err = unfold(capsys, "--input", md, "--output", result, *options)
        assert status == 2
        return err

    assert "nested deeper than the 1 level" in error("--columns", "Area")
    assert "unique" in error("--columns", "A", "B", "--details", "A")
    assert "Invalid sheet name" in error("--sheet", "a/b")
    md.write_text("- A\nsome note\n")
    assert "Line 2 is neither a heading nor a list item" in error()
    md.write_text("---\nxlsx-outline: [1\n---\n- A\n")
    assert "not valid YAML" in error()
    md.write_text("")
    assert "no items" in error()
    md.write_text("- A\n")
    assert unfold(capsys, "--input", md, "--output", result)[0] == 0
    assert "--overwrite" in error()
    assert unfold(capsys, "--input", md, "--output", result, "--overwrite")[0] == 0
    status, _, err = unfold(capsys, "--input", md, "--output", tmp_path / "out.csv")
    assert status == 2 and ".xlsx" in err
    assert (
        unfold(capsys, "--input", tmp_path / "missing.md", "--output", result, "--overwrite")[0]
        == 3
    )


MERGE_ROWS = [
    ["Area", "Team", "Item", "Notes", "Size"],
    ["Platform", "Core", "Scheduler", "keep me", 3],
    ["Platform", "Core", "Memory", None, 5],
    ["Platform", "Net", "Routing", "=1+1", None],
    [None, None, None, "orphan note", None],
    ["Apps", "Mobile", "Login", None, 8],
]


@pytest.fixture
def original(tmp_path):
    path = tmp_path / "original.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in MERGE_ROWS:
        ws.append(row)
    ws["C2"].font = Font(bold=True)
    wb.create_sheet("Other")["A1"] = "untouched"
    wb.save(path)
    return path


def tagged(original, tmp_path, capsys, *options):
    md = tmp_path / "plan.md"
    argv = ["--input", original, "--output", md, "--columns", "Area", "Team", "Item", "Size"]
    argv += ["--levels", "3", "--overwrite", *options]
    assert xlsx_outline.main([str(a) for a in argv]) == 0
    capsys.readouterr()
    return md


def merged(original, md, tmp_path, capsys, *options):
    result = tmp_path / "merged.xlsx"
    status, out, err = unfold(
        capsys, "--input", md, "--into", original, "--output", result, *options
    )
    assert status == 0, err
    return load_workbook(result), out


def values(wb, sheet="Data"):
    return [list(row) for row in wb[sheet].iter_rows(values_only=True)]


def test_merge_without_edits_keeps_the_sheet(original, tmp_path, capsys):
    md = tagged(original, tmp_path, capsys, "--row-ids", "--front-matter")
    assert "    - Scheduler\n      - Size: 3 <!-- rows: 2 -->\n" in md.read_text()
    before = original.read_bytes()
    wb, out = merged(original, md, tmp_path, capsys)
    assert values(wb) == [*MERGE_ROWS[:4], MERGE_ROWS[5], MERGE_ROWS[4]]
    assert "0 updated, 0 added, 0 deleted, 1 without outline values kept at the end" in out
    assert wb["Data"]["C2"].font.bold
    assert wb["Data"]["D4"].data_type == "f"
    assert wb["Other"]["A1"].value == "untouched"
    assert original.read_bytes() == before


def test_merge_applies_mind_map_edits_to_rows(original, tmp_path, capsys):
    md = tagged(original, tmp_path, capsys, "--row-ids", "--front-matter")
    text = md.read_text()
    edited = (
        text.replace("  - Core\n", "  - Kernel\n")
        .replace("    - Memory\n      - Size: 5 <!-- rows: 3 -->\n", "")
        .replace("  - Net\n    - Routing <!-- rows: 4 -->\n", "")
        .replace("- Apps\n", "- Apps\n  - Net\n    - Routing <!-- rows: 4 -->\n    - Caching\n")
        .replace("Size: 8", "Size: 13")
    )
    md.write_text(edited)
    wb, out = merged(original, md, tmp_path, capsys)
    assert values(wb) == [
        MERGE_ROWS[0],
        ["Platform", "Kernel", "Scheduler", "keep me", 3],
        ["Apps", "Net", "Routing", "=1+1", None],
        ["Apps", "Net", "Caching", None, None],
        ["Apps", "Mobile", "Login", None, "13"],
        [None, None, None, "orphan note", None],
    ]
    assert "Merged 4 rows into Data: 3 updated, 1 added, 1 deleted" in out
    # Formatting and unselected columns travel with their row.
    assert wb["Data"]["C2"].font.bold


def test_merge_without_front_matter_uses_options(original, tmp_path, capsys):
    md = tagged(original, tmp_path, capsys, "--row-ids")
    md.write_text(md.read_text().replace("Scheduler", "Dispatcher"))
    wb, _ = merged(
        original,
        md,
        tmp_path,
        capsys,
        "--columns",
        "Area",
        "Team",
        "Item",
        "--details",
        "Size",
        "--sheet",
        "Data",
    )
    assert values(wb)[1] == ["Platform", "Core", "Dispatcher", "keep me", 3]


def test_merge_keeps_filled_down_parents_blank(tmp_path, capsys):
    path = tmp_path / "sparse.xlsx"
    wb = Workbook()
    for row in [["Team", "Task"], ["Alpha", "One"], [None, "Two"], ["Beta", "Three"]]:
        wb.active.append(row)
    wb.save(path)
    md = tmp_path / "sparse.md"
    argv = ["--input", path, "--output", md, "--fill-down", "--row-ids", "--front-matter"]
    assert xlsx_outline.main([str(a) for a in argv]) == 0
    md.write_text(md.read_text().replace("Three", "Four"))
    result = tmp_path / "merged.xlsx"
    assert unfold(capsys, "--input", md, "--into", path, "--output", result)[0] == 0
    assert values(load_workbook(result), "Sheet") == [
        ["Team", "Task"],
        ["Alpha", "One"],
        [None, "Two"],
        ["Beta", "Four"],
    ]


def test_row_tags_restore_merged_duplicates_in_a_new_sheet(tmp_path, capsys):
    rows = [["Team", "Task"], ["Alpha", "One"], ["Alpha", "One"], ["Alpha", None]]
    md = outline(tmp_path, capsys, "--row-ids", rows=rows)
    assert "- Alpha <!-- rows: 4 -->\n  - One <!-- rows: 2 3 -->\n" in md.read_text()
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result)[0] == 0
    assert table(result) == [rows[0], rows[3], rows[1], rows[2]]


def test_merge_validation(original, tmp_path, capsys):
    result = tmp_path / "merged.xlsx"

    def error(md, *options):
        status, _, err = unfold(
            capsys, "--input", md, "--into", original, "--output", result, *options
        )
        assert status == 2
        return err

    untagged = tagged(original, tmp_path, capsys, "--front-matter")
    assert "no row tags" in error(untagged)
    md = tagged(original, tmp_path, capsys, "--row-ids")
    assert "level column names" in error(md)
    assert "exactly once" in error(md, "--columns", "Area", "Team", "Nope", "--details", "Size")
    md.write_text(
        md.read_text().replace("rows: 2 ", "rows: 99 ").replace("rows: 2 -->", "rows: 99 -->")
    )
    assert "outside the data rows 2-6" in error(
        md, "--columns", "Area", "Team", "Item", "--details", "Size"
    )
    md = tagged(original, tmp_path, capsys, "--row-ids", "--front-matter")
    wb = load_workbook(original)
    wb["Data"]["A2"] = "Changed"
    wb.save(original)
    assert "changed since the outline was written" in error(md)
    status, _, err = unfold(capsys, "--input", md, "--into", original, "--output", original)
    assert status == 2 and "new file" in err


def round_trip(tmp_path, capsys, rows, *options):
    """Outline rows with tags and front matter, then unfold and merge without edits."""
    path = tmp_path / "round.xlsx"
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    md = tmp_path / "round.md"
    argv = ["--input", path, "--output", md, "--row-ids", "--front-matter", "--overwrite"]
    assert xlsx_outline.main([str(a) for a in [*argv, *options]]) == 0
    new, merged_path = tmp_path / "new.xlsx", tmp_path / "merged.xlsx"
    for extra, output in [([], new), (["--into", path], merged_path)]:
        status, _, err = unfold(capsys, "--input", md, "--output", output, "--overwrite", *extra)
        assert status == 0, err
    capsys.readouterr()
    return md.read_text(), table(new), values(load_workbook(merged_path), "Sheet")


@pytest.mark.parametrize("position", ["middle", "end"])
def test_merge_keeps_rows_with_uncached_formulas(tmp_path, capsys, position):
    rows = [["Area", "Task"], ["A", "One"], [None, "=1+1"], ["B", "Two"]]
    if position == "end":
        rows = [rows[0], rows[1], rows[3], rows[2]]
    path = tmp_path / "formulas.xlsx"
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    md = tmp_path / "formulas.md"
    argv = ["--input", path, "--output", md, "--row-ids", "--front-matter"]
    assert xlsx_outline.main([str(a) for a in argv]) == 0
    md.write_text(md.read_text() + "- C\n")
    result = tmp_path / "merged.xlsx"
    status, out, _ = unfold(capsys, "--input", md, "--into", path, "--output", result)
    assert status == 0
    assert "1 added, 0 deleted, 1 without outline values kept at the end" in out
    merged_rows = values(load_workbook(result), "Sheet")
    assert merged_rows[-1] == [None, "=1+1"]
    assert sorted(map(str, merged_rows[1:-1])) == sorted(
        map(str, [["A", "One"], ["B", "Two"], ["C", None]])
    )


def test_detail_values_containing_later_labels_round_trip(tmp_path, capsys):
    rows = [
        ["Area", "Owner", "Due"],
        ["A", "Ask; Due: tomorrow", None],
        ["B", "Ask; Due: soon", "today"],
        ["C", "x\\", "y"],
    ]
    text, new, merged = round_trip(tmp_path, capsys, rows, "--levels", "1")
    assert "- Owner: Ask\\; Due: tomorrow <!-- rows: 2 -->" in text
    assert new == merged == rows


def test_hierarchy_items_that_look_like_details_round_trip(tmp_path, capsys):
    rows = [["Area", "Item", "Owner"], ["A", "Owner: Ann", None], ["A", None, "Bob"]]
    text, new, merged = round_trip(tmp_path, capsys, rows, "--levels", "2")
    assert "  - Owner\\: Ann <!-- rows: 2 -->\n  - Owner: Bob <!-- rows: 3 -->\n" in text
    assert new == merged == rows


def test_trailing_hashes_in_headings_round_trip(tmp_path, capsys):
    rows = [["Area #", "Task"], ["Feature #", "One"], ["C ##", "Two"], ["#", "Three"]]
    expected = {
        (): "# Feature \\#",
        ("--label-levels",): "# Area #: Feature \\#",
        ("--title", "Plan #"): "# Plan \\#\n\n## Feature \\#",
    }
    for options, heading in expected.items():
        text, new, merged = round_trip(tmp_path, capsys, rows, "--heading-levels", "1", *options)
        assert heading in text
        assert new == merged == rows


def test_blank_label_literal_text_round_trips(tmp_path, capsys):
    rows = [["Area", "Item"], [None, "X"], ["[empty]", "Y"], ["(empty)", "Z"]]
    text, new, merged = round_trip(tmp_path, capsys, rows, "--blank-label", "(empty)")
    assert "- (empty)\n  - X" in text and "- \\(empty)\n  - Z" in text
    assert new == merged == rows


def test_blank_label_must_be_reversible(tmp_path, capsys):
    path = source(tmp_path)
    for label in ["EMPTY", "[empty]", " (x)", ""]:
        status = xlsx_outline.main(["--input", str(path), "--blank-label", label])
        assert status == 2
        assert "--blank-label must start with ASCII punctuation" in capsys.readouterr().err


def test_detail_names_containing_semicolons_round_trip(tmp_path, capsys):
    rows = [
        ["Area", "Owner", "Due; phase"],
        ["A", "Ann", "Soon"],
        ["B", None, "Later"],
        ["C", "Cy", None],
    ]
    text, new, merged = round_trip(tmp_path, capsys, rows, "--levels", "1")
    assert "- Owner: Ann; Due\; phase: Soon <!-- rows: 2 -->" in text
    assert "- Due\; phase: Later <!-- rows: 3 -->" in text
    assert new == merged == rows


TYPED_ROWS = [
    ["Kind", "Value"],
    ["whole", 42],
    ["negative", -7],
    ["fraction", 0.1],
    ["date", datetime(2024, 1, 2)],
    ["date-time", datetime(2024, 1, 2, 3, 4, 5, 123000)],
    ["time", time(3, 4, 5)],
    ["code", "007"],
    ["id", "1234567890123456"],
    ["flag", True],
]


def test_retype_restores_numbers_and_dates(tmp_path, capsys):
    md = outline(tmp_path, capsys, "--levels", "1", rows=TYPED_ROWS)
    plain, typed = tmp_path / "plain.xlsx", tmp_path / "typed.xlsx"
    assert unfold(capsys, "--input", md, "--output", plain)[0] == 0
    assert unfold(capsys, "--input", md, "--output", typed, "--retype")[0] == 0
    # Without the option every value stays text, as before.
    assert table(plain) == [[k, xlsx_outline.text(v)] for k, v in TYPED_ROWS]
    # Leading zeros, IDs beyond Excel's precision and booleans stay text.
    assert table(typed) == [
        *TYPED_ROWS[:7],
        *[[k, xlsx_outline.text(v)] for k, v in TYPED_ROWS[7:]],
    ]
    sheet = load_workbook(typed).active
    assert sheet["B5"].number_format == "yyyy-mm-dd"
    assert [sheet[f"A{row}"].data_type for row in range(1, 11)] == ["s"] * 10


def test_retype_keeps_other_spellings_as_text(tmp_path, capsys):
    spellings = [
        "1.50", "1e3", "-0", "+5", "1_000", "inf", "=1+1", "12:30", "2024-1-2",
        "2024-02-30", "2024-01-02T03:04:05", "2024-01-02 00:00:00", "1899-12-31",
        "2024-01-02 03:04:05+00:00", "2024-01-02 03:04:05.123456", "TRUE",
        "0.30000000000000004", "1234567890123456.5", "5e-324",
    ]  # fmt: skip
    md = tmp_path / "hand.md"
    md.write_text("".join(f"- {value}\n" for value in ["2024", *spellings]))
    result = tmp_path / "back.xlsx"
    assert unfold(capsys, "--input", md, "--output", result, "--retype")[0] == 0
    sheet = load_workbook(result).active
    assert (sheet["A1"].value, sheet["A1"].data_type) == ("Level 1", "s")
    assert sheet["A2"].value == 2024
    assert [(cell.value, cell.data_type) for (cell,) in sheet.iter_rows(min_row=3)] == [
        (value, "s") for value in spellings
    ]


def test_merge_retypes_only_edited_values(tmp_path, capsys):
    path = tmp_path / "original.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Item", "Size", "Due"], ["A", 3, date(2024, 1, 2)], ["B", 5, "soon"]]:
        ws.append(row)
    ws["B3"].number_format = "0.00"
    wb.save(path)
    md = tmp_path / "plan.md"
    argv = ["--input", path, "--output", md, "--levels", "1", "--row-ids", "--front-matter"]
    assert xlsx_outline.main([str(a) for a in argv]) == 0
    capsys.readouterr()
    edited = (
        md.read_text().replace("Due: 2024-01-02", "Due: 12").replace("Due: soon", "Due: 2025-03-04")
    )
    md.write_text(edited.replace("Size: 5", "Size: 2.5"))
    plain, _ = merged(path, md, tmp_path, capsys)
    assert values(plain) == [["Item", "Size", "Due"], ["A", 3, "12"], ["B", "2.5", "2025-03-04"]]
    typed, out = merged(path, md, tmp_path, capsys, "--overwrite", "--retype")
    assert values(typed) == [
        ["Item", "Size", "Due"],
        ["A", 3, 12],
        ["B", 2.5, datetime(2025, 3, 4)],
    ]
    assert "2 updated" in out
    sheet = typed["Data"]
    # A number written over a date no longer displays as a date; other formats stay.
    assert (sheet["C2"].number_format, sheet["B3"].number_format) == ("General", "0.00")
    assert sheet["C3"].number_format == "yyyy-mm-dd"


def merge_one_row(tmp_path, capsys, cells, edits, epoch=None):
    """Outline a one-row sheet with row tags, apply text edits and merge with --retype."""
    path = tmp_path / "formats.xlsx"
    wb = Workbook()
    if epoch is not None:
        wb.epoch = epoch
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", *cells])
    ws.append(["A", *(value for value, _ in cells.values())])
    for column, (_, number_format) in enumerate(cells.values(), 2):
        ws.cell(2, column).number_format = number_format
    wb.save(path)
    md = tmp_path / "plan.md"
    argv = ["--input", path, "--output", md, "--levels", "1", "--row-ids", "--front-matter"]
    assert xlsx_outline.main([str(a) for a in argv]) == 0
    capsys.readouterr()
    content = md.read_text()
    for old, new in edits:
        assert old in content
        content = content.replace(old, new)
    md.write_text(content)
    return merged(path, md, tmp_path, capsys, "--retype")[0]["Data"]


def test_merge_retype_fits_the_cell_format(tmp_path, capsys):
    cells = {
        "Day": (date(2024, 1, 2), "yyyy-mm-dd"),
        "Clock": (time(1, 2, 3), "h:mm:ss"),
        "Stamp": (datetime(2024, 1, 2, 3, 4), "yyyy-mm-dd h:mm"),
        "Moment": (date(2024, 1, 2), "d mmm yyyy"),
        "Code": ("1042", "@"),
    }
    sheet = merge_one_row(
        tmp_path,
        capsys,
        cells,
        [
            ("Day: 2024-01-02", "Day: 03:04:05"),
            ("Clock: 01:02:03", "Clock: 2024-05-06"),
            ("Stamp: 2024-01-02 03:04:00", "Stamp: 2024-01-02 15:45:00"),
            ("Moment: 2024-01-02", "Moment: 2024-01-03"),
            ("Code: 1042", "Code: 1043"),
        ],
    )
    assert [cell.value for cell in sheet[2]][1:] == [
        time(3, 4, 5),
        datetime(2024, 5, 6),
        datetime(2024, 1, 2, 15, 45),
        datetime(2024, 1, 3),
        "1043",
    ]
    # A format showing other parts than the value is replaced; a matching one stays.
    assert [cell.number_format for cell in sheet[2]][1:] == [
        "h:mm:ss",
        "yyyy-mm-dd",
        "yyyy-mm-dd h:mm",
        "d mmm yyyy",
        "@",
    ]


def test_merge_retype_respects_the_1904_epoch(tmp_path, capsys):
    cells = {"Early": ("x", "General"), "Epoch": ("y", "General"), "Late": ("z", "General")}
    sheet = merge_one_row(
        tmp_path,
        capsys,
        cells,
        [("Early: x", "Early: 1901-06-01"), ("Epoch: y", "Epoch: 1904-01-01"),
         ("Late: z", "Late: 1904-01-02")],
        epoch=MAC_EPOCH,
    )  # fmt: skip
    assert [cell.value for cell in sheet[2]][1:] == [
        "1901-06-01",
        "1904-01-01",
        datetime(1904, 1, 2),
    ]
