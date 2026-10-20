"""Behavior of team-roster: title classes, hierarchical AreaPath mapping and headcount."""

import io

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from workflow_automations import team_roster

STAFF = {
    "email": ["Ann@x.example", "bob@x.example", None, "dan@x.example", "eve@x.example", "fay@x"],
    "corporatetitle": [
        "Senior Developer",
        "Lead Developer",
        "Test Engineer",
        None,
        "Chief Architect",
        "Intern",
    ],
    "department_name": ["Core", "Core", "QA", "Net", "Core", "Nowhere"],
    "fullpath": ["Org/Eng/Core", "Org/Eng/Core", "Org/Eng/QA", "Org/Net", "Org/Eng/Core", "Org/X"],
    "approved": [True, True, True, True, False, True],
}


def workbook(path, sheets):
    with pd.ExcelWriter(path) as writer:
        for name, data in sheets.items():
            pd.DataFrame(data).to_excel(writer, sheet_name=name, index=False)
    return path


def sheets(path):
    return pd.read_excel(path, sheet_name=None)


@pytest.fixture
def inputs(tmp_path):
    staff = workbook(
        tmp_path / "staff.xlsx",
        {
            "staff": STAFF,
            "person_to_area": {"email": ["ann@x.example "], "area_path": ["P\\Special"]},
        },
    )
    mapping = workbook(
        tmp_path / "mapping.xlsx",
        {
            "department_to_area": {
                "department_name": ["core", "QA"],
                "area_path": ["P\\Core", "P\\QA"],
            },
            "person_to_area": {"email": ["bob@x.example"], "area_path": ["P\\Ignored"]},
        },
    )
    return staff, mapping


def run(*argv):
    return team_roster.main([str(arg) for arg in argv])


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Lead Developer", "lead developer"),
        ("Senior Developer", "developer"),
        ("Junior Software Developer", "developer"),
        ("Principal Research Developer", "research developer"),
        ("Senior Test Engineer", "testing engineer"),
        ("Testing Engineer II", "testing engineer"),
        ("Software Engineer in Test", "sdet"),
        ("Security Analyst", "security analyst"),
        ("System Analyst", "analyst"),
        ("Head of Development", "head"),
        ("Team Lead", "team lead"),
        ("Solution Architect", "architect"),
        ("Development Manager", "manager"),
        ("  LEAD DEVELOPER ", "lead developer"),
        ("Intern", "Intern"),
        (None, "Unknown"),
        (float("nan"), "Unknown"),
        (42, "Unknown"),
    ],
)
def test_classify_title_first_matching_rule_wins(title, expected):
    assert team_roster.classify_title(title) == expected


def test_area_path_prefers_person_then_department_then_fullpath(inputs, tmp_path):
    staff, mapping = inputs
    workbook(
        staff,
        {
            "staff": STAFF,
            "person_to_area": {"email": ["ann@x.example "], "area_path": ["P\\Special"]},
            "fullpath_to_area": {"full_path": ["org/net"], "area_path": ["P\\Net"]},
        },
    )
    output = tmp_path / "out.xlsx"
    assert run(staff, output, "--mapping", mapping) == 0
    main = sheets(output)["staff"]
    assert main["area_path"].tolist()[:4] == ["P\\Special", "P\\Core", "P\\QA", "P\\Net"]
    assert main["area_path"].isna().tolist()[4:] == [False, True]
    assert main["classified_title"].tolist() == [
        "developer",
        "lead developer",
        "testing engineer",
        "Unknown",
        "architect",
        "Intern",
    ]


def test_input_mapping_sheet_wins_over_mapping_file(inputs, tmp_path):
    staff, mapping = inputs
    output = tmp_path / "out.xlsx"
    assert run(staff, output, "-m", mapping) == 0
    book = sheets(output)
    # bob is mapped by department: the mapping file's person_to_area is not used.
    assert book["staff"].loc[1, "area_path"] == "P\\Core"
    assert book["person_to_area"]["email"].tolist() == ["ann@x.example "]
    assert list(book) == [
        "staff",
        "dev_per_area",
        "person_to_area",
        "department_to_area",
        "fullpath_to_area",
    ]
    assert list(book["fullpath_to_area"].columns) == ["full_path", "area_path"]


def test_headcount_counts_approved_people_with_area_and_email(inputs, tmp_path):
    staff, mapping = inputs
    output = tmp_path / "out.xlsx"
    assert run(staff, output, "-m", mapping) == 0
    pivot = sheets(output)["dev_per_area"].set_index("area_path")
    # eve is not approved; dan and fay have no area; the QA tester has no email.
    assert pivot.to_dict("index") == {
        "P\\Core": {"developer": 0, "lead developer": 1, "testing engineer": 0},
        "P\\QA": {"developer": 0, "lead developer": 0, "testing engineer": 0},
        "P\\Special": {"developer": 1, "lead developer": 0, "testing engineer": 0},
    }


@pytest.mark.parametrize(
    ("value", "approved"),
    [(True, True), (1, True), (1.0, True), ("TRUE", True), ("true", True), ("1", True)]
    + [(False, False), (0, False), ("yes", False), ("FALSE", False), (None, False)],
)
def test_is_approved(value, approved):
    assert team_roster.is_approved(value) is approved


def test_without_approved_column_every_row_counts():
    df = pd.DataFrame(
        {"email": ["a", "b"], "area_path": ["P", "P"], "classified_title": ["developer"] * 2}
    )
    assert team_roster.headcount_pivot(df).to_dict("records") == [
        {"area_path": "P", "developer": 2}
    ]


def test_blank_email_does_not_match_a_nan_key():
    book = {"person_to_area": pd.DataFrame({"email": ["nan"], "area_path": ["P\\Wrong"]})}
    team_roster.ensure_mapping_sheets(book, {})
    df = pd.DataFrame({"email": [None], "corporatetitle": ["Developer"]})
    stats = team_roster.assign_area_paths(df, book)
    assert df["area_path"].isna().all()
    assert stats.counts["no_matches"] == 1


def test_existing_classified_title_is_kept(tmp_path):
    staff = workbook(
        tmp_path / "staff.xlsx",
        {"s": {"email": ["a"], "corporatetitle": ["Developer"], "classified_title": ["custom"]}},
    )
    output = tmp_path / "out.xlsx"
    assert run(staff, output) == 0
    assert sheets(output)["s"]["classified_title"].tolist() == ["custom"]


def test_missing_mapping_sheets_are_created_and_widths_stay_valid(tmp_path):
    staff = workbook(tmp_path / "staff.xlsx", {"s": {"email": ["a"], "corporatetitle": [None]}})
    output = tmp_path / "out.xlsx"
    assert run(staff, output) == 0
    wb = load_workbook(output)  # a NaN width made the original script's files unreadable
    assert wb.sheetnames == ["s", "dev_per_area", *(m.sheet for m in team_roster.MAPPINGS)]
    assert wb["fullpath_to_area"].column_dimensions["A"].width == len("full_path") + 2
    assert wb["s"].column_dimensions["B"].width == len("corporatetitle") + 2
    assert wb["dev_per_area"].max_row == 1  # nobody has an area: empty pivot


def test_column_width_is_capped():
    assert team_roster.column_width(pd.Series(["x" * 80]), "h") == team_roster.MAX_COLUMN_WIDTH


def test_formula_like_titles_stay_text(tmp_path):
    staff = tmp_path / "staff.xlsx"
    wb = Workbook()
    wb.active.append(["email", "corporatetitle"])
    wb.active.append(["a", "=1+1"])
    wb.active["B2"].data_type = "s"
    wb.save(staff)
    output = tmp_path / "out.xlsx"
    assert run(staff, output) == 0
    cell = load_workbook(output)["Sheet"]["B2"]
    assert (cell.value, cell.data_type) == ("=1+1", "s")


def test_no_corporatetitle_sheet_fails_and_writes_nothing(tmp_path):
    staff = workbook(tmp_path / "staff.xlsx", {"s": {"email": ["a"]}})
    output = tmp_path / "out.xlsx"
    assert run(staff, output) == 1
    assert not output.exists()


@pytest.mark.parametrize("missing", ["input", "mapping"])
def test_missing_inputs_fail(inputs, tmp_path, missing):
    staff, mapping = inputs
    args = {"input": (tmp_path / "nope.xlsx", "-m", mapping), "mapping": (staff, "-m", "no.xlsx")}
    source, flag, mapfile = args[missing]
    assert run(source, tmp_path / "out.xlsx", flag, mapfile) == 1
    assert not (tmp_path / "out.xlsx").exists()


def test_existing_output_needs_force_when_not_interactive(inputs, tmp_path, monkeypatch):
    staff, _ = inputs
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"old")
    monkeypatch.setattr("sys.stdin", io.StringIO(""))
    assert run(staff, output) == 1
    assert output.read_bytes() == b"old"
    assert run(staff, output, "--force") == 0
    assert "dev_per_area" in sheets(output)


@pytest.mark.parametrize(("answer", "replaced"), [("y", True), ("n", False)])
def test_interactive_prompt_decides_overwrite(inputs, tmp_path, monkeypatch, answer, replaced):
    staff, _ = inputs
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"old")
    monkeypatch.setattr("sys.stdin.isatty", lambda: True, raising=False)
    monkeypatch.setattr("builtins.input", lambda prompt: answer)
    assert run(staff, output) == 0
    assert (output.read_bytes() != b"old") is replaced


def test_output_may_not_replace_an_input(inputs):
    staff, mapping = inputs
    before = staff.read_bytes()
    assert run(staff, staff, "--force") == 1
    assert run(staff, mapping, "-m", mapping, "--force") == 1
    assert staff.read_bytes() == before


def test_dry_run_writes_nothing(inputs, tmp_path, caplog):
    staff, _ = inputs
    output = tmp_path / "out.xlsx"
    with caplog.at_level("INFO", logger="team_roster"):
        assert run(staff, output, "--dry-run") == 0
    assert not output.exists()
    assert "'Lead Developer' -> 'lead developer'" in caplog.text
    assert "Rows with approved=TRUE: 5/6" in caplog.text


def test_input_is_not_modified(inputs, tmp_path):
    staff, mapping = inputs
    before = (staff.read_bytes(), mapping.read_bytes())
    assert run(staff, tmp_path / "out.xlsx", "-m", mapping) == 0
    assert (staff.read_bytes(), mapping.read_bytes()) == before
