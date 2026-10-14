"""Repositories from a sheet column are matched to their closest parent in a tree."""

import pytest
from openpyxl import Workbook, load_workbook

from workflow_automations import gerrit_parents

TREE = """\
* all-projects
   * some-projects
      * repo1
   * other-projects
      * repo3
      * nested-other-projects
         * repo2
"""


def workbook(tmp_path, rows):
    path = tmp_path / "repos.xlsx"
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.active.title = "Repos"
    wb.save(path)
    return path


def tree(tmp_path, text=TREE):
    path = tmp_path / "tree.md"
    path.write_text(text, encoding="utf-8")
    return path


def run(capsys, *argv):
    status = gerrit_parents.main([str(arg) for arg in argv])
    captured = capsys.readouterr()
    return status, captured.out, captured.err


def table(path):
    wb = load_workbook(path)
    rows = [list(row) for row in wb.active.iter_rows(values_only=True)]
    wb.close()
    return rows


def test_example_from_the_request(tmp_path, capsys):
    source = workbook(tmp_path, [["repos"], ["repo1"], ["repo2"]])
    status, out, err = run(capsys, "--input", source, "--tree", tree(tmp_path))
    assert status == 0
    assert out == "- some-projects\n  - repo1\n- nested-other-projects\n  - repo2\n"
    assert "Matched 2 repositories to 2 parents" in err


def test_workbook_keeps_columns_and_adds_parent_path_status(tmp_path, capsys):
    source = workbook(
        tmp_path, [["repos", "Owner"], ["repo2", "Ann"], [None, "Bob"], ["repo1", None]]
    )
    output = tmp_path / "out.xlsx"
    status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path), "--output", output)
    assert status == 0
    assert out == ""
    assert table(output) == [
        ["repos", "Owner", "Parent", "Path", "Status"],
        [
            "repo2",
            "Ann",
            "nested-other-projects",
            "all-projects / other-projects / nested-other-projects",
            "matched",
        ],
        [None, "Bob", None, None, None],
        ["repo1", None, "some-projects", "all-projects / some-projects", "matched"],
    ]
    assert load_workbook(output).sheetnames == ["Repos"]


def test_outline_deduplicates_and_follows_tree_order(tmp_path, capsys):
    source = workbook(tmp_path, [["repos"], ["repo2"], ["repo3"], ["repo1"], ["repo2"]])
    outline = tmp_path / "parents.md"
    status, _, _ = run(capsys, "--input", source, "--tree", tree(tmp_path), "--outline", outline)
    assert status == 0
    assert outline.read_text(encoding="utf-8") == (
        "- some-projects\n  - repo1\n"
        "- other-projects\n  - repo3\n"
        "- nested-other-projects\n  - repo2\n"
    )


def test_unresolved_names_are_reported_not_guessed(tmp_path, capsys):
    text = TREE + "   * third\n      * repo2\n"
    source = workbook(tmp_path, [["repos"], ["repo1"], ["repo2"], ["ghost"], ["all-projects"]])
    output = tmp_path / "out.xlsx"
    status, out, err = run(
        capsys,
        "--input",
        source,
        "--tree",
        tree(tmp_path, text),
        "--output",
        output,
        "--outline",
        tmp_path / "out.md",
    )
    assert status == 4
    assert (tmp_path / "out.md").read_text(encoding="utf-8") == (
        "- some-projects\n  - repo1\n"
        "- (top level, no parent)\n  - all-projects\n"
        "- (under more than one parent)\n  - repo2\n"
        "    - all-projects / other-projects / nested-other-projects\n"
        "    - all-projects / third\n"
        "- (not found in tree)\n  - ghost\n"
    )
    rows = table(output)
    assert rows[2] == [
        "repo2",
        None,
        "all-projects / other-projects / nested-other-projects; all-projects / third",
        "ambiguous",
    ]
    assert rows[3] == ["ghost", None, None, "not found"]
    assert rows[4] == ["all-projects", None, None, "no parent"]
    assert "Under more than one parent (1): repo2" in err
    assert "Not found in tree (1): ghost" in err


def test_duplicate_listing_under_same_parent_is_one_project(tmp_path, capsys):
    text = TREE + "   * some-projects\n      * repo1\n"
    source = workbook(tmp_path, [["repos"], ["repo1"]])
    status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path, text))
    assert status == 0
    assert out == "- some-projects\n  - repo1\n"


def test_parents_follow_their_own_tree_position(tmp_path, capsys):
    # ParentA's own repository comes after the subtree of ParentB, nested inside it.
    text = "- All-Projects\n  - ParentA\n    - ParentB\n      - repoB\n    - repoA\n"
    for order in (["repoA", "repoB"], ["repoB", "repoA"]):
        source = workbook(tmp_path, [["repos"], *([name] for name in order)])
        status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path, text))
        assert status == 0
        assert out == "- ParentA\n  - repoA\n- ParentB\n  - repoB\n"


@pytest.mark.parametrize("children_first", [False, True])
def test_duplicate_listing_with_children_is_not_a_leaf(tmp_path, capsys, children_first):
    listings = ["  - repoA\n", "  - repoA\n    - repoB\n"]
    if children_first:
        listings.reverse()
    source = workbook(tmp_path, [["repos"], ["repoA"]])
    text = "- All-Projects\n" + "".join(listings)
    status, out, err = run(capsys, "--input", source, "--tree", tree(tmp_path, text))
    assert status == 0
    assert out == "- All-Projects\n  - repoA\n"
    assert "have children of their own (1): repoA" in err


def test_same_parent_name_in_two_branches_shows_paths(tmp_path, capsys):
    text = "- a\n  - common\n    - x\n- b\n  - common\n    - y\n"
    source = workbook(tmp_path, [["repos"], ["x"], ["y"]])
    status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path, text))
    assert status == 0
    assert out == "- a / common\n  - x\n- b / common\n  - y\n"


def test_headings_escapes_and_front_matter_in_tree(tmp_path, capsys):
    text = (
        '---\nxlsx-outline:\n  levels: ["A"]\n---\n'
        "# Projects\n\n## all\\_projects\n\n- tools\n  - my\\_repo\n"
    )
    source = workbook(tmp_path, [["Name"], ["my_repo"], ["tools"]])
    status, out, err = run(
        capsys, "--input", source, "--tree", tree(tmp_path, text), "--title", "--column", "Name"
    )
    assert status == 0
    assert out == "- all\\_projects\n  - tools\n- tools\n  - my\\_repo\n"
    assert "have children of their own (1): tools" in err


def test_column_selection_and_outputs_are_protected(tmp_path, capsys):
    source = workbook(tmp_path, [["Owner", "repos"], ["Ann", "repo1"]])
    status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path), "--column", "repos")
    assert (status, out) == (0, "- some-projects\n  - repo1\n")
    status, _, err = run(capsys, "--input", source, "--tree", tree(tmp_path), "--column", "x")
    assert status == 2 and "Unknown columns: x" in err
    existing = tmp_path / "out.md"
    existing.write_text("keep", encoding="utf-8")
    status, _, err = run(
        capsys,
        "--input",
        source,
        "--tree",
        tree(tmp_path),
        "--column",
        "repos",
        "--outline",
        existing,
    )
    assert status == 2 and "--overwrite" in err
    assert existing.read_text(encoding="utf-8") == "keep"
    status, _, err = run(capsys, "--input", source, "--tree", tree(tmp_path), "--output", source)
    assert status == 2 and "new files" in err


@pytest.mark.parametrize(
    ("rows", "message"),
    [
        ([["repos"], [None]], "no repository names"),
        ([["repos", "Parent"], ["repo1", None]], "already has columns named Parent"),
    ],
)
def test_validation(tmp_path, capsys, rows, message):
    source = workbook(tmp_path, rows)
    status, _, err = run(
        capsys, "--input", source, "--tree", tree(tmp_path), "--output", tmp_path / "o.xlsx"
    )
    assert status == 2
    assert message in err


def test_tree_must_be_an_outline(tmp_path, capsys):
    source = workbook(tmp_path, [["repos"], ["repo1"]])
    status, _, err = run(capsys, "--input", source, "--tree", tree(tmp_path, "just prose\n"))
    assert status == 2
    assert "neither a heading nor a list item" in err


PATHS_TREE = """\
- All-Projects
  - platform/parent
    - platform/tools/repo
  - device/parent
    - repo
    - vendor/lib
"""


def test_full_path_entry_beats_same_named_repository_elsewhere(tmp_path, capsys):
    source = workbook(tmp_path, [["repos"], ["platform/tools/repo"], ["repo"]])
    output = tmp_path / "out.xlsx"
    status, out, err = run(
        capsys, "--input", source, "--tree", tree(tmp_path, PATHS_TREE), "--output", output
    )
    assert status == 0
    assert table(output)[1:] == [
        ["platform/tools/repo", "platform/parent", "All-Projects / platform/parent", "matched"],
        ["repo", "device/parent", "All-Projects / device/parent", "matched"],
    ]
    assert "last path segment" not in err


def test_last_segment_matches_path_and_bare_names(tmp_path, capsys):
    # A bare name finds a path ending in it; a path finds a bare item named after it.
    source = workbook(tmp_path, [["repos"], ["lib"], ["other/tools/repo"]])
    status, out, err = run(capsys, "--input", source, "--tree", tree(tmp_path, PATHS_TREE))
    assert status == 0
    assert out == "- device/parent\n  - other/tools/repo\n  - lib\n"
    assert "Matched by last path segment (2): lib (vendor/lib), other/tools/repo (repo)" in err


def test_different_full_paths_are_different_projects(tmp_path, capsys):
    source = workbook(tmp_path, [["repos"], ["other/lib"]])
    status, out, _ = run(capsys, "--input", source, "--tree", tree(tmp_path, PATHS_TREE))
    assert status == 4
    assert out == "- (not found in tree)\n  - other/lib\n"


def test_last_segment_under_several_paths_is_ambiguous(tmp_path, capsys):
    text = "- All-Projects\n  - a/repo\n  - b/repo\n- c/parent\n  - c/repo\n"
    source = workbook(tmp_path, [["repos"], ["repo"]])
    output = tmp_path / "out.xlsx"
    status, out, _ = run(
        capsys,
        "--input",
        source,
        "--tree",
        tree(tmp_path, text),
        "--output",
        output,
        "--outline",
        tmp_path / "out.md",
    )
    assert status == 4
    assert (tmp_path / "out.md").read_text(encoding="utf-8") == (
        "- (under more than one parent)\n  - repo\n"
        "    - All-Projects (a/repo)\n    - All-Projects (b/repo)\n    - c/parent (c/repo)\n"
    )
    assert table(output)[1][2:] == [
        "All-Projects (a/repo); All-Projects (b/repo); c/parent (c/repo)",
        "ambiguous",
    ]
