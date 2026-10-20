"""Shared Excel cell safety and all-or-nothing publication for the ADO tools."""

import datetime as dt

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from workflow_automations import ado_triage
from workflow_automations.ado_common import excel
from workflow_automations.ado_graph import report

PLUS_TWO = dt.timezone(dt.timedelta(hours=2))


def test_excel_safe_strips_control_characters_from_text_only():
    frame = pd.DataFrame({"Title": ["ok", "bell\x07 and\x1b esc", None], "N": [1, 2, 3]})
    safe = excel.excel_safe(frame)
    assert safe["Title"].tolist()[:2] == ["ok", "bell and esc"]
    assert pd.isna(safe["Title"].iloc[2])
    assert safe["N"].tolist() == [1, 2, 3]
    assert frame["Title"].iloc[1] == "bell\x07 and\x1b esc"  # the input is not modified


def test_excel_safe_makes_zoned_datetimes_naive_utc():
    zoned = dt.datetime(2026, 10, 14, 12, 0, tzinfo=PLUS_TWO)
    frame = pd.DataFrame(
        {
            "Column": pd.to_datetime(["2026-10-14T12:00:00+02:00"], utc=True),
            "Mixed": pd.Series([zoned], dtype=object),
        }
    )
    frame.loc[1] = [pd.NaT, "not a date"]
    safe = excel.excel_safe(frame)
    assert safe["Column"].iloc[0] == pd.Timestamp("2026-10-14 10:00")
    assert safe["Column"].dt.tz is None
    assert safe["Mixed"].tolist() == [pd.Timestamp("2026-10-14 10:00"), "not a date"]


def test_excel_safe_returns_a_clean_frame_itself():
    frame = pd.DataFrame({"Title": ["a", "b"], "When": [dt.datetime(2026, 1, 1)] * 2})
    assert excel.excel_safe(frame) is frame


def test_excel_safe_output_round_trips_through_openpyxl(tmp_path):
    frame = pd.DataFrame(
        {
            "Title": ["tab\tand\x0bvertical"],
            "When": [dt.datetime(2026, 10, 14, 12, 0, tzinfo=PLUS_TWO)],
        }
    )
    path = tmp_path / "safe.xlsx"
    excel.excel_safe(frame).to_excel(path, index=False)
    back = pd.read_excel(path)
    assert back.loc[0, "Title"] == "tab\tandvertical"
    assert back.loc[0, "When"] == pd.Timestamp("2026-10-14 10:00")


def test_keep_text_literal_leaves_headings_and_stores_formula_text_as_text(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["=heading"])
    sheet.append(["=SUM(A1:A2)"])
    excel.keep_text_literal(sheet, min_row=2)
    assert sheet["A1"].data_type == "f"
    assert sheet["A2"].data_type == "s" and sheet["A2"].quotePrefix


def test_publish_all_refuses_existing_targets_with_the_callers_error(tmp_path):
    existing, fresh = tmp_path / "a.txt", tmp_path / "b.txt"
    existing.write_text("old")

    class ToolError(Exception):
        pass

    def write(path):
        path.write_text("new")

    with pytest.raises(ToolError, match="pass --overwrite"):
        excel.publish_all([(fresh, write), (existing, write)], False, exists_error=ToolError)
    with pytest.raises(FileExistsError):
        excel.publish_all([(existing, write)], False)
    assert existing.read_text() == "old" and not fresh.exists()


def test_exclusive_publish_never_replaces_a_target_created_while_staging(tmp_path, monkeypatch):
    target = tmp_path / "a.txt"

    def write_and_race(path):
        path.write_text("ours")
        target.write_text("theirs")

    with pytest.raises(FileExistsError, match="not replaced"):
        excel.publish_all([(target, write_and_race)], False, exclusive=True)
    assert target.read_text() == "theirs"
    assert [p.name for p in tmp_path.iterdir()] == ["a.txt"]

    target.unlink()
    monkeypatch.setattr(excel.os, "link", lambda *a: (_ for _ in ()).throw(PermissionError()))
    excel.publish_all([(target, lambda p: p.write_text("ours"))], False, exclusive=True)
    assert target.read_text() == "ours"  # no hard links: falls back to a checked rename
    assert [p.name for p in tmp_path.iterdir()] == ["a.txt"]


def test_publish_all_failure_leaves_every_target_and_no_temp_files(tmp_path):
    first, second = tmp_path / "a.txt", tmp_path / "out" / "b.txt"
    first.write_text("old")

    def write(path):
        path.write_text("new")

    def fail(path):
        path.write_text("partial")
        raise OSError("disk full")

    with pytest.raises(OSError, match="disk full"):
        excel.publish_all([(first, write), (second, fail)], True)
    assert first.read_text() == "old" and not second.exists()
    assert sorted(p.name for p in tmp_path.rglob("*") if p.is_file()) == ["a.txt"]

    excel.publish_all([(first, write), (second, write)], True)
    assert first.read_text() == second.read_text() == "new"


def test_both_tools_publish_through_the_shared_primitive(tmp_path):
    assert report.publish_all is excel.publish_all
    target = tmp_path / "r.xlsx"
    target.write_text("old")
    with pytest.raises(ado_triage.TriageError, match="pass --overwrite"):
        ado_triage.publish_all([(target, lambda p: None)], False)
    assert target.read_text() == "old"


def test_graph_export_writes_control_characters_and_formula_text_safely(tmp_path):
    nodes = pd.DataFrame({"WI_ID": [1, 2], "Title": ["=1+1", "bell\x07"]})
    edges = pd.DataFrame({"Source_ID": [1], "Target_ID": [2], "Link_Type": ["Child"]})
    report._export_dataframes(nodes, edges, tmp_path, ["excel"])
    sheet = load_workbook(tmp_path / "work_items.xlsx")["Nodes"]
    assert [sheet["B2"].value, sheet["B3"].value] == ["=1+1", "bell"]
    assert sheet["B2"].data_type == "s"


def test_triage_report_keeps_formula_like_titles_as_text(tmp_path):
    frame = pd.DataFrame({"System.Title": ['=HYPERLINK("x")', "ok\x07"]})
    special = {ado_triage.SHEET_SUMMARY, ado_triage.SHEET_AGENDA, "18. Stale Items"}
    name = next(spec.name for spec in ado_triage.SHEETS if spec.name not in special)
    built = ado_triage.Report(
        master=frame, relations=pd.DataFrame(), sheets={name: frame}, omitted=[], notes=[]
    )
    path = tmp_path / "report.xlsx"
    ado_triage.write_excel(path, built, {})
    sheet = load_workbook(path)[name]
    first = ado_triage.DATA_START_ROW
    assert sheet.cell(row=first, column=1).value == '=HYPERLINK("x")'
    assert sheet.cell(row=first, column=1).data_type == "s"
    assert sheet.cell(row=first + 1, column=1).value == "ok"
