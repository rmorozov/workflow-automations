"""Offline behavior tests for the ado-graph analysis workbook."""

import pytest
from ado_graph_data import edges, fake_client, nodes, run_export, wi
from openpyxl import load_workbook

from workflow_automations import ado_graph as tool
from workflow_automations.ado_graph import analyze, report


def write_inputs(tmp_path):
    node_file, edge_file = tmp_path / "nodes.parquet", tmp_path / "edges.parquet"
    nodes().to_parquet(node_file, index=False)
    edges().to_parquet(edge_file, index=False)
    return node_file, edge_file


def run_analyze(node_file, edge_file, output, *extra):
    tool.main(
        [
            "analyze",
            "--nodes-file",
            str(node_file),
            "--edges-file",
            str(edge_file),
            "--output-file",
            str(output),
            "--effort-field",
            "Effort",
            "--steps",
            "orphans",
            *extra,
        ]
    )


def test_analyze_parquet_to_all_report_sheets(tmp_path):
    node_file, edge_file = tmp_path / "nodes.parquet", tmp_path / "edges.parquet"
    nodes().to_parquet(node_file, index=False)
    edges().to_parquet(edge_file, index=False)
    output = tmp_path / "report.xlsx"
    assert (
        tool.main(
            [
                "analyze",
                "--nodes-file",
                str(node_file),
                "--edges-file",
                str(edge_file),
                "--output-file",
                str(output),
                "--effort-field",
                "Effort",
            ]
        )
        is None
    )
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == list(analyze.ALL_STEPS)
        assert workbook["orphans"]["A2"].value == 4
        assert workbook["cr_scope_impact"]["C2"].value == 5
    finally:
        workbook.close()


def test_analyze_refuses_existing_report_without_overwrite(tmp_path, capsys):
    node_file, edge_file = write_inputs(tmp_path)
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"keep me")
    with pytest.raises(SystemExit) as exit_info:
        run_analyze(node_file, edge_file, output)
    assert exit_info.value.code == 2
    assert "--overwrite" in capsys.readouterr().err
    assert output.read_bytes() == b"keep me"

    run_analyze(node_file, edge_file, output, "--overwrite")
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["orphans"]
    finally:
        workbook.close()


@pytest.mark.parametrize("which", ["nodes", "edges"])
def test_analyze_never_writes_over_an_input(tmp_path, which):
    node_file, edge_file = write_inputs(tmp_path)
    target = node_file if which == "nodes" else edge_file
    before = target.read_bytes()
    alias = tmp_path / "alias.parquet"
    alias.symlink_to(target)
    for output in (target, alias):
        with pytest.raises(SystemExit) as exit_info:
            run_analyze(node_file, edge_file, output, "--overwrite")
        assert exit_info.value.code == 2
    assert target.read_bytes() == before


def test_failed_analysis_write_keeps_existing_report(tmp_path, monkeypatch):
    node_file, edge_file = write_inputs(tmp_path)
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"previous report")

    def fail(_results, path):
        path.write_bytes(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(report, "_write_workbook", fail)
    with pytest.raises(OSError, match="disk full"):
        run_analyze(node_file, edge_file, output, "--overwrite")
    assert output.read_bytes() == b"previous report"
    assert sorted(p.name for p in tmp_path.iterdir()) == [
        "edges.parquet",
        "nodes.parquet",
        "report.xlsx",
    ]


def test_export_refuses_existing_files_before_fetching(tmp_path, monkeypatch, capsys):
    client, log = fake_client({1: wi(1)}, [1])
    (tmp_path / "edges.parquet").write_bytes(b"keep me")
    with pytest.raises(SystemExit) as exit_info:
        run_export(monkeypatch, tmp_path, client)
    assert exit_info.value.code == 2
    assert "edges.parquet exists" in capsys.readouterr().err
    assert log["fetched"] == []
    assert (tmp_path / "edges.parquet").read_bytes() == b"keep me"

    exported, _ = run_export(monkeypatch, tmp_path, client, "--overwrite")
    assert exported["WI_ID"].tolist() == [1]


def test_failed_export_leaves_the_whole_previous_bundle(tmp_path, monkeypatch):
    previous = {name: name.encode() for name in ("nodes.csv", "edges.csv", "work_items.xlsx")}
    for name, content in previous.items():
        (tmp_path / name).write_bytes(content)

    def fail(_nodes, _edges, path):
        path.write_bytes(b"partial")
        raise ValueError("cannot serialize")

    monkeypatch.setattr(report, "_write_excel", fail)
    with pytest.raises(ValueError, match="cannot serialize"):
        report._export_dataframes(nodes(), edges(), tmp_path, ["csv", "excel"])
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == previous


def test_export_publishes_each_format_once_and_skips_unknown(tmp_path):
    report._export_dataframes(nodes(), edges(), tmp_path, ["csv", "bogus", "csv"])
    assert sorted(p.name for p in tmp_path.iterdir()) == ["edges.csv", "nodes.csv"]
    assert report.export_paths(tmp_path, ["excel", "bogus", "excel"]) == [
        tmp_path / "work_items.xlsx"
    ]
