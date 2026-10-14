"""Offline behavior tests for the ado-graph analysis workbook."""

from ado_graph_data import edges, nodes
from openpyxl import load_workbook

from workflow_automations import ado_graph as tool
from workflow_automations.ado_graph import analyze


def test_analyze_parquet_to_all_report_sheets(tmp_path):
    node_file, edge_file = tmp_path / "nodes.parquet", tmp_path / "edges.parquet"
    nodes().to_parquet(node_file, index=False)
    edges().to_parquet(edge_file, index=False)
    output = tmp_path / "report.xlsx"
    assert (
        tool.main(
            [
                "analyze",
                "--nodes-file",
                str(node_file),
                "--edges-file",
                str(edge_file),
                "--output-file",
                str(output),
                "--effort-field",
                "Effort",
            ]
        )
        is None
    )
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == list(analyze.ALL_STEPS)
        assert workbook["orphans"]["A2"].value == 4
        assert workbook["cr_scope_impact"]["C2"].value == 5
    finally:
        workbook.close()
