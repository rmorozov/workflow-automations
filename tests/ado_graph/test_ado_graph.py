"""Offline behavior tests for the ADO graph/export tool using SDK-like fakes."""

from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

from workflow_automations import ado_graph as tool


def nodes():
    return pd.DataFrame(
        [
            {"WI_ID": 1, "Title": "CR", "WorkItemType": "Change Request", "Effort": 2},
            {"WI_ID": 2, "Title": "Child", "WorkItemType": "Feature", "Effort": 3},
            {"WI_ID": 3, "Title": "Prerequisite", "WorkItemType": "Feature", "Effort": 5},
            {"WI_ID": 4, "Title": "Isolated", "WorkItemType": "Feature", "Effort": 7},
        ]
    )


def edges():
    return pd.DataFrame(
        [(1, 2, "Parent->Child"), (1, 3, "Successor->Predecessor"), (1, 99, "Blocks")],
        columns=["Source_ID", "Target_ID", "Link_Type"],
    )


def test_reverse_edges_and_absent_endpoints():
    analyzer = tool.GraphAnalyzer(nodes(), edges(), "https://ado.example", "Effort")
    assert set(analyzer.G.edges) == {(1, 2), (3, 1)}
    assert analyzer.G[3][1]["link_type"] == "Predecessor->Successor"
    assert set(analyzer.find_orphans()["WI_ID"]) == {4}
    assert set(analyzer.analyze_cr_scope_and_impact().iloc[0]["Impacted_WI_IDs"].split(",")) == {
        "2"
    }
    assert analyzer.analyze_cr_scope_and_impact().iloc[0]["Total_Effort"] == 5
    assert analyzer.analyze_cr_prerequisites().iloc[0]["Prerequisite_WI_IDs"] == "3"
    # Existing has_cr_dep means CR + ancestors, unlike downstream scope/impact.
    assert set(analyzer.find_has_cr_dep()["WI_ID"]) == {1, 3}


def test_parallel_relation_labels_currently_collapse():
    links = pd.DataFrame([(1, 2, "Parent->Child"), (1, 2, "Blocks")], columns=edges().columns)
    analyzer = tool.GraphAnalyzer(nodes(), links)
    assert analyzer.G.number_of_edges() == 1
    assert analyzer.G[1][2]["link_type"] == "Blocks"


def test_dependency_depth_is_not_longest_chain_depth():
    links = pd.DataFrame(
        [(1, 2, "Blocks"), (2, 3, "Blocks"), (1, 3, "Blocks")], columns=edges().columns
    )
    analyzer = tool.GraphAnalyzer(nodes(), links)
    metrics = analyzer.compute_metrics().set_index("WI_ID")
    assert metrics.loc[3, "Dependency_Depth"] == 1
    assert analyzer.find_longest_chain().attrs["longest_path_length"] == 2


def test_analyze_parquet_to_all_report_sheets(tmp_path):
    node_file, edge_file = tmp_path / "nodes.parquet", tmp_path / "edges.parquet"
    nodes().to_parquet(node_file, index=False)
    edges().to_parquet(edge_file, index=False)
    output = tmp_path / "report.xlsx"
    assert (
        tool.main(
            [
                "analyze",
                "--nodes-file",
                str(node_file),
                "--edges-file",
                str(edge_file),
                "--output-file",
                str(output),
                "--effort-field",
                "Effort",
            ]
        )
        is None
    )
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == list(tool.ALL_STEPS)
        assert workbook["orphans"]["A2"].value == 4
        assert workbook["cr_scope_impact"]["C2"].value == 5
    finally:
        workbook.close()


def test_saved_query_flat_and_link_results():
    client = object.__new__(tool.ADOClient)
    flat = SimpleNamespace(work_items=[SimpleNamespace(id=5)], work_item_relations=None)
    client._wit_client = SimpleNamespace(query_by_id=lambda **kwargs: flat)
    assert client.run_query_by_id("Example", "query") == [5]
    links = SimpleNamespace(
        work_items=None,
        work_item_relations=[
            SimpleNamespace(source=None, target=SimpleNamespace(id=1)),
            SimpleNamespace(source=SimpleNamespace(id=1), target=SimpleNamespace(id=2)),
        ],
    )
    client._wit_client = SimpleNamespace(query_by_id=lambda **kwargs: links)
    assert client.run_query_by_id("Example", "query") == [1, 2]


def test_export_filtered_bridge_is_pruned_before_traversal(tmp_path, monkeypatch):
    """AG-001: accepted 4 is lost when the skipped bridge 3 is pruned."""
    calls = []

    def item(wi_id, wi_type="Feature", targets=()):
        return SimpleNamespace(
            id=wi_id,
            fields={"System.WorkItemType": wi_type, "System.TeamProject": "Example"},
            relations=[
                SimpleNamespace(rel="System.LinkTypes.Hierarchy-Forward", url=f"https://x/{i}")
                for i in targets
            ],
        )

    class Client:
        def __init__(self, _url):
            pass

        def run_query_by_id(self, _project, _query):
            return [1, 2, 3]

        def filter_ids_via_wiql(self, ids, _projects, _types):
            return set(ids) - {3}

        def get_work_items_batch(self, ids):
            calls.extend(ids)
            items = {1: item(1), 2: item(2), 3: item(3, "Task", [4]), 4: item(4)}
            return [items[i] for i in ids]

    monkeypatch.setattr(tool, "ADOClient", Client)
    monkeypatch.setattr(tool, "configure_ca_bundle", lambda _path: None)
    tool.main(
        [
            "export",
            "--ado-url",
            "https://ado.example",
            "--query-id",
            "synthetic",
            "--batch-size",
            "1",
            "--output-dir",
            str(tmp_path),
            "--formats",
            "parquet",
        ]
    )
    exported = pd.read_parquet(tmp_path / "nodes.parquet")
    assert calls == [1, 2]
    assert set(exported["WI_ID"]) == {1, 2}
    assert 4 not in set(exported["WI_ID"])


def test_queue_survivors_are_repeatedly_sent_to_wiql(tmp_path, monkeypatch):
    """AG-014: quantify repeated work independently of filtered-bridge loss."""
    filtered = []
    fetched = []

    class Client:
        def __init__(self, _url):
            pass

        def run_query_by_id(self, _project, _query):
            return list(range(1, 2001))

        def filter_ids_via_wiql(self, ids, _projects, _types):
            filtered.extend(ids)
            return set(ids)

        def get_work_items_batch(self, ids):
            fetched.extend(ids)
            return [
                SimpleNamespace(
                    id=wi_id,
                    fields={"System.WorkItemType": "Feature", "System.TeamProject": "Example"},
                    relations=[],
                )
                for wi_id in ids
            ]

    monkeypatch.setattr(tool, "ADOClient", Client)
    monkeypatch.setattr(tool, "configure_ca_bundle", lambda _path: None)
    tool.main(
        [
            "export",
            "--ado-url",
            "https://ado.example",
            "--query-id",
            "synthetic",
            "--batch-size",
            "100",
            "--output-dir",
            str(tmp_path),
            "--formats",
            "parquet",
        ]
    )
    assert len(fetched) == len(set(fetched)) == 2000
    assert len(filtered) == 18900
    assert len(set(filtered)) == 1900


def test_filtered_bridge_breaks_analysis_even_without_pruning(tmp_path, monkeypatch):
    """AG-015: traversal reaches the descendant, but exported edges lose the path."""
    fetched = []

    class Client:
        def __init__(self, _url):
            pass

        def run_query_by_id(self, _project, _query):
            return [1]

        def filter_ids_via_wiql(self, *_args):
            raise AssertionError("Small-queue fixture must not use pruning")

        def get_work_items_batch(self, ids):
            fetched.extend(ids)
            return [
                SimpleNamespace(
                    id=wi_id,
                    fields={
                        "System.WorkItemType": "Task" if wi_id == 2 else "Feature",
                        "System.TeamProject": "Example",
                    },
                    relations=[
                        SimpleNamespace(
                            rel="System.LinkTypes.Hierarchy-Forward",
                            url=f"https://ado.example/{wi_id + 1}",
                        )
                    ]
                    if wi_id < 3
                    else [],
                )
                for wi_id in ids
            ]

    monkeypatch.setattr(tool, "ADOClient", Client)
    monkeypatch.setattr(tool, "configure_ca_bundle", lambda _path: None)
    tool.main(
        [
            "export",
            "--ado-url",
            "https://ado.example",
            "--query-id",
            "synthetic",
            "--output-dir",
            str(tmp_path),
            "--formats",
            "parquet",
        ]
    )
    assert fetched == [1, 2, 3]
    exported_nodes = pd.read_parquet(tmp_path / "nodes.parquet")
    exported_edges = pd.read_parquet(tmp_path / "edges.parquet")
    assert set(exported_nodes["WI_ID"]) == {1, 3}
    assert list(exported_edges[["Source_ID", "Target_ID"]].itertuples(index=False, name=None)) == [
        (1, 2)
    ]
    analyzer = tool.GraphAnalyzer(exported_nodes, exported_edges)
    assert analyzer.G.number_of_edges() == 0
    assert set(analyzer.find_orphans()["WI_ID"]) == {1, 3}
