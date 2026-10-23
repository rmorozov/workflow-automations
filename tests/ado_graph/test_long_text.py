"""Text over Excel's cell limit: split, trim or keep, in export, replay and analyze."""

import pandas as pd
import pytest
from ado_graph_data import edges, fake_client, nodes, run_export, wi
from openpyxl import load_workbook

from workflow_automations import ado_graph as tool
from workflow_automations.ado_graph import export, report

LIMIT = report.EXCEL_MAX_CELL_CHARS
# 80,001 Excel characters: emoji count 2, and one straddles the first part boundary.
LONG = "ab😀" * 20_000 + "x"


def excel_cells(path, sheet):
    return {cell.coordinate: cell for row in load_workbook(path)[sheet].iter_rows() for cell in row}


def test_excel_length_counts_utf16_units_and_cuts_never_split_an_emoji():
    assert report.excel_len("a😀") == 3
    assert report._cut("a😀b", 2) == "a"
    parts = report._parts(LONG, LIMIT)
    assert "".join(parts) == LONG
    lengths = [report.excel_len(p) for p in parts]
    assert lengths[0] == LIMIT - 1  # the emoji at the boundary moves to the next part
    assert sum(lengths) == 80_001 and max(lengths) <= LIMIT


def long_nodes():
    frame = nodes()
    frame["Description"] = ["short", LONG, None, "=" + "y" * LIMIT]
    return frame


def test_export_splits_long_text_into_a_linked_long_text_sheet(tmp_path):
    report._export_dataframes(long_nodes(), edges(), tmp_path, ["csv", "excel"])
    path = tmp_path / "work_items.xlsx"
    assert load_workbook(path).sheetnames == ["Nodes", "Edges", "Long_Text"]

    cells = excel_cells(path, "Nodes")
    trimmed = cells["E3"]
    assert report.excel_len(trimmed.value) <= LIMIT
    assert trimmed.value.startswith("ab😀")
    assert trimmed.value.endswith("[80,001 characters; full text in Long_Text row 2]")
    assert trimmed.hyperlink.location == "'Long_Text'!A2"
    assert cells["E5"].value.endswith("full text in Long_Text row 3]")
    assert cells["E5"].data_type == "s"  # still literal text, never a formula
    assert cells["E2"].value == "short"

    sheet = pd.read_excel(path, sheet_name="Long_Text")
    assert list(sheet.columns[:6]) == ["Sheet", "Cell", "Column", "Key", "Characters", "Parts"]
    assert sheet[["Sheet", "Cell", "Column", "Key"]].values.tolist() == [
        ["Nodes", "E3", "Description", 2],
        ["Nodes", "E5", "Description", 4],
    ]
    assert sheet["Parts"].tolist() == [3, 2]
    first = sheet.iloc[0]
    assert "".join(first[f"Part_{i}"] for i in range(1, 4)) == LONG
    assert first["Characters"] == 80_001
    assert sheet.iloc[1]["Part_1"] + sheet.iloc[1]["Part_2"] == "=" + "y" * LIMIT
    back = excel_cells(path, "Long_Text")
    assert back["B2"].hyperlink.location == "'Nodes'!E3"
    assert back["G3"].data_type == "s"

    # CSV keeps every text whole.
    assert pd.read_csv(tmp_path / "nodes.csv")["Description"][1] == LONG


def test_trim_keeps_the_start_and_a_marker_within_max_cell_chars(tmp_path):
    policy = report.LongText("trim", 300)
    report._export_dataframes(long_nodes(), edges(), tmp_path, ["excel"], long_text=policy)
    path = tmp_path / "work_items.xlsx"
    assert load_workbook(path).sheetnames == ["Nodes", "Edges"]
    value = excel_cells(path, "Nodes")["E3"].value
    assert report.excel_len(value) <= 300
    assert value.endswith(" … [80,001 characters; trimmed]")
    assert excel_cells(path, "Nodes")["E3"].hyperlink is None


def test_keep_refuses_a_long_cell_before_writing(tmp_path):
    with pytest.raises(report.ExcelLimitError, match="--long-text split or trim"):
        report._export_dataframes(
            long_nodes(), edges(), tmp_path, ["csv", "excel"], long_text=report.LongText("keep")
        )
    assert list(tmp_path.iterdir()) == []


def long_client():
    item = wi(1, "Epic", targets=[2])
    item.fields["System.Description"] = LONG
    return fake_client({1: item, 2: wi(2)}, [1])[0]


def test_export_and_replay_take_the_mode_from_options_and_flags(monkeypatch, tmp_path):
    options = tmp_path / "graph.ini"
    options.write_text("long-text = trim\nmax-cell-chars = 1000\n")
    snapshot = tmp_path / "snap.json"
    run_export(
        monkeypatch,
        tmp_path / "live",
        long_client(),
        "--formats",
        "parquet,excel",
        "--options",
        str(options),
        "--save-snapshot",
        str(snapshot),
    )
    workbook = pd.read_excel(tmp_path / "live" / "work_items.xlsx", sheet_name=None)
    assert list(workbook) == ["Nodes", "Edges"]
    text = workbook["Nodes"].set_index("WI_ID").loc[1].astype(str)
    trimmed = [v for v in text if v.endswith("characters; trimmed]")]
    assert len(trimmed) == 1 and report.excel_len(trimmed[0]) <= 1000
    assert LONG in pd.read_parquet(tmp_path / "live" / "nodes.parquet").values

    monkeypatch.setattr(export, "ADOClient", None)  # offline from here
    out = tmp_path / "replayed"
    tool.main(
        ["replay", "--snapshot", str(snapshot), "--output-dir", str(out), "--formats", "excel"]
        + ["--options", str(options), "--long-text", "split"]
    )
    sheet = pd.read_excel(out / "work_items.xlsx", sheet_name="Long_Text")
    assert sheet["Key"].tolist() == [1]
    assert "".join(sheet.iloc[0][[f"Part_{i}" for i in range(1, 4)]]) == LONG


def analyze_long(tmp_path, *extra):
    node_file, edge_file = tmp_path / "nodes.parquet", tmp_path / "edges.parquet"
    long_nodes().to_parquet(node_file, index=False)
    edges().to_parquet(edge_file, index=False)
    output = tmp_path / "report.xlsx"
    argv = ["analyze", "--nodes-file", str(node_file), "--edges-file", str(edge_file)]
    argv += ["--output-file", str(output), "--steps", "orphans", "--meta-field", "Description"]
    tool.main([*argv, *extra])
    return output


def test_analyze_splits_long_meta_text_and_reports_it(tmp_path):
    output = analyze_long(tmp_path, "--max-cell-chars", "500")
    workbook = pd.read_excel(output, sheet_name=None)
    assert list(workbook)[-1] == "Long_Text"
    orphan = workbook["orphans"].iloc[0]
    assert orphan["WI_ID"] == 4
    assert report.excel_len(orphan["Description"]) <= 500
    long = workbook["Long_Text"].iloc[0]
    assert (long["Sheet"], long["Cell"], long["Key"], long["Parts"]) == ("orphans", "G2", 4, 2)
    assert long["Part_1"] + long["Part_2"] == "=" + "y" * LIMIT
    assert orphan["Description"].startswith("=yyy")  # literal text, not a formula
    summary = workbook["Summary"].set_index("Sheet").loc["orphans"]
    assert "1 cell(s) over 500 characters trimmed; full text in Long_Text" in summary["Note"]
    assert bool(summary["Complete"])
    config = workbook["Config"].set_index("Setting")["Value"]
    assert (config["long_text"], str(config["max_cell_chars"])) == ("split", "500")


def test_analyze_trim_marks_the_sheet_incomplete(tmp_path):
    output = analyze_long(tmp_path, "--long-text", "trim")
    workbook = pd.read_excel(output, sheet_name=None)
    assert "Long_Text" not in workbook
    summary = workbook["Summary"].set_index("Sheet").loc["orphans"]
    assert "text cut" in summary["Note"] and not bool(summary["Complete"])


def test_analyze_keep_spills_the_sheet_to_csv_as_before(tmp_path):
    output = analyze_long(tmp_path, "--long-text", "keep")
    workbook = pd.read_excel(output, sheet_name=None)
    assert "Long_Text" not in workbook
    assert workbook["orphans"].iloc[0]["Spilled_To"] == "report.orphans.csv"
    assert pd.read_csv(tmp_path / "report.orphans.csv")["Description"][0] == "=" + "y" * LIMIT


@pytest.mark.parametrize("value", ["255", "32768", "many"])
def test_max_cell_chars_is_checked(tmp_path, value, capsys):
    with pytest.raises(SystemExit) as exit_info:
        analyze_long(tmp_path, "--max-cell-chars", value)
    assert exit_info.value.code == 2
    assert "--max-cell-chars" in capsys.readouterr().err


def test_written_options_name_the_long_text_settings(tmp_path):
    path = tmp_path / "analyze.ini"
    tool.main(["analyze", "--long-text", "trim", "--write-options", str(path)])
    text = path.read_text()
    assert "long-text = trim" in text
    assert f"max-cell-chars = {LIMIT}" in text


def test_links_beyond_the_hyperlink_budget_stay_text(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "EXCEL_MAX_HYPERLINKS", 1)
    report._export_dataframes(long_nodes(), edges(), tmp_path, ["excel"])
    cells = excel_cells(tmp_path / "work_items.xlsx", "Nodes")
    assert cells["E3"].hyperlink is None
    assert cells["E3"].value.endswith("full text in Long_Text row 2]")
